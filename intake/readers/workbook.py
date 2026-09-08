"""The Workbook Reader: finds the sheets that matter - and says what it skipped.

Three stages, and only the middle one may involve a model:

1. **Survey.** Python reads sheet names, visibility, dimensions, the header rows and a
   count of numeric cells. Structure only; the full workbook never enters a model.
2. **Select.** A `SheetSelector` returns which sheets to read, each with a reason, and
   which to skip, each with a reason. `ModelSheetSelector` is the one model call the
   blueprint allows here; `StructuralSelector` is the deterministic default and what a
   test uses. Whatever the selector says, `Selection.accounting_for` makes sure every
   sheet in the survey is either selected or skipped with a reason - silent
   under-selection is the dangerous failure, and a sheet the selector forgot to mention
   is exactly that.
3. **Normalise.** Deterministic code turns each selected sheet into rows: one block per
   numeric column, locator `Sheet!C{first}:C{last}`, each row carrying the label from the
   sheet's label column, the value, and the cell's row number as its ordinal - so a
   figure can cite `TB_FY25!B44:B47` and the resolver finds those four rows.
"""

from __future__ import annotations

import io
import json
import re
from dataclasses import dataclass, field
from typing import Any, Protocol

from ops.source_parse import Row

__all__ = [
    "ModelSheetSelector",
    "NormalizedSheet",
    "ScriptedSelector",
    "Selection",
    "SheetChoice",
    "SheetInfo",
    "SheetSelector",
    "StructuralSelector",
    "WorkbookSurvey",
    "normalize_workbook",
    "survey_workbook",
]

HEADER_ROWS = 12
"""How far down a sheet the title block may run. An accountant\'s statement opens with the
company, the statement, the period and the scale - four or five rows - and the scale cue
is the one that matters most and sits lowest. Header collection stops at the first row
with a number in it, so this is a cap rather than a count.
"""


@dataclass(frozen=True)
class SheetInfo:
    name: str
    state: str
    rows: int
    columns: int
    headers: tuple[str, ...]
    numeric_cells: int
    kind: str = "worksheet"

    @property
    def readable(self) -> bool:
        return self.kind == "worksheet" and self.state == "visible" and self.numeric_cells > 0


@dataclass(frozen=True)
class WorkbookSurvey:
    sheets: tuple[SheetInfo, ...]

    def as_prompt(self) -> str:
        return json.dumps(
            [
                {
                    "sheet": s.name,
                    "state": s.state,
                    "kind": s.kind,
                    "rows": s.rows,
                    "columns": s.columns,
                    "numeric_cells": s.numeric_cells,
                    "headers": list(s.headers),
                }
                for s in self.sheets
            ],
            indent=1,
        )


@dataclass(frozen=True)
class SheetChoice:
    sheet: str
    reason: str


@dataclass(frozen=True)
class Selection:
    selected: tuple[SheetChoice, ...]
    skipped: tuple[SheetChoice, ...]
    by: str = "structural"

    def accounting_for(self, survey: WorkbookSurvey) -> Selection:
        """Every sheet in the survey, selected or skipped with a reason. A sheet the
        selector did not mention is skipped and says so; a sheet it named that does not
        exist is dropped."""
        names = {s.name for s in survey.sheets}
        selected = tuple(c for c in self.selected if c.sheet in names)
        mentioned = {c.sheet for c in selected} | {c.sheet for c in self.skipped}
        skipped = [
            c
            for c in self.skipped
            if c.sheet in names and c.sheet not in {s.sheet for s in selected}
        ]
        for sheet in survey.sheets:
            if sheet.name not in mentioned:
                skipped.append(SheetChoice(sheet.name, f"not selected by the {self.by} selector"))
        return Selection(selected=selected, skipped=tuple(skipped), by=self.by)


class SheetSelector(Protocol):
    def select(self, survey: WorkbookSurvey) -> Selection: ...


class StructuralSelector:
    """No model: every visible worksheet with a number in it is read."""

    def select(self, survey: WorkbookSurvey) -> Selection:
        selected: list[SheetChoice] = []
        skipped: list[SheetChoice] = []
        for sheet in survey.sheets:
            if sheet.kind != "worksheet":
                skipped.append(SheetChoice(sheet.name, f"{sheet.kind}: no cells to read"))
            elif sheet.state != "visible":
                skipped.append(SheetChoice(sheet.name, f"sheet is {sheet.state}"))
            elif sheet.numeric_cells == 0:
                skipped.append(SheetChoice(sheet.name, "no numeric cells"))
            else:
                selected.append(SheetChoice(sheet.name, f"{sheet.numeric_cells} numeric cell(s)"))
        return Selection(tuple(selected), tuple(skipped), by="structural")


class ScriptedSelector:
    """Says what a test told it to say."""

    def __init__(self, selected: dict[str, str], skipped: dict[str, str] | None = None) -> None:
        self._selected = selected
        self._skipped = skipped or {}

    def select(self, survey: WorkbookSurvey) -> Selection:
        return Selection(
            tuple(SheetChoice(s, r) for s, r in self._selected.items()),
            tuple(SheetChoice(s, r) for s, r in self._skipped.items()),
            by="scripted",
        )


class ModelSheetSelector:
    """The one model call: the survey goes in, wrapped as data; a selection with reasons
    comes back. The workbook's cells never do."""

    SYSTEM = (
        "You select which sheets of a financial workbook a valuation analyst needs: trial "
        "balances, income statements, balance sheets, general ledger detail, schedules of "
        "add-backs. Skip charts, notes, instructions, lookup tables and scratch sheets. "
        'Answer with JSON only: {"selected": [{"sheet": name, "reason": why}], '
        '"skipped": [{"sheet": name, "reason": why}]}. Name every sheet exactly once.'
    )

    def __init__(
        self,
        door: Any,
        identifiers: Any,
        boundary: Any,
        pin: Any,
        run_id: str,
        step_id: str | None = None,
    ) -> None:
        self.door = door
        self.identifiers = identifiers
        self.boundary = boundary
        self.pin = pin
        self.run_id = run_id
        self.step_id = step_id
        self.calls = 0

    def select(self, survey: WorkbookSurvey) -> Selection:
        from guards.untrusted import SENTINEL_RULE
        from ports.llm import LLMRequest, Message

        wrapped = self.boundary.wrap(
            survey.as_prompt(),
            source_id="sha256:" + "0" * 64,
            label="workbook-survey",
            run_id=self.run_id,
            engagement_id=self.identifiers.engagement_id,
            step_id=self.step_id,
        )
        request = LLMRequest(
            pin=self.pin,
            system=f"{self.SYSTEM}\n\n{SENTINEL_RULE}",
            messages=(Message("user", wrapped.text),),
            max_tokens=2000,
            metadata={"step_id": self.step_id or "", "purpose": "sheet-selection"},
        )
        call = self.door.send(
            request, self.identifiers, self.run_id, step_id=self.step_id, estimated_cost_usd=0.01
        )
        self.calls += 1
        return _parse_selection(call.response.text)


def _parse_selection(text: str) -> Selection:
    start, end = text.find("{"), text.rfind("}")
    body = json.loads(text[start : end + 1]) if start >= 0 and end > start else {}
    return Selection(
        tuple(
            SheetChoice(str(c["sheet"]), str(c.get("reason", ""))) for c in body.get("selected", [])
        ),
        tuple(
            SheetChoice(str(c["sheet"]), str(c.get("reason", ""))) for c in body.get("skipped", [])
        ),
        by="model",
    )


# ---------------------------------------------------------------- survey


def survey_workbook(data: bytes) -> WorkbookSurvey:
    from openpyxl import load_workbook

    book = load_workbook(io.BytesIO(data), data_only=True, read_only=False)
    sheets: list[SheetInfo] = []
    for ws in book.worksheets:
        headers: list[str] = []
        numeric = 0
        in_title_block = True
        for r, row in enumerate(ws.iter_rows(values_only=True), start=1):
            row_has_number = any(
                isinstance(c, int | float) and not isinstance(c, bool) for c in row
            )
            for cell in row:
                if isinstance(cell, int | float) and not isinstance(cell, bool):
                    numeric += 1
                elif (
                    in_title_block
                    and not row_has_number
                    and r <= HEADER_ROWS
                    and isinstance(cell, str)
                    and cell.strip()
                ):
                    headers.append(cell.strip())
            # The title block ends at the first row with a figure in it. Everything above
            # is where the company, the statement, the period and the scale are stated.
            if row_has_number:
                in_title_block = False
        sheets.append(
            SheetInfo(
                name=ws.title,
                state=ws.sheet_state,
                rows=ws.max_row,
                columns=ws.max_column,
                headers=tuple(headers),
                numeric_cells=numeric,
            )
        )
    for cs in book.chartsheets:
        sheets.append(SheetInfo(cs.title, cs.sheet_state, 0, 0, (), 0, kind="chartsheet"))
    return WorkbookSurvey(tuple(sheets))


# ------------------------------------------------------------- normalise


@dataclass
class NormalizedSheet:
    sheet: str
    label_column: str
    blocks: dict[str, tuple[Row, ...]] = field(default_factory=dict)
    first_rows: list[str] = field(default_factory=list)

    def as_text(self) -> str:
        lines = [f"[sheet {self.sheet}]"]
        for locator, rows in self.blocks.items():
            lines.append(f"{locator}:")
            lines.extend(f"  {r.ordinal}: {r.label} = {r.raw}" for r in rows)
        return "\n".join(lines)


def normalize_workbook(data: bytes, sheets: list[str]) -> list[NormalizedSheet]:
    from openpyxl import load_workbook
    from openpyxl.utils import get_column_letter

    book = load_workbook(io.BytesIO(data), data_only=True)
    out: list[NormalizedSheet] = []
    for name in sheets:
        ws = book[name]
        grid = list(ws.iter_rows(values_only=True))
        if not grid:
            continue
        width = max(len(r) for r in grid)
        label_col = _label_column(grid, width)
        code_col = _code_column(grid, label_col)
        sheet = NormalizedSheet(sheet=name, label_column=get_column_letter(label_col + 1))
        for c in range(width):
            if c in (label_col, code_col):
                continue
            rows: list[Row] = []
            for r, values in enumerate(grid, start=1):
                cell = values[c] if c < len(values) else None
                if isinstance(cell, int | float) and not isinstance(cell, bool):
                    label = values[label_col] if label_col < len(values) else None
                    text = str(label).strip() if label is not None else ""
                    if code_col is not None and code_col < len(values):
                        code = values[code_col]
                        if code is not None and str(code).strip():
                            text = f"{_code_text(code)} {text}".strip()
                    rows.append(
                        Row(
                            label=text,
                            value=float(cell),
                            raw=_display(cell),
                            ordinal=r,
                        )
                    )
            if not rows:
                continue
            col = get_column_letter(c + 1)
            locator = f"{_quote(name)}!{col}{rows[0].ordinal}:{col}{rows[-1].ordinal}"
            sheet.blocks[locator] = tuple(rows)
            if not sheet.first_rows:
                sheet.first_rows = [f"{r.label} {r.raw}".strip() for r in rows[:5]]
        out.append(sheet)
    return out


def _label_column(grid: list[tuple[Any, ...]], width: int) -> int:
    """The leftmost column that is mostly text, voted over the data rows.

    Over the data rows and not the title block: a short trial balance with three title
    rows in column A and three account numbers below them voted A "mostly text", read every
    row as `6000 = 265,000`, and dropped the account-name column - which is the one thing
    every skill downstream needs, and the one place a hostile sentence would sit. The
    title rows say what the sheet is; they do not say which column names the lines.
    """
    data = [row for row in grid if _has_number(row)] or grid
    for c in range(width):
        values = [row[c] for row in data if c < len(row) and row[c] is not None]
        if values and sum(isinstance(v, str) for v in values) >= len(values) / 2:
            return c
    return 0


def _code_column(grid: list[tuple[Any, ...]], label_col: int) -> int | None:
    """An account-number column immediately left of the label column, if there is one.

    `Account # | Account name | Debit | Credit` is the shape a trial balance nearly always
    arrives in. The number is not the label, but an analyst writes the two together -
    `6000 Officer compensation` - and so does this.
    """
    if label_col == 0:
        return None
    c = label_col - 1
    data = [row for row in grid if _has_number(row)]
    values = [row[c] for row in data if c < len(row) and row[c] is not None]
    if not values:
        return None
    codes = sum(
        1 for v in values
        if (isinstance(v, int | float) and not isinstance(v, bool) and float(v).is_integer())
        or (isinstance(v, str) and v.strip().isdigit())
    )
    return c if codes >= len(values) * 0.8 else None


def _has_number(row: tuple[Any, ...]) -> bool:
    return any(isinstance(c, int | float) and not isinstance(c, bool) for c in row)


def _code_text(code: Any) -> str:
    """An account number as written - `6000`, never `6,000`. It is a code, not a figure."""
    if isinstance(code, int | float) and not isinstance(code, bool) and float(code).is_integer():
        return str(int(code))
    return str(code).strip()


def _display(value: float) -> str:
    return f"{int(value):,}" if float(value).is_integer() else f"{value:,.2f}"


def _quote(sheet: str) -> str:
    return f"'{sheet}'" if re.search(r"[^A-Za-z0-9_]", sheet) else sheet
