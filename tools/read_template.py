"""Read a workbook and draft the input schema this build would need to write into it.

The thing standing between this system and the firm's own model is one file and about four
hundred lines of YAML nobody has written. `BUILD/OPEN.md` has carried it since P24 - *"the
firm's real model: `Model Template - MASTER 2026`. The input-help workbook has to match
its input format, and the roll-forward's Excel half has to bind its named ranges"* - and
`templates/inputs/Chartwell_Input_Help_v1@1.0.yaml` says in its own comment that when the
real template arrives *"this schema is re-derived from its Input tab"*.

This is the tool that re-derives it. Point it at the workbook and it writes two things:

* **the inventory** - every defined name in the file, the sheet and cell it points at,
  whether that cell is empty, holds a number or holds a formula, and whether the name
  spans one cell or a range. That answers the question CLAUDE.md says is *"not knowable
  from a photograph"*, and the roll-forward's Excel half needs it as much as the schema
  does.
* **a draft schema**, in the shape `deliver/registry.TemplateInputs.load` reads, with
  every name it believes is an input and the `line` and `period` it proposes for each.

**The second half is a proposal and says so, loudly.** Deciding that a range called
`Revenue_FY1` means the line *Revenue* in period *FY1* is reading a naming convention,
and a convention is a thing a firm has rather than a thing a file states. So every
proposed `line` is marked, the draft will not load until a person has been through it, and
the tool prints how many it guessed at. A schema this had silently filled in would bind
the firm's model by inference - and the first symptom would be a figure written into the
wrong row of a workbook that opens perfectly.

**It never writes into the workbook.** It reads bytes. The one file it writes is the YAML
it was asked for, and it refuses to overwrite one that exists.

    python -m uv run python tools/read_template.py "Model Template - MASTER 2026.xlsx"
    python -m uv run python tools/read_template.py <book.xlsx> --sheet Input --out draft.yaml
    python -m uv run python tools/read_template.py <book.xlsx> --inventory

What to do with the draft: read every `line:` in it, correct the ones that are wrong,
delete the `REVIEW` markers as you go, move anything the model computes for itself out of
`required`, then save it as `templates/inputs/<file>@<version>.yaml` and bind it at a
scope on the Setup page. The binding is a decision and it is recorded there.
"""

from __future__ import annotations

import argparse
import json
import re
import sys
from collections.abc import Sequence
from dataclasses import dataclass, field
from io import BytesIO
from pathlib import Path
from typing import Any

ROOT = Path(__file__).resolve().parents[1]
if str(ROOT) not in sys.path:  # this is run as a script, from wherever somebody is standing
    sys.path.insert(0, str(ROOT))

from tools._environment import needs  # noqa: E402

needs("openpyxl", "yaml", script="tools/read_template.py")

import yaml  # noqa: E402
from openpyxl import load_workbook  # noqa: E402
from openpyxl.utils.cell import range_boundaries  # noqa: E402

from deliver.registry import version_key  # noqa: E402

__all__ = [
    "NEEDS_A_PERSON",
    "Drafted",
    "Name",
    "draft",
    "inventory",
    "names_in",
    "report",
]

# The marker every proposed line carries. Written into the draft so it cannot be loaded
# and registered without somebody having read it: `TemplateInputs.load` has no idea what
# this means, which is the point - a schema full of these is a schema that was generated
# and not reviewed, and the person who generated it is the only one who knows that.
NEEDS_A_PERSON = "REVIEW"

# How a name that carries a period usually ends. Not an exhaustive grammar - a convention
# this tool has seen, used only to *propose*, and the proposal is marked.
PERIOD = re.compile(r"[_\s]?(FY\d{1,2}|Y\d{1,2}|LTM|TTM|BUD|ACT|\d{4})$", re.I)

# Names that are not an input at all, whatever cell they point at: the stamps this build
# writes itself, and the print and scroll areas Excel keeps as defined names.
NOT_AN_INPUT = (
    "Print_Area",
    "Print_Titles",
    "_xlnm",
    "Template_Version",
    "Run_Id",
)


@dataclass(frozen=True)
class Name:
    """One defined name in the workbook, and what is behind it."""

    name: str
    sheet: str
    ref: str
    cells: int
    holds: str
    """`empty`, `number`, `text`, `formula`, or `range` when the name spans more than one
    cell and a single answer would be a guess."""
    value: Any = None

    @property
    def is_a_single_cell(self) -> bool:
        return self.cells == 1

    @property
    def looks_like_an_input(self) -> bool:
        """An empty single cell, on a sheet, under a name this build did not reserve.

        Empty is the signal and it is a good one: a template is empty by definition, so a
        cell somebody means to be filled in is a cell with nothing in it. A cell holding a
        formula is the model's own work and writing into it would replace the model.
        """
        if any(mark in self.name for mark in NOT_AN_INPUT):
            return False
        return self.is_a_single_cell and self.holds == "empty"

    def as_row(self) -> dict[str, Any]:
        return {
            "name": self.name,
            "sheet": self.sheet,
            "ref": self.ref,
            "cells": self.cells,
            "holds": self.holds,
            "input": self.looks_like_an_input,
        }


@dataclass(frozen=True)
class Drafted:
    """The draft schema, and an honest account of how much of it was guessed."""

    schema: dict[str, Any]
    names: tuple[Name, ...]
    proposed: tuple[str, ...]
    """Every input whose `line` this tool proposed rather than read. Which is all of
    them: a workbook states a name and a cell, and never what the name means."""
    skipped: tuple[str, ...] = field(default_factory=tuple)

    @property
    def inputs(self) -> list[dict[str, Any]]:
        required: list[dict[str, Any]] = self.schema.get("required", [])
        return required


# --------------------------------------------------------------------- reading the file


def names_in(data: bytes, *, sheet: str | None = None) -> tuple[tuple[Name, ...], tuple[str, ...]]:
    """Every defined name in the workbook, resolved to a sheet, a cell and what it holds.

    Both the workbook's global names and each sheet's local ones, because a model's Input
    tab often defines its names locally and a reader that asked only the workbook would
    report a file with four hundred names as having none.
    """
    wb = load_workbook(BytesIO(data), data_only=False)
    skipped: list[str] = []
    found: list[Name] = []

    defined: list[tuple[str, Any]] = [(k, v) for k, v in wb.defined_names.items()]
    for ws in wb.worksheets:
        defined.extend((k, v) for k, v in ws.defined_names.items())

    sheets = {w.title for w in wb.worksheets}
    for name, entry in defined:
        destinations = list(getattr(entry, "destinations", []) or [])
        if not destinations:
            skipped.append(f"{name}: names no destination ({getattr(entry, 'attr_text', '')!r})")
            continue
        if len(destinations) > 1:
            skipped.append(f"{name}: names {len(destinations)} destinations, so it is not one cell")
            continue
        where, ref = destinations[0]
        if where not in sheets:
            skipped.append(f"{name}: points at sheet {where!r}, which this workbook has not got")
            continue
        if sheet is not None and where != sheet:
            continue
        try:
            min_col, min_row, max_col, max_row = range_boundaries(str(ref).replace("$", ""))
        except Exception:  # noqa: BLE001 - an unparseable ref is reported, never guessed at
            skipped.append(f"{name}: {ref!r} is not a reference this can follow")
            continue
        if None in (min_col, min_row, max_col, max_row):
            skipped.append(f"{name}: {ref!r} is a whole row or column, not a cell")
            continue
        cells = (int(max_col) - int(min_col) + 1) * (int(max_row) - int(min_row) + 1)
        value = wb[where].cell(row=int(min_row), column=int(min_col)).value if cells == 1 else None
        found.append(
            Name(
                name=name,
                sheet=where,
                ref=str(ref).replace("$", ""),
                cells=cells,
                holds=_holds(value) if cells == 1 else "range",
                value=value if cells == 1 else None,
            )
        )

    if sheet is not None and not found:
        skipped.append(
            f"no defined name points at sheet {sheet!r}. The sheets in this workbook are: "
            + ", ".join(sorted(sheets))
        )
    return tuple(sorted(found, key=lambda n: (n.sheet, n.name))), tuple(skipped)


def _holds(value: Any) -> str:
    if value is None or (isinstance(value, str) and not value.strip()):
        return "empty"
    if isinstance(value, str):
        return "formula" if value.startswith("=") else "text"
    if isinstance(value, bool):
        return "text"
    return "number"


def inventory(data: bytes, *, sheet: str | None = None) -> dict[str, Any]:
    """The inventory, as data. What is in this workbook and what each name points at."""
    names, skipped = names_in(data, sheet=sheet)
    return {
        "names": [n.as_row() for n in names],
        "sheets": sorted({n.sheet for n in names}),
        "inputs": sum(1 for n in names if n.looks_like_an_input),
        "skipped": list(skipped),
    }


# -------------------------------------------------------------------- drafting a schema


def draft(
    data: bytes,
    *,
    file: str,
    template_version: str,
    sheet: str | None = None,
    scale: int = 1,
    currency: str = "USD",
) -> Drafted:
    """A schema in the registry's shape, with every `line` marked as a proposal.

    `scale` defaults to 1 and not to 1000. The firm's convention is that statements go
    into the model in thousands, and this reads a workbook - which states a layout and
    never its own units. Defaulting to the convention would put the one assumption nobody
    can see underneath every figure in the file; `--scale` says it out loud instead.
    """
    names, skipped = names_in(data, sheet=sheet)
    inputs = [n for n in names if n.looks_like_an_input]

    required: list[dict[str, Any]] = []
    proposed: list[str] = []
    for one in inputs:
        line, period = _propose(one.name)
        required.append(
            {
                "name": one.name,
                "line": f"{NEEDS_A_PERSON}: {line}",
                "period": period or f"{NEEDS_A_PERSON}: no period in the name",
            }
        )
        proposed.append(one.name)

    schema = {
        "template": file,
        "template_version": template_version,
        "scale": scale,
        "currency": currency,
        "stamps": ["Template_Version", "Run_Id"],
        "sources_beside_values": False,
        "required": required,
    }
    return Drafted(
        schema=schema, names=names, proposed=tuple(proposed), skipped=tuple(skipped)
    )


def _propose(name: str) -> tuple[str, str]:
    """A line and a period, read off a name's spelling.

    `Officer_Comp_FY2` becomes ("Officer Comp", "FY2"). That is a reading of a convention
    and not a fact about the workbook, which is why everything this returns is marked
    before it reaches the file.
    """
    match = PERIOD.search(name)
    period = match.group(1).upper() if match else ""
    body = name[: match.start()] if match else name
    line = re.sub(r"[_\s]+", " ", body).strip()
    return line or name, period


HEADER = """# DRAFT, generated by tools/read_template.py from {file}.
#
# **Every `line:` below is a proposal, not a reading.** This was derived from {count}
# defined name(s) that point at an empty single cell on {where} - which is what an input
# looks like in a template, because a template is empty by definition. What a name
# *means* is not in the file: `{example}` was read as a line and a period by splitting on
# the trailing period marker, which is a convention a firm has rather than something the
# workbook states.
#
# So this will not load. `{marker}:` markers are there so it cannot be registered by
# accident, and removing one is somebody saying they read that line.
#
# What to do, in order:
#
#   1. Read every `line:`. Correct the ones that are wrong and drop the `{marker}:`
#      prefix as you go. A line has to match what `line-item-tagging` calls it, or
#      model-load cannot bind a figure to it.
#   2. Move anything the model computes for itself out of `required`. A required input
#      with no source stops a build for a person, which is right for a figure somebody
#      has to supply and wrong for one the model derives.
#   3. Move anything optional into an `optional:` block of the same shape.
#   4. Check `scale:`. It is {scale} here because a workbook states a layout and never
#      its own units, and the firm's convention is thousands - so if these cells expect
#      thousands, say so.
#   5. Save it as
#        templates/inputs/{key}.yaml
#      and bind it at a scope on the Setup page, where the binding is recorded against
#      you. That filename is the key a run pins, so it has to be exactly that.
#
# Never edit a published version in place. A changed schema is a new file beside it.
"""


def as_yaml(drafted: Drafted) -> str:
    schema = drafted.schema
    example = drafted.proposed[0] if drafted.proposed else "Revenue_FY1"
    header = HEADER.format(
        file=schema["template"],
        key=version_key(str(schema["template"]), str(schema["template_version"])),
        count=len(drafted.proposed),
        where=", ".join(sorted({n.sheet for n in drafted.names})) or "no sheet",
        example=example,
        marker=NEEDS_A_PERSON,
        scale=schema["scale"],
    )
    body = yaml.safe_dump(schema, sort_keys=False, default_flow_style=False, width=96)
    return header + "\n" + str(body)


# ------------------------------------------------------------------------------ report


def report(drafted: Drafted, *, as_json: bool = False, full: bool = False) -> str:
    if as_json:
        return json.dumps(
            {
                "names": [n.as_row() for n in drafted.names],
                "proposed": list(drafted.proposed),
                "skipped": list(drafted.skipped),
                "schema": drafted.schema,
            },
            indent=2,
            default=str,
        )
    names = drafted.names
    by_sheet: dict[str, list[Name]] = {}
    for one in names:
        by_sheet.setdefault(one.sheet, []).append(one)

    lines = [
        f"  {len(names)} defined name(s) over {len(by_sheet)} sheet(s), "
        f"{len(drafted.proposed)} of which look like inputs"
    ]
    for sheet, group in sorted(by_sheet.items()):
        holds: dict[str, int] = {}
        for one in group:
            holds[one.holds] = holds.get(one.holds, 0) + 1
        lines.append(
            f"\n  {sheet}  ({len(group)} name(s): "
            + ", ".join(f"{n} {k}" for k, n in sorted(holds.items()))
            + ")"
        )
        shown = group if full else group[:12]
        for one in shown:
            mark = "input " if one.looks_like_an_input else "      "
            span = "" if one.is_a_single_cell else f"  [{one.cells} cells]"
            lines.append(f"    {mark}{one.name:34} {one.ref:10} {one.holds}{span}")
        if len(group) > len(shown):
            lines.append(f"    ... and {len(group) - len(shown)} more (--full to see them)")

    if drafted.skipped:
        lines.append("\n  Names this could not follow")
        for note in drafted.skipped:
            lines.append(f"    - {note}")

    lines.append(
        f"\n  Every `line` in the draft is a proposal. {len(drafted.proposed)} of them, "
        f"each marked {NEEDS_A_PERSON}, because a workbook states a name and a cell and "
        f"never what the name means."
    )
    return "\n".join(lines)


# ------------------------------------------------------------------------------- CLI


def main(argv: Sequence[str] | None = None) -> int:
    parser = argparse.ArgumentParser(
        prog="read_template.py",
        description="Read a workbook's named ranges and draft the input schema for it.",
    )
    parser.add_argument("workbook", type=Path)
    parser.add_argument(
        "--sheet",
        help="only names pointing at this sheet - a model's Input tab, usually. "
        "Without it, every sheet.",
    )
    parser.add_argument("--out", type=Path, help="write the draft schema here")
    parser.add_argument("--version", default="0.1", help="the template version to draft under")
    parser.add_argument(
        "--scale",
        type=int,
        default=1,
        help="what these cells expect: 1 for units, 1000 for thousands. Defaults to 1 "
        "because a workbook states a layout and never its own units.",
    )
    parser.add_argument("--currency", default="USD")
    parser.add_argument("--inventory", action="store_true", help="the names as JSON, nothing else")
    parser.add_argument("--full", action="store_true", help="every name, not the first twelve")
    parser.add_argument("--json", action="store_true", dest="as_json")
    args = parser.parse_args(argv)

    if not args.workbook.is_file():
        print(f"  no such workbook: {args.workbook}", file=sys.stderr)
        return 2
    try:
        data = args.workbook.read_bytes()
        if args.inventory:
            print(json.dumps(inventory(data, sheet=args.sheet), indent=2, default=str))
            return 0
        drafted = draft(
            data,
            file=args.workbook.name,
            template_version=args.version,
            sheet=args.sheet,
            scale=args.scale,
            currency=args.currency,
        )
    except Exception as exc:  # noqa: BLE001 - a workbook that will not open says so
        print(f"  {args.workbook.name} could not be read: {exc}", file=sys.stderr)
        return 2

    print(report(drafted, as_json=args.as_json, full=args.full))

    if args.out:
        if args.out.exists():
            print(
                f"\n  {args.out} is already there and this will not write over it. A "
                f"schema somebody has been through is worth more than a fresh draft.",
                file=sys.stderr,
            )
            return 2
        args.out.write_text(as_yaml(drafted), encoding="utf-8")
        key = version_key(args.workbook.name, args.version)
        print(f"\n  draft written to {args.out}")
        print(
            f"  It will not load until the {len(drafted.proposed)} {NEEDS_A_PERSON} "
            f"marker(s) are gone. The header says what to do, in order."
        )
        print(
            f"  Its registered name is templates/inputs/{key}.yaml - that is the key a "
            f"run pins, so the file has to be called exactly that."
        )
    elif drafted.proposed:
        print("\n  --out <path> to write the draft.")
    return 0 if drafted.proposed else 1


if __name__ == "__main__":  # pragma: no cover - the CLI's own line
    raise SystemExit(main())
