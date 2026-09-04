"""The Model Template Writer: writes values into a copy of the pinned template.

openpyxl against the registered template's bytes, opened as a copy - the master is never
written to. Every destination is a named range, never a cell address; a destination not
in the mapping schema is refused, so a hallucinated target fails loudly at write time
rather than landing in a cell nobody meant. Formulas, other names and formatting are
left as the template had them.

Runs inside assembly, after unmasking and after every guard: the values it is handed
are already re-identified, tied out and locked, and the Numeric Integrity Lock brackets
the write by reading the result back. Pure code; the mapping it executes came from the
model-load skill and the Model Builder resolved every figure before calling this.

No Excel, no COM, no VBA. Excel recalculates when a person opens the file.
"""

from __future__ import annotations

from collections.abc import Collection, Mapping
from io import BytesIO
from typing import Any

from openpyxl import load_workbook
from openpyxl.utils.cell import range_boundaries
from openpyxl.workbook.workbook import Workbook

from registries.errors import CELL_ADDRESS_BINDING, WRITE_FAILED, RegistryError, diagnose

__all__ = ["cell_for", "read_named", "write_template"]

Value = float | int | str


def cell_for(wb: Workbook, name: str) -> tuple[str, str]:
    """The (sheet, coordinate) a defined name points at. One cell, or a refusal: a
    figure has one destination, and a name spanning a range is not one."""
    defined = wb.defined_names.get(name)
    if defined is None:
        for ws in wb.worksheets:
            if name in ws.defined_names:
                defined = ws.defined_names[name]
                break
    if defined is None or not defined.attr_text:
        raise RegistryError(
            diagnose(
                WRITE_FAILED,
                f"{name!r} is not a defined name in this workbook",
                locator=name,
            )
        )
    destinations = list(defined.destinations)
    if len(destinations) != 1:
        raise RegistryError(
            diagnose(
                WRITE_FAILED,
                f"{name!r} names {len(destinations)} destinations; a figure has one",
                locator=name,
            )
        )
    sheet, ref = destinations[0]
    ref = ref.replace("$", "")
    min_col, min_row, max_col, max_row = range_boundaries(ref)
    if (min_col, min_row) != (max_col, max_row):
        raise RegistryError(
            diagnose(
                WRITE_FAILED,
                f"{name!r} spans {ref}; a figure has one cell",
                locator=name,
            )
        )
    return str(sheet), ref


def write_template(
    template: bytes,
    writes: Mapping[str, Value],
    declared: Collection[str],
    stamps: Mapping[str, str],
) -> bytes:
    """Execute the writes against a copy of the template and return the new bytes.

    `writes` is named range -> value. `declared` is the mapping schema's input names;
    a write to anything else is refused before the workbook is opened. `stamps` are the
    template-version and run-id cells, which the writer owns and the mapping never names.
    """
    for name in writes:
        if name not in declared:
            raise RegistryError(
                diagnose(
                    WRITE_FAILED,
                    f"{name!r} is not in the mapping schema; the writer refuses a "
                    f"destination the mapping did not declare",
                    locator=name,
                )
            )
        if name in stamps:
            raise RegistryError(
                diagnose(
                    WRITE_FAILED,
                    f"{name!r} is a stamp; the writer sets it, the mapping does not",
                    locator=name,
                )
            )
    for name in (*writes, *stamps):
        if _looks_like_a_cell(name):
            raise RegistryError(
                diagnose(
                    CELL_ADDRESS_BINDING,
                    f"{name!r} is a cell address, not a named range",
                    locator=name,
                    found=name,
                    expected="a named range",
                )
            )

    wb = load_workbook(BytesIO(template))  # a copy: the master's bytes are never touched
    for name, value in writes.items():
        sheet, ref = cell_for(wb, name)
        wb[sheet][ref] = value
    for name, text in stamps.items():
        sheet, ref = cell_for(wb, name)
        wb[sheet][ref] = str(text)
    out = BytesIO()
    wb.save(out)
    return out.getvalue()


def read_named(data: bytes, names: Collection[str]) -> dict[str, Any]:
    """What the named cells hold, from the bytes. The lock's after-image."""
    wb = load_workbook(BytesIO(data))
    found: dict[str, Any] = {}
    for name in names:
        sheet, ref = cell_for(wb, name)
        found[name] = wb[sheet][ref].value
    return found


def _looks_like_a_cell(name: str) -> bool:
    import re

    return bool(re.match(r"^\$?[A-Za-z]{1,3}\$?[0-9]+$", name))
