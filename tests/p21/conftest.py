"""Fixtures for P21: the registry, the writer, the builder, QC and release, over the
Henderson case from P20 carried one step further - into a workbook.

`world` wires the P20 world plus the template registry, the schema lookup tool, the
release gate over a local outbox, and the mapping the model-load skill is expected to
propose for the ESOP template. The template on disk is synthetic, built by
`tools/synthesize_templates.py` from `templates/inputs/`.
"""

from __future__ import annotations

import json
from collections.abc import Iterator
from io import BytesIO
from pathlib import Path
from typing import Any

import pytest
from cryptography.fernet import Fernet
from openpyxl import Workbook, load_workbook
from openpyxl.workbook.defined_name import DefinedName

import ports
from deliver import DeliverableQC, ModelBuilder, ReleaseGate, TemplateRegistry
from deliver.registry import TemplateInputs, version_key
from graph.compile import CompiledGraph, PlanContext, Step, compile_recipe
from graph.recipes import load_recipe
from graph.resolve import ResolvedScope, ScopeResolver
from guards.untrusted import UntrustedBoundary
from harness.budget import BudgetGovernor
from harness.context import ContextBuilder
from harness.dispatch import SkillDispatcher
from harness.door import ModelDoor
from harness.policy import PolicyGate, RunScope
from harness.procedures.model_load import (
    PRODUCED_BY,
    PROPOSAL_CONTRACT,
    ModelLoad,
    ModelLoadInputs,
    catalogue_from,
    inputs_from,
)
from harness.runtime import BoundTools, GuardedRuntime
from ops.blobs import LocalContentStore
from ops.cost import CostMeter
from ops.events import EventBus
from ops.models import RunState
from ops.run_store import SqlRunStore, open_store
from ops.source_parse import SourceParseStore
from ports.delivery.local_folder import LocalFolderDelivery
from ports.llm import LLMRequest, LLMResponse, ModelPin, Usage
from precedent import PrecedentStore
from privacy.mask import Masker
from privacy.payload_scan import PayloadScan
from privacy.vault import KEY_NAME, open_vault
from registries.constraints import ConstraintRegistry
from registries.contracts import ContractRegistry
from registries.skills import SkillTree
from tests.p15.conftest import profile
from tests.p20.conftest import blocks_of, load_case, timing_output, write_blocks
from tools.template_schema import template_schema_lookup

ROOT = Path(__file__).resolve().parents[2]
ENGAGEMENT = "HEND-2026-04"
CLIENT = "client_henderson"
SCOPE_PATH = "valuation/esop/repeat"
RUN_ID = "run_p21"
PIN = ModelPin.for_tier("mid")
TAXONOMY = "coa-taxonomy@1.0.0"
SKILL = "model-load@0.1.0"
TEMPLATE_FILE = "Chartwell_ESOP_Model_v9.xlsx"
TEMPLATE_VERSION = "9.4"
TEMPLATE_KEY = version_key(TEMPLATE_FILE, TEMPLATE_VERSION)
SLOTS = {"FY2023": "FY1", "FY2024": "FY2", "FY2025": "FY3"}

# The historical table's lines against the template's, for the expected mapping. The
# officer lines are two on the table and one on the template: salary is bound and the
# bonus is unrouted with the reason, per template-input-map.md.
LINE_TO_INPUT = {
    "Revenue": "Revenue",
    "Cost of sales": "COGS",
    "Officer salary": "Officer_Comp",
    "Wages": "Wages",
    "Payroll taxes": "Payroll_Taxes",
    "Rent - facilities": "Rent",
    "Repairs and maintenance": "Repairs",
    "Utilities": "Utilities",
    "Insurance": "Insurance",
    "Professional fees": "Professional_Fees",
    "Depreciation": "Depreciation",
    "Travel": "Travel",
    "Office supplies": "Office",
    "Charitable contributions": "Other_Opex",
}


def expected_proposal(inputs: ModelLoadInputs) -> dict[str, Any]:
    """What the model-load skill is expected to propose for Henderson on the ESOP
    template: every required input bound to the matching line and slot."""
    bindings = []
    routed = set()
    for ref in inputs.catalogue:
        if ref.source_output == "profile":
            bindings.append({"named_range": "Owner_Comp_Benchmark", "figure_ref": ref.figure_ref})
            routed.add(ref.figure_ref)
            continue
        target = LINE_TO_INPUT.get(ref.description)
        if target is None:
            continue
        bindings.append({"named_range": f"{target}_{ref.period}", "figure_ref": ref.figure_ref})
        routed.add(ref.figure_ref)
    unrouted = [
        {"figure_ref": ref.figure_ref, "why": "the template has no line for it"}
        for ref in inputs.catalogue
        if ref.figure_ref not in routed
    ]
    return {
        "schema_version": PROPOSAL_CONTRACT,
        "produced_by": PRODUCED_BY,
        "bindings": bindings,
        "unmapped_template_inputs": [],
        "unrouted_upstream_figures": unrouted,
    }


class ProposalPort:
    name = "proposal"

    def __init__(self, proposal: dict[str, Any]) -> None:
        self.proposal = proposal
        self.requests: list[LLMRequest] = []

    def complete(self, request: LLMRequest) -> LLMResponse:
        self.requests.append(request)
        return LLMResponse(
            text=json.dumps(self.proposal),
            model=request.pin.model,
            stop_reason="end_turn",
            usage=Usage(input_tokens=5000, output_tokens=1500),
        )


def workbook_with(names: dict[str, str], cells: dict[str, Any] | None = None) -> bytes:
    """A small workbook declaring `names` (name -> Inputs!$B$n) and holding `cells`."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Inputs"
    for name, ref in names.items():
        wb.defined_names[name] = DefinedName(name, attr_text=ref)
    for coordinate, value in (cells or {}).items():
        sheet, cell = coordinate.split("!")
        target = wb[sheet] if sheet in wb.sheetnames else wb.create_sheet(sheet)
        target[cell] = value
    out = BytesIO()
    wb.save(out)
    return out.getvalue()


def named_values(data: bytes, names: list[str]) -> dict[str, Any]:
    from tools.write_template import read_named

    return read_named(data, names)


def sheet_values(data: bytes, sheet: str) -> dict[str, Any]:
    wb = load_workbook(BytesIO(data))
    return {
        c.coordinate: c.value for row in wb[sheet].iter_rows() for c in row if c.value is not None
    }


# ------------------------------------------------------------------ registries


@pytest.fixture(scope="session")
def contracts() -> ContractRegistry:
    return ContractRegistry.load()


@pytest.fixture(scope="session")
def constraints() -> ConstraintRegistry:
    return ConstraintRegistry.load()


@pytest.fixture(scope="session")
def tree() -> SkillTree:
    return SkillTree.load()


@pytest.fixture(scope="session")
def scope() -> ResolvedScope:
    return ScopeResolver().resolve(SCOPE_PATH)


@pytest.fixture(scope="session")
def plan(scope: ResolvedScope) -> CompiledGraph:
    return compile_recipe(
        load_recipe("monthly-refresh"),
        PlanContext(
            engagement_id=ENGAGEMENT,
            scope=scope,
            taxonomy_version=TAXONOMY,
            precedent_snapshot="precedent_snapshot_empty",
        ),
    )


@pytest.fixture(scope="session")
def model_step(plan: CompiledGraph) -> Step:
    return plan.step("model")


@pytest.fixture(scope="session")
def case() -> dict[str, Any]:
    return load_case()


@pytest.fixture(scope="session")
def schema() -> TemplateInputs:
    return TemplateInputs.load(TEMPLATE_KEY)


# ------------------------------------------------------------------ the record


@pytest.fixture
def store(tmp_path: Path) -> Iterator[SqlRunStore]:
    opened = open_store(f"sqlite:///{(tmp_path / 'p21.db').as_posix()}")
    yield opened
    opened.close()


@pytest.fixture
def blobs(tmp_path: Path) -> LocalContentStore:
    return LocalContentStore(tmp_path / "blobs")


@pytest.fixture
def bus(store: SqlRunStore, contracts: ContractRegistry) -> EventBus:
    return EventBus(store, contracts=contracts)


@pytest.fixture
def parse_store(store: SqlRunStore) -> SourceParseStore:
    return SourceParseStore(store)


@pytest.fixture
def started(store: SqlRunStore, plan: CompiledGraph) -> str:
    store.start_run(
        run_id=RUN_ID,
        engagement_id=ENGAGEMENT,
        plan_hash=plan.plan_hash,
        versions={
            "model_pin": str(PIN),
            "skills": [SKILL],
            "contracts": "1.0.0",
            "constraints": "1.0.0",
            "taxonomy": TAXONOMY,
            "template_version": TEMPLATE_KEY,
        },
        tieout_tolerance=1.0,
    )
    store.transition(RUN_ID, RunState.RUNNING)
    store.start_clock(RUN_ID)
    return RUN_ID


@pytest.fixture
def registry(scope: ResolvedScope) -> TemplateRegistry:
    reg = TemplateRegistry()
    reg.register(TEMPLATE_FILE, TEMPLATE_VERSION)
    return reg


@pytest.fixture
def outbox(tmp_path: Path) -> Path:
    return tmp_path / "outbox"


# ------------------------------------------------------------------ world


@pytest.fixture
def world(
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
    store: SqlRunStore,
    bus: EventBus,
    blobs: LocalContentStore,
    contracts: ContractRegistry,
    constraints: ConstraintRegistry,
    tree: SkillTree,
    scope: ResolvedScope,
    parse_store: SourceParseStore,
    registry: TemplateRegistry,
    outbox: Path,
    case: dict[str, Any],
    started: str,
) -> Iterator[dict[str, Any]]:
    monkeypatch.setenv(KEY_NAME, Fernet.generate_key().decode())
    ports.reset()
    vault = open_vault(ports.get_secrets(), f"sqlite:///{(tmp_path / 'vault.db').as_posix()}")
    meter = CostMeter(store)
    governor = BudgetGovernor(constraints, meter, bus, store)
    runtime = GuardedRuntime(
        PolicyGate(constraints, bus),
        constraints,
        bus,
        {"template_schema_lookup": template_schema_lookup},
    )
    run_scope = RunScope(
        run_id=started,
        engagement_id=ENGAGEMENT,
        engagement_root=f"engagements/{ENGAGEMENT}",
        scope_path=SCOPE_PATH,
        versions={"chart_of_accounts_lookup": TAXONOMY, "template_schema_lookup": TEMPLATE_KEY},
        budget_remaining_usd=governor.remaining_usd(started),
    )
    masker = Masker(vault)
    engagement_profile = blobs.put_json(profile(), name="engagement_profile")
    write_blocks(parse_store, started, case)

    timing = timing_output(case)
    template = registry.get(TEMPLATE_KEY)
    schema_rows = [i.as_row() for i in template.inputs.inputs]
    catalogue = catalogue_from(timing, profile(), period_slots=SLOTS)
    inputs = inputs_from(
        TEMPLATE_KEY, TEMPLATE_VERSION, schema_rows, catalogue, periods=("FY1", "FY2", "FY3")
    )

    def dispatcher_for(port: Any) -> SkillDispatcher:
        return SkillDispatcher(
            tree=tree,
            context=ContextBuilder(tree, blobs, bus, store=store, precedents=PrecedentStore(store)),
            door=ModelDoor(port, PayloadScan(bus), governor, meter, bus),
            runtime=runtime,
            contracts=contracts,
            boundary=UntrustedBoundary(bus),
            masker=masker,
            bus=bus,
            blobs=blobs,
        )

    def tools_for(step_id: str = "model") -> BoundTools:
        return runtime.bound(run_scope, SKILL, step_id)

    delivery = LocalFolderDelivery(outbox)
    yield {
        "store": store,
        "bus": bus,
        "blobs": blobs,
        "vault": vault,
        "masker": masker,
        "runtime": runtime,
        "run_scope": run_scope,
        "engagement_profile": engagement_profile,
        "dispatcher_for": dispatcher_for,
        "tools_for": tools_for,
        "procedure": ModelLoad(),
        "parse_store": parse_store,
        "registry": registry,
        "template": template,
        "timing": timing,
        "blocks": blocks_of(case),
        "inputs": inputs,
        "proposal": expected_proposal(inputs),
        "builder": ModelBuilder(bus),
        "qc": DeliverableQC(bus),
        "delivery": delivery,
        "outbox": outbox,
        "release": ReleaseGate(store, bus, delivery, DeliverableQC(bus)),
        "run_id": started,
    }
    vault.close()


def dispatch(world: dict[str, Any], step: Step, scope: ResolvedScope, port: Any) -> Any:
    return world["dispatcher_for"](port).run(
        step=step,
        run_id=world["run_id"],
        engagement_id=ENGAGEMENT,
        client_id=CLIENT,
        scope=scope,
        engagement_profile=world["engagement_profile"],
        inputs=world["inputs"],
        identifiers=(),
        run_scope=world["run_scope"],
        procedure=world["procedure"],
        pin=PIN,
    )


def mapping_for(world: dict[str, Any]) -> dict[str, Any]:
    """The mapping the procedure emits from the expected proposal, without a model call."""
    out: dict[str, Any] = world["procedure"].apply(
        world["proposal"], world["tools_for"](), world["inputs"], "pin"
    )
    return out


def build(world: dict[str, Any], mapping: dict[str, Any] | None = None) -> Any:
    return world["builder"].build(
        mapping or mapping_for(world),
        world["inputs"].figures(),
        world["template"],
        run_id=world["run_id"],
        engagement_id=ENGAGEMENT,
    )


def events_of(store: SqlRunStore, run_id: str, event_type: str) -> list[Any]:
    return [e for e in store.events(run_id) if e.type == event_type]
