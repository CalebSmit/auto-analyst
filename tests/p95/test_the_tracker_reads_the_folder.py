"""P95 - The tracker reads the folder and updates itself.

Most of this file is about what an update must *not* do. Matching a trial balance to the
item that asks for one is the easy half; the half that decides whether the thing gets used
is that a refresh leaves everything a person wrote exactly where they put it.
"""

from __future__ import annotations

from pathlib import Path
from typing import Any

import pytest
from openpyxl import Workbook, load_workbook

from intake.request_list import RequestItem, RequestList, Status, for_scope
from intake.tracker import (
    DERIVED,
    THEIRS,
    Seen,
    Tracker,
    look_at,
    read_tracker,
    write_tracker,
)
from ports.storage.local_folder import LocalFolderStorage
from tests.p17.conftest import build_pdf

ROOT = Path(__file__).resolve().parents[2]


def a_workbook(path: Path, *rows: tuple[Any, ...], title: str = "Sheet1") -> None:
    book = Workbook()
    sheet = book.active
    sheet.title = title
    for row in rows:
        sheet.append(list(row))
    book.save(path)


@pytest.fixture
def folder(tmp_path: Path) -> Path:
    """A pile of documents shaped like what a client sends: ordinary names, no lettered
    folders, a mix of formats, and one thing nobody asked for."""
    inbox = tmp_path / "inbox"
    inbox.mkdir()
    a_workbook(
        inbox / "TB 2025 final.xlsx",
        ("Account", "Debit", "Credit"),
        ("1000 Cash", 1000, 0),
        title="TB",
    )
    (inbox / "FS 2021-2025.pdf").write_bytes(
        build_pdf(
            [
                ("Consolidated Balance Sheet", 12),
                ("Total assets", 10),
                ("Statement of Operations", 12),
                ("Net income", 10),
            ]
        )
    )
    (inbox / "2024 return.pdf").write_bytes(
        build_pdf([("Form 1120S U.S. Income Tax Return", 12), ("Ordinary business income", 10)])
    )
    (inbox / "holiday photos.zip").write_bytes(b"PK not a document at all")
    return inbox


def tracker_for(folder: Path) -> Tracker:
    listing = for_scope("valuation/non-esop")
    assert listing is not None
    return Tracker(listing, look_at(LocalFolderStorage(folder)))


def sheet_rows(path: Path, name: str = "Request list") -> dict[str, dict[str, Any]]:
    book = load_workbook(path)
    try:
        sheet = book[name]
        header = [c.value for c in sheet[1]]
        out = {}
        for row in sheet.iter_rows(min_row=2, values_only=True):
            out[row[0]] = dict(zip(header, row, strict=False))
        return out
    finally:
        book.close()


# ============================================================ reading the folder


def test_a_folder_of_ordinary_filenames_marks_the_right_items_provided(
    folder: Path,
) -> None:
    """The point of the whole phase. A client sends `TB 2025 final.xlsx`, not
    `A. Financial Records/A.3/`, and the tracker still knows what it is."""
    report = tracker_for(folder).report()
    rows = {row["ref"]: row for row in report["items"]}
    assert rows["A.3"]["status"] == Status.PROVIDED.value, "the trial balance"
    assert rows["A.1"]["status"] == Status.PROVIDED.value, "the financial statements"
    assert rows["A.5"]["status"] == Status.PROVIDED.value, "the tax return"


def test_a_pdf_is_opened_because_most_of_what_arrives_is_one(folder: Path) -> None:
    """Without a one-page peek, a set of financial statements lands in "not matched"
    beside the holiday photos - which is the state this phase would have shipped in."""
    kinds = {Path(s.path).name: s.kind for s in look_at(LocalFolderStorage(folder))}
    assert kinds["FS 2021-2025.pdf"] == "financial_statements"
    assert kinds["2024 return.pdf"] == "tax_return"


def test_a_workbook_is_read_for_its_sheet_names_and_first_rows(folder: Path) -> None:
    kinds = {Path(s.path).name: s.kind for s in look_at(LocalFolderStorage(folder))}
    assert kinds["TB 2025 final.xlsx"] == "trial_balance"


def test_a_file_that_answers_nothing_is_reported_not_dropped(folder: Path) -> None:
    """The thing a person most needs to see: either a document nobody asked for, or an
    item the matching failed on. A coverage report only ever lists items, so neither
    would otherwise be visible at all."""
    report = tracker_for(folder).report()
    assert report["unmatched_files"] == ["holiday photos.zip"]
    assert report["files_seen"] == 4


def test_an_unreadable_file_is_still_seen(tmp_path: Path) -> None:
    """Dropping it would make an unreadable document indistinguishable from one that is
    not there, and the first is something to chase."""
    (tmp_path / "corrupt.xlsx").write_bytes(b"not a workbook")
    seen = look_at(LocalFolderStorage(tmp_path))
    assert [s.path for s in seen] == ["corrupt.xlsx"]


def test_a_file_the_client_filed_under_its_number_is_matched_by_it(tmp_path: Path) -> None:
    (tmp_path / "B.6 prior appraisal.pdf").write_bytes(build_pdf([("Appraisal", 12)]))
    report = tracker_for(tmp_path).report()
    rows = {row["ref"]: row for row in report["items"]}
    assert rows["B.6"]["status"] == Status.PROVIDED.value
    assert rows["B.6"]["inferred"] is False


# ============================================================ the workbook


def test_the_tracker_is_written_as_a_workbook_he_can_open(folder: Path, tmp_path: Path) -> None:
    """The 2026-09-18 decision: the request workbook *is* the tracker. One artefact, and
    for an engagement with no list at all it is also how the list comes into existence."""
    path = write_tracker(tracker_for(folder), tmp_path / "DD tracker.xlsx")
    assert path.is_file()
    book = load_workbook(path)
    try:
        assert book.sheetnames == ["Request list", "Not matched"]
        header = [c.value for c in book["Request list"][1]]
        for column in (*THEIRS, *DERIVED):
            assert column in header, column
    finally:
        book.close()


def test_an_inferred_match_is_visible_in_the_sheet_not_only_in_the_code(
    folder: Path, tmp_path: Path
) -> None:
    """The person reading the tracker is the one who can settle whether a document really
    answers an item. Telling them only in a docstring is telling nobody."""
    path = write_tracker(tracker_for(folder), tmp_path / "t.xlsx")
    rows = sheet_rows(path)
    assert rows["A.3"]["Matched by"] == "inferred from the document"


def test_a_file_filed_under_its_number_says_so_instead(tmp_path: Path) -> None:
    (tmp_path / "B.6 appraisal.pdf").write_bytes(build_pdf([("Appraisal", 12)]))
    path = write_tracker(tracker_for(tmp_path), tmp_path / "t.xlsx")
    assert sheet_rows(path)["B.6"]["Matched by"] == "filed under its number"


def test_the_unmatched_files_get_their_own_sheet(folder: Path, tmp_path: Path) -> None:
    """Not appended to the list, because they are not items. "What is still owed" and
    "what is this" are asked at different times by the same person."""
    path = write_tracker(tracker_for(folder), tmp_path / "t.xlsx")
    book = load_workbook(path)
    try:
        names = [row[0] for row in book["Not matched"].iter_rows(min_row=2, values_only=True)]
        assert names == ["holiday photos.zip"]
    finally:
        book.close()


def test_the_sheet_says_which_items_cannot_wait(folder: Path, tmp_path: Path) -> None:
    path = write_tracker(tracker_for(folder), tmp_path / "t.xlsx")
    rows = sheet_rows(path)
    assert rows["A.1"]["Essential"] == "Y"
    assert rows["C.2"]["Essential"] in ("", None)


# ============================================================ his notes survive


def write_into(path: Path, ref: str, column: str, value: str) -> None:
    """A person opening the tracker and typing in it."""
    book = load_workbook(path)
    try:
        sheet = book["Request list"]
        header = [c.value for c in sheet[1]]
        for row in sheet.iter_rows(min_row=2):
            if row[0].value == ref:
                row[header.index(column)].value = value
        book.save(path)
    finally:
        book.close()


def test_a_refresh_keeps_every_column_a_person_owns(folder: Path, tmp_path: Path) -> None:
    """The rule that decides whether this gets used. An update that cleared them would
    mean re-entering everything after every refresh, so they would stop being entered."""
    path = write_tracker(tracker_for(folder), tmp_path / "t.xlsx")
    write_into(path, "D.5", "Notes", "no union - confirmed with the controller")
    write_into(path, "D.5", "Not Applicable", "Y")
    write_into(path, "A.3", "In Model?", "Y")

    (folder / "extra.pdf").write_bytes(build_pdf([("Something new", 12)]))
    write_tracker(tracker_for(folder), path)

    rows = sheet_rows(path)
    assert rows["D.5"]["Notes"] == "no union - confirmed with the controller"
    assert rows["D.5"]["Not Applicable"] == "Y"
    assert rows["A.3"]["In Model?"] == "Y"


def test_a_person_s_not_applicable_decides_the_status(folder: Path, tmp_path: Path) -> None:
    """Their decision about the engagement, and the tool has no business forming a view
    on it - including when a file turns up that would otherwise satisfy the item."""
    path = write_tracker(tracker_for(folder), tmp_path / "t.xlsx")
    write_into(path, "A.3", "Not Applicable", "Y")

    held = read_tracker(path)
    refreshed = tracker_for(folder)
    refreshed.theirs = held
    write_tracker(refreshed, path)

    assert sheet_rows(path)["A.3"]["Status"] == Status.NOT_APPLICABLE.value


def test_the_two_sets_of_columns_do_not_overlap() -> None:
    """The other half of the preservation rule, and the reason it holds structurally
    rather than by care: a column in both sets would be written by the tool and expected
    by a person, and whichever ran last would win.

    Asserted on the two tuples because that is where it is decided. The behaviour below
    follows from it; this is what would have to change for the behaviour to stop.
    """
    assert not set(THEIRS) & set(DERIVED)


def test_a_stale_status_somebody_typed_over_is_replaced(
    folder: Path, tmp_path: Path
) -> None:
    """What the tool owns, it owns. A status left standing because somebody typed into
    the cell would be a tracker that lies, quietly, in the direction of being finished."""
    path = write_tracker(tracker_for(folder), tmp_path / "t.xlsx")
    write_into(path, "A.9", "Status", "provided")
    write_tracker(tracker_for(folder), path)
    assert sheet_rows(path)["A.9"]["Status"] == Status.OUTSTANDING.value


def test_a_first_run_has_nothing_to_preserve_and_says_so(tmp_path: Path) -> None:
    """Not a failure. An engagement that has never had a tracker is the ordinary case."""
    assert read_tracker(tmp_path / "nothing-here.xlsx") == {}


def test_a_workbook_that_is_not_a_tracker_is_not_read_as_one(tmp_path: Path) -> None:
    """Pointed at the wrong file, it takes nothing from it rather than reading column
    three as somebody's notes."""
    a_workbook(tmp_path / "something-else.xlsx", ("Name", "Amount"), ("Cash", 100))
    assert read_tracker(tmp_path / "something-else.xlsx") == {}


def test_two_refreshes_with_no_change_produce_the_same_thing(
    folder: Path, tmp_path: Path
) -> None:
    """Pressing it twice is the ordinary case - he presses update, then again because he
    is not sure the first one did anything."""
    path = write_tracker(tracker_for(folder), tmp_path / "t.xlsx")
    first = sheet_rows(path)
    write_tracker(tracker_for(folder), path)
    second = sheet_rows(path)
    for ref in first:
        for column in ("Status", "Matched by", "Evidence", *THEIRS):
            assert first[ref][column] == second[ref][column], (ref, column)


# ============================================================ what it does not do


def test_nothing_here_reaches_a_model_or_an_ocr_engine() -> None:
    """Shallow on purpose. A tracker that took four minutes is a tracker nobody
    refreshes, and OCR on a button press over sixty documents is not a button press."""
    import ast

    tree = ast.parse((ROOT / "intake" / "tracker.py").read_text(encoding="utf-8"))
    named: set[str] = set()
    for node in ast.walk(tree):
        if isinstance(node, ast.Import):
            named.update(a.name for a in node.names)
        elif isinstance(node, ast.ImportFrom) and node.module:
            named.add(node.module)
    assert named
    for forbidden in ("ports.llm", "ports.ocr", "harness", "skills"):
        assert not any(n == forbidden or n.startswith(forbidden + ".") for n in named), forbidden


def test_the_evidence_is_the_path_so_provided_can_be_checked(
    folder: Path, tmp_path: Path
) -> None:
    path = write_tracker(tracker_for(folder), tmp_path / "t.xlsx")
    assert "TB 2025 final.xlsx" in sheet_rows(path)["A.3"]["Evidence"]


def test_a_list_that_is_a_starting_point_is_still_reported_as_one(folder: Path) -> None:
    """Carried all the way through: a coverage report against a guessed list and one
    against the firm's own list are different sentences."""
    assert tracker_for(folder).report()["is_a_starting_point"] is True


def test_an_empty_folder_leaves_everything_outstanding(tmp_path: Path) -> None:
    """A Monday morning before anything has arrived. Every item outstanding is the
    correct answer and no items at all would not be."""
    report = tracker_for(tmp_path).report()
    assert report["files_seen"] == 0
    assert report["counts"][Status.PROVIDED.value] == 0
    assert report["counts"][Status.OUTSTANDING.value] > 30


def test_a_list_with_one_item_still_writes_a_tracker(tmp_path: Path) -> None:
    """The degenerate case, because a firm's real list will arrive and be shorter or
    stranger than the default and must not need code changed for it."""
    from intake.request_list import ItemRef

    tiny = RequestList([RequestItem(ItemRef("A", 1), "anything")], name="tiny")
    path = write_tracker(Tracker(tiny, (Seen("x.pdf", None),)), tmp_path / "t.xlsx")
    assert list(sheet_rows(path)) == ["A.1"]


# ============================================================ the button


def test_the_page_offers_a_way_to_run_it_and_to_see_what_it_would_do() -> None:
    """A preview, because the first thing anybody wants to know about a tool that writes
    into their workbook is what it is about to do to it."""
    from admin.page import PAGE

    for action in ("updateTracker", "previewTracker"):
        assert f'data-act="{action}"' in PAGE, action
        assert f"ACTIONS.{action}" in PAGE, action


def test_the_page_says_that_what_a_person_typed_is_left_alone() -> None:
    """The promise the whole design rests on belongs where somebody will read it before
    pressing, not in a docstring."""
    from admin.page import PAGE

    assert "left exactly" in PAGE
    assert "Not Applicable" in PAGE


def test_the_page_warns_when_the_list_is_only_a_starting_point() -> None:
    """Otherwise "nine of twenty provided" reads as a fact about the engagement rather
    than a fact about a guess."""
    from admin.page import PAGE

    assert "starting point, not the firm" in PAGE


def console_over(folder: Path) -> Any:
    """The smallest console that can run the tracker over a real folder.

    Built rather than fixtured because everything else it does is beside the point here -
    what matters is that the route's `dry_run` reaches the write and stops it.
    """
    from unittest.mock import Mock

    from admin.console import AdminConsole

    entry = Mock()
    entry.scope_path = "valuation/non-esop"
    preflight = Mock()
    preflight.storage = LocalFolderStorage(folder)
    console = AdminConsole(Mock(), Mock(), preflight, constraints=Mock())
    console._folder_of = Mock(return_value=(entry, ""))  # type: ignore[method-assign]
    return console


def test_a_dry_run_reads_the_folder_and_writes_no_file(folder: Path, tmp_path: Path) -> None:
    """The preview. Somebody pressing it wants to know what is about to happen to their
    workbook, and finding out by having it happen is not that."""
    target = tmp_path / "would-be.xlsx"
    report = console_over(folder).dd_tracker("ENG-1", str(target), dry_run=True)

    assert report["dry_run"] is True
    assert report["counts"][Status.PROVIDED.value] > 0, "it really read the folder"
    assert not target.exists(), "a preview that writes is not a preview"


def test_the_same_call_without_dry_run_writes_it(folder: Path, tmp_path: Path) -> None:
    """And the preview is the same path minus the write, so the two cannot disagree about
    what would have happened."""
    # Separate files on purpose. Pointed at the same one, a preview that wrongly wrote
    # would leave a file for the real run to be credited with, and this would pass.
    previewed = tmp_path / "previewed.xlsx"
    written = tmp_path / "written.xlsx"
    shown = console_over(folder).dd_tracker("ENG-1", str(previewed), dry_run=True)
    done = console_over(folder).dd_tracker("ENG-1", str(written), dry_run=False)

    assert written.is_file()
    assert not previewed.exists()
    assert done["counts"] == shown["counts"]
    assert done["unmatched_files"] == shown["unmatched_files"]


def test_the_route_is_a_post_because_it_writes_a_file() -> None:
    source = (ROOT / "admin" / "app.py").read_text(encoding="utf-8")
    assert '@app.post("/api/admin/engagements/{engagement_id}/dd-tracker")' in source


def test_a_scope_with_no_list_is_refused_rather_than_given_another_division_s() -> None:
    """A request list is what the client was asked for. Borrowing one would put items in
    front of a client that nobody at the firm chose to ask for."""
    from unittest.mock import Mock

    from admin.console import AdminConsole, AdminRefused

    console = AdminConsole(Mock(), Mock(), Mock(), constraints=Mock())
    entry = Mock()
    entry.scope_path = "tax/compliance"
    console._folder_of = Mock(return_value=(entry, "somewhere"))  # type: ignore[method-assign]
    console.config.document = Mock(return_value={})
    with pytest.raises(AdminRefused) as refused:
        console.dd_tracker("ENG-1")
    assert "does not borrow another division" in refused.value.diagnostic.detail
