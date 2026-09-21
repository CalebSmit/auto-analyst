"""P91 - A report is a document with named places in it.

Two halves, and they fail differently. Getting the *place* wrong puts Appendix F into
Appendix E and looks fine. Getting the *move* wrong produces a document that opens without
complaint and is laid out wrong, and also looks fine. Neither is caught by reading the
result on screen, which is why they are caught here.
"""

from __future__ import annotations

import ast
from io import BytesIO
from pathlib import Path
from typing import Any

import pytest
from docx import Document

from registries.errors import RegistryError
from report.anchors import BOOKMARK, HEADING, anchors_in, find_anchor, normalise, span_of
from report.appendices import APPENDICES, appendix, registry_from, updatable
from report.transplant import Transplant, transplant
from tests.docs import a_report, appendix_source, bookmark, list_ids_in, text_of

ROOT = Path(__file__).resolve().parents[2]

C = "Appendix C - Economic Conditions"
E = "Appendix E - Summary of Major Market Indices"


def blocks_of(document: Any) -> tuple[Any, ...]:
    return tuple(b for b in document.element.body.iterchildren() if isinstance(b.tag, str))


def styles_in(document: Any) -> set[str]:
    W = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"
    return {
        s.get(f"{W}styleId") for s in document.styles.element.findall(f"{W}style")
    }


def reopened(document: Any) -> Any:
    """Save and load it again. A document that only holds together in memory is not one."""
    buffer = BytesIO()
    document.save(buffer)
    return Document(BytesIO(buffer.getvalue()))


# ============================================================ binding by name


def test_an_appendix_is_found_by_its_heading() -> None:
    found = find_anchor(a_report(), C)
    assert found.kind == HEADING
    assert found.name == C


def test_a_bookmark_beats_a_heading_of_the_same_name() -> None:
    """A bookmark is what somebody put there deliberately; a heading is what happened to
    be written. When both exist the deliberate one wins."""
    found = find_anchor(a_report(bookmarks=True), C)
    assert found.kind == BOOKMARK


def test_an_anchor_that_is_not_there_is_a_refusal_naming_it() -> None:
    """Never the closest thing and never the end of the document. The whole value of
    binding by name is that it is not guessed at."""
    with pytest.raises(RegistryError) as refused:
        find_anchor(a_report(), "Appendix Q - Something Nobody Wrote")
    d = refused.value.diagnostic
    assert d.code == "CONFIG.BINDING_MISSING"
    assert "Appendix Q" in d.detail
    assert "guessed" in d.detail


def test_the_refusal_shows_what_is_actually_there() -> None:
    """Usually the anchor is spelled differently, and somebody reading the error needs to
    see that rather than be told to go and look."""
    with pytest.raises(RegistryError) as refused:
        find_anchor(a_report(), "Appendix C - Economic Condition")
    assert "Appendix C" in (refused.value.diagnostic.found or "")


def test_two_years_of_the_same_heading_are_the_same_heading() -> None:
    """Case, the kind of dash, and the amount of whitespace differ between one year's copy
    and the next without meaning anything."""
    assert normalise("Appendix C — Economic  Conditions") == normalise(
        "appendix c - economic conditions"
    )


def test_matching_is_not_fuzzy_beyond_that() -> None:
    """A near-match that quietly picked a different heading would put an appendix in the
    wrong place, which is exactly what binding by name exists to prevent."""
    assert normalise("Appendix C - Economic Conditions") != normalise(
        "Appendix C - Economic Outlook"
    )
    with pytest.raises(RegistryError):
        find_anchor(a_report(), "Appendix C")


def test_a_paragraph_index_is_not_an_anchor() -> None:
    """The rule, stated as a property of the code rather than of a docstring: `find_anchor`
    takes a name, and there is no integer path into it."""
    with pytest.raises(RegistryError):
        find_anchor(a_report(), "3")

    source = (ROOT / "report" / "anchors.py").read_text(encoding="utf-8")
    tree = ast.parse(source)
    signature = next(
        node
        for node in ast.walk(tree)
        if isinstance(node, ast.FunctionDef) and node.name == "find_anchor"
    )
    named = [a.arg for a in signature.args.args]
    assert named == ["document", "name"], "no index argument exists to pass"
    assert signature.args.args[1].annotation is not None
    assert getattr(signature.args.args[1].annotation, "id", "") == "str"


# ============================================================ what an anchor owns


def test_a_heading_owns_its_sub_headings_but_not_the_next_appendix() -> None:
    """A sub-heading inside Appendix C belongs to Appendix C. Appendix E does not, and an
    update that swallowed it would delete an appendix nobody asked it to touch."""
    report = a_report()
    owned = span_of(report, find_anchor(report, C))
    text = "\n".join(
        "".join(t.text or "" for t in b.iter("{http://schemas.openxmlformats.org/wordprocessingml/2006/main}t"))
        for b in owned
    )
    assert "A sub-heading inside the appendix" in text
    assert "Appendix E" not in text


def test_the_anchor_s_own_block_is_part_of_what_it_owns() -> None:
    """Replacing an appendix replaces its heading, because the new content brings one."""
    report = a_report()
    anchor = find_anchor(report, C)
    assert span_of(report, anchor)[0] is anchor.element


def test_a_bookmarked_span_stops_at_the_next_bookmark() -> None:
    """A bookmark has no depth to compare, so "this bit" ends at the next named place of
    any kind - and that includes another bookmark with no heading between them.

    Built here rather than taken from the report fixture, because in that document every
    bookmark sits on a heading and the heading rule stops the span first. A test that used
    it would pass with the bookmark rule deleted, which is the shape of test this repo
    keeps having to rewrite.
    """
    document = Document()
    first = document.add_paragraph("The part the first bookmark marks.")
    bookmark(first, "FIRST", number=1)
    document.add_paragraph("Still the first part.")
    second = document.add_paragraph("Where the second bookmark starts.")
    bookmark(second, "SECOND", number=2)
    document.add_paragraph("Which is not the first part.")

    owned = span_of(document, find_anchor(document, "FIRST"))
    text = _text(owned)
    assert "Still the first part." in text
    assert "Where the second bookmark starts." not in text
    assert len(owned) == 2


def _W(tag: str) -> str:
    return f"{{http://schemas.openxmlformats.org/wordprocessingml/2006/main}}{tag}"


def _text(blocks: Any) -> str:
    """The text of a span, for asserting what an anchor owns and what it does not."""
    return chr(10).join("".join(t.text or "" for t in b.iter(_W("t"))) for b in blocks)


def test_every_named_place_is_listed_in_document_order() -> None:
    names = [a.name for a in anchors_in(a_report())]
    assert names.index(C) < names.index(E)


def test_word_s_own_hidden_bookmark_is_not_a_named_place() -> None:
    """`_GoBack` is in most documents and means nothing. Treating it as an anchor would
    put it in the list a person chooses from."""
    report = a_report()
    bookmark(report.paragraphs[1], "_GoBack", number=99)
    assert "_GoBack" not in [a.name for a in anchors_in(report)]


# ============================================================ the move


def test_the_content_arrives_where_the_anchor_was() -> None:
    report, source = a_report(), appendix_source()
    anchor = find_anchor(report, C)
    transplant(source, report, blocks_of(source), after=anchor.element)

    body = text_of(report)
    assert "Economic conditions in the period under review." in body
    # After the anchor it was bound to, and before the next appendix. Both halves matter:
    # inserting before the anchor would put Appendix C's content under the report body.
    assert body.index("Appendix C") < body.index("Economic conditions in the period")
    assert body.index("Economic conditions in the period") < body.index("Appendix E")


def test_a_style_the_report_does_not_define_arrives_with_it() -> None:
    """A paragraph carries a style *id*, not a style. Without the definition Word renders
    it as Normal - so a market commentary arrives as body text and nobody can say why."""
    report, source = a_report(), appendix_source()
    assert "MarketCommentary" not in styles_in(report)

    carried = transplant(source, report, blocks_of(source), after=find_anchor(report, C).element)
    assert "MarketCommentary" in carried.styles
    assert "MarketCommentary" in styles_in(reopened(report))


def test_a_style_the_report_already_has_is_left_alone() -> None:
    """The destination's definition is the firm's formatting and the arriving content is
    the guest. Overwriting Heading 2 with the source's would reformat the whole report to
    match one appendix."""
    report, source = a_report(), appendix_source()
    before = len(styles_in(report))
    carried = transplant(source, report, blocks_of(source), after=find_anchor(report, C).element)
    assert "Heading2" not in carried.styles
    assert len(styles_in(report)) == before + len(carried.styles)


def test_an_arriving_list_does_not_join_a_list_already_in_the_report() -> None:
    """The failure this is all for. Both documents claim numId 10 for different lists;
    left alone, Appendix C's first bullet becomes item 3 of the report's own list, which
    reads as a mistake somebody made."""
    report, source = a_report(), appendix_source()
    assert set(list_ids_in(report)) == set(list_ids_in(source)), "the ids really do collide"

    transplant(source, report, blocks_of(source), after=find_anchor(report, C).element)
    ids = list_ids_in(report)
    assert len(set(ids)) == 2, "the report's list and the arriving one are different lists"


def test_the_arriving_list_still_has_numbering_at_all() -> None:
    """The other way it goes wrong: remapped to nothing, so the numbers vanish."""
    report, source = a_report(), appendix_source()
    carried = transplant(source, report, blocks_of(source), after=find_anchor(report, C).element)
    assert carried.numbering
    assert len(list_ids_in(reopened(report))) == 5


def test_an_image_arrives_pointing_at_its_own_picture() -> None:
    """A relationship id is per-document. Unremapped, the image in an appendix becomes
    whatever the report's rId7 was, which is usually the letterhead."""
    report, source = a_report(), appendix_source()
    before = len(report.part.rels)
    carried = transplant(source, report, blocks_of(source), after=find_anchor(report, C).element)
    assert carried.images
    assert len(report.part.rels) > before
    assert reopened(report) is not None


def test_a_table_survives_the_move() -> None:
    report, source = a_report(), appendix_source()
    transplant(source, report, blocks_of(source), after=find_anchor(report, C).element)
    assert "A broad market index" in text_of(reopened(report))


def test_every_reference_in_the_saved_file_resolves() -> None:
    """What "Word would open it" actually means, checked rather than hoped.

    A relationship id that survives into the body without a matching entry in the saved
    package is the failure that renders as an empty box, or as the letterhead where an
    exhibit should be. Saving and reloading is the cheapest way to ask, and asking about
    the *ids* rather than about the text is what makes it a real question.
    """
    R = "{http://schemas.openxmlformats.org/officeDocument/2006/relationships}"
    report, source = a_report(), appendix_source()
    transplant(source, report, blocks_of(source), after=find_anchor(report, C).element)

    again = reopened(report)
    referenced = {
        str(value)
        for node in again.element.body.iter()
        for key, value in node.attrib.items()
        if key.startswith(R) and str(value).startswith("rId")
    }
    assert referenced, "no references survived at all, so this proves nothing"
    assert referenced <= set(again.part.rels), sorted(referenced - set(again.part.rels))
    assert "Economic conditions in the period under review." in text_of(again)
    assert "Appendix H - Certification" in text_of(again)


def test_two_appendices_moved_in_one_pass_share_one_list(tmp_path: Path) -> None:
    """The reason `Transplant` holds state: two paragraphs of the same list have to end up
    with the same new numId, and a function that rebuilt the table per block would split
    one list into three."""
    report, source = a_report(), appendix_source()
    move = Transplant(source, report)
    carried = move.move(blocks_of(source), after=find_anchor(report, C).element)
    arriving = [i for i in list_ids_in(report) if i not in set(list_ids_in(a_report()))]
    assert carried.numbering
    assert len(set(arriving)) == 1


# ============================================================ what could not be carried


def test_something_that_cannot_be_carried_is_named_rather_than_dropped() -> None:
    """Silently losing an exhibit is how a report goes out with a hole in it. The
    relationship kinds this code knows are images and hyperlinks; anything else is
    reported with enough context to put it in by hand.
    """
    W = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"
    R = "{http://schemas.openxmlformats.org/officeDocument/2006/relationships}"
    report, source = a_report(), appendix_source(image=False)

    paragraph = source.add_paragraph("A chart lives here.")
    holder = paragraph._p.makeelement(f"{W}object", {})
    holder.set(f"{R}id", "rIdNoSuchThing")
    paragraph._p.append(holder)

    carried = transplant(source, report, blocks_of(source), after=find_anchor(report, C).element)
    assert not carried.complete
    assert any("rIdNoSuchThing" in note for note in carried.left_behind)


def test_a_carried_result_says_what_it_did_in_words() -> None:
    report, source = a_report(), appendix_source()
    carried = transplant(source, report, blocks_of(source), after=find_anchor(report, C).element)
    assert "block(s)" in carried.summary()
    assert carried.complete


# ============================================================ the registry


def test_three_appendices_have_an_update_and_three_do_not() -> None:
    """G, H and I carry forward in the copy of last year's report. Listed rather than
    omitted, so "there is no button for the certification" is an answer rather than a gap.
    """
    assert {a.id for a in updatable()} == {"economic", "indices", "ma"}
    assert {a.letter for a in APPENDICES if not a.updated} == {"G", "H", "I"}


def test_nothing_in_the_package_can_write_the_certification() -> None:
    """The hardest non-negotiable in the blueprint, satisfied by construction.

    Appendix H is not marked updatable, and the update path only ever reaches an appendix
    it got from `updatable()`. There is no code that writes it - which is stronger than a
    guard that refuses to.
    """
    certification = appendix("certification")
    assert certification.letter == "H"
    assert not certification.updated
    assert certification not in updatable()


def test_an_unknown_appendix_is_refused_and_the_message_tells_the_two_cases_apart() -> None:
    """"There is no appendix called that" and "that appendix is not one anything updates"
    send a person to different places."""
    with pytest.raises(RegistryError) as refused:
        appendix("nonsense")
    detail = refused.value.diagnostic.detail
    assert "economic" in detail
    assert "carry forward" in detail


def test_the_registry_can_be_rebuilt_from_configuration() -> None:
    """The real report has not been seen. When a heading turns out to be worded
    differently, correcting it has to be a setting rather than an edit to a module."""
    built = registry_from(
        [
            {"id": "economic", "letter": "C", "title": "Economic", "anchor": "APPX_C"},
            {"id": "indices", "letter": "E", "title": "Indices", "anchor": "APPX_E",
             "updated": False},
        ]
    )
    assert [a.anchor for a in built] == ["APPX_C", "APPX_E"]
    assert [a.id for a in updatable(built)] == ["economic"]


def test_an_appendix_with_no_anchor_is_refused() -> None:
    """It would have nowhere to put its content, and the failure would arrive later and
    read as a missing bookmark in the report."""
    with pytest.raises(RegistryError) as refused:
        registry_from([{"id": "x", "letter": "X", "title": "X", "anchor": ""}])
    assert "anchor" in refused.value.diagnostic.detail


def test_two_appendices_cannot_share_an_id() -> None:
    """The id is what a button names, so a duplicate makes one unreachable and the other
    ambiguous."""
    with pytest.raises(RegistryError):
        registry_from(
            [
                {"id": "same", "letter": "C", "title": "A", "anchor": "a"},
                {"id": "same", "letter": "E", "title": "B", "anchor": "b"},
            ]
        )
