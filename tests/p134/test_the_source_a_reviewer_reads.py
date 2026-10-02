"""P134: the review modal told you to read a source it did not show.

Opened a real review item on the console. The modal asked *"IS formulas.xlsx: what is this
document?"*, gave the reason - *"no recognisable cue in the headers or first rows"* - and
then said, in the box above the decision buttons:

    Read the source first.

There was no source above it. The excerpt block is conditional on `item.source_excerpt`
and that sentence was not, so on an item with no excerpt the modal gave advice nobody
could follow, on the one question that can *only* be answered by looking at the file.

**And the excerpt was there all along.** `intake/readers/__init__.py` writes
`"headers": headers[:12]` into the item's detail - the twelve headers the classifier
itself looked at - and `ReviewSurface.show` reads the excerpt from `source_excerpt`. Two
names for one thing, so twelve headers were captured for a reviewer and never reached one.
This build's most repeated defect, in the one place where a person is being asked to
disagree with the system and cannot see what the system saw.

Both halves are fixed here, and both are tested: the producer writes the excerpt under the
name the surface reads, and the modal stops giving advice it cannot support when an item
genuinely has none.
"""

from __future__ import annotations

import re

from ops.console.app import PAGE as ACTIVITY


def script() -> str:
    return "\n".join(re.findall(r"<script>(.*?)</script>", ACTIVITY, re.S))


# ---------------------------------------------- the headers reach the reviewer


def test_the_classifier_writes_the_excerpt_under_the_name_the_surface_reads() -> None:
    """Asserted across the two modules, because the defect was that each was correct on
    its own: the producer wrote `headers`, the surface read `source_excerpt`, and nothing
    compared them."""
    import ast
    from pathlib import Path

    root = Path(__file__).resolve().parents[2]
    produced = (root / "intake" / "readers" / "__init__.py").read_text(encoding="utf-8")
    assert '"source_excerpt"' in produced, "the classifier records no excerpt"
    assert '"locator"' in produced, "and no locator, so the file cannot be named"

    surface = (root / "review" / "surface.py").read_text(encoding="utf-8")
    tree = ast.parse(surface)
    read = {
        node.args[0].value
        for node in ast.walk(tree)
        if isinstance(node, ast.Call)
        and isinstance(node.func, ast.Attribute)
        and node.func.attr == "get"
        and node.args
        and isinstance(node.args[0], ast.Constant)
    }
    assert "source_excerpt" in read, "the surface no longer reads an excerpt"
    assert "locator" in read


def test_a_classified_document_carries_the_headers_it_was_judged_on(
    started: str, store: object, parse_store: object, boundary: object, bus: object
) -> None:
    """End to end through the real reader and the real surface.

    A workbook whose headers say nothing recognisable is classified `other` and goes to a
    person - and the person can now see the headers it was judged on, which is the whole
    of what they need to disagree with it.
    """
    from io import BytesIO

    from openpyxl import Workbook

    from intake.readers import read_source
    from review.surface import ReviewSurface
    from tests.p17.conftest import intake_row

    wb = Workbook()
    sheet = wb.active
    assert sheet is not None
    sheet.title = "Sheet1"
    for column, header in enumerate(("Widget", "Qty", "Zone", "Ref"), start=1):
        sheet.cell(row=1, column=column, value=header)
    sheet.cell(row=2, column=1, value="blue")
    sheet.cell(row=2, column=2, value=4)
    buffer = BytesIO()
    wb.save(buffer)
    data = buffer.getvalue()

    report = read_source(
        started,
        "HEND-2026-04",
        intake_row("Widgets.xlsx", data, "spreadsheet"),
        data,
        store=store,  # type: ignore[arg-type]
        parse_store=parse_store,  # type: ignore[arg-type]
        boundary=boundary,  # type: ignore[arg-type]
        bus=bus,  # type: ignore[arg-type]
    )
    assert report.classification.document_type == "other", (
        "the fixture is meant to be unrecognisable; it classified as "
        f"{report.classification.document_type}"
    )
    assert report.review_item_id, "it did not go to a person"

    surface = ReviewSurface(store, bus)  # type: ignore[arg-type]
    shown = surface.show(started, report.review_item_id)
    assert shown.source_excerpt, "the reviewer is shown no source at all"
    for header in ("Widget", "Qty", "Zone", "Ref"):
        assert header in shown.source_excerpt, shown.source_excerpt
    assert "Widgets.xlsx" in shown.locator, "the excerpt cannot be traced to a file"
    assert shown.reasons, "and it still says why it came to a person"


def test_the_excerpt_is_the_headers_joined_rather_than_a_repr() -> None:
    """A reviewer reads a line, not a Python list. `['Widget', 'Qty']` on screen is the
    shape of the data leaking into the one place it has to be prose."""
    from pathlib import Path

    produced = (Path(__file__).resolve().parents[2] / "intake" / "readers" / "__init__.py")
    body = produced.read_text(encoding="utf-8")
    assert '" | ".join(h for h in headers[:12] if h.strip())' in body
    assert '"source_excerpt": headers' not in body, "the list itself would render as a repr"


# ---------------------------------------- and the modal stops promising a source


def test_the_modal_only_says_read_the_source_when_there_is_one() -> None:
    body = script()
    found = re.search(r'<div id="rv-proposal">(.*?)</div></div>', body, re.S)
    assert found, "the reveal box is gone"
    reveal = found.group(1)
    assert "item.source_excerpt" in reveal, "the sentence does not depend on there being one"
    assert "Read the source first." in reveal


def test_when_there_is_no_excerpt_it_says_what_to_do_instead() -> None:
    """Not just silence. An item with nothing above it still has a file behind it, and the
    answer is to go and open it."""
    body = script()
    found = re.search(r'<div id="rv-proposal">(.*?)</div></div>', body, re.S)
    assert found
    reveal = found.group(1)
    assert "No excerpt was recorded" in reveal
    assert "open the file itself" in reveal


def test_the_reveal_discipline_is_untouched() -> None:
    """The whole point of that box is that answering without looking is visible later.
    Changing what it says must not change what it gates."""
    body = script()
    assert 'data-act="reveal"' in body
    assert 'data-d="accepted" disabled' in body
    assert 'data-d="corrected" disabled' in body
    assert "Show the proposal first" in body


def test_the_excerpt_is_shown_with_the_locator_when_the_item_has_one() -> None:
    """The file the headers came from. An excerpt with no locator is a line of text a
    reviewer cannot trace back to anything."""
    body = script()
    assert "What the source says" in body
    assert "item.locator" in body
