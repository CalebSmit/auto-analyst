"""P101 - version 1's own recipe, ending at a workbook nobody has to trust.

The analyst asked to set an engagement up at the dashboard, point it at a client folder
and a request list, set a spend limit, and get a workbook. Auditing that against the build
found there was no path at all, for three separate reasons:

* **All three shipped recipes end at `qc-tieout`**, which CLAUDE.md puts outside version 1.
* **`qc-tieout` and `ddr-tagging` are skeletons** - frontmatter and no mechanical half. The
  driver's own docstring says a run reaching one "stops there with the step named". So
  `full-ddr-intake` could never finish either, for a second and unrelated reason.
* **The two recipes that reach a release step write through the firm's model template**,
  and version 1 must never touch the real model.

`valuation-input-prep` is version 1's scope as a graph: intake, masking, timing, tagging,
and a mapping onto an input-help workbook. Every skill in it has a procedure. It ends at a
workbook that opens on a sheet saying it is not the model.

What this phase does **not** do is make the output good. It makes a run able to finish, so
that what comes out can be looked at and argued with. A recipe that cannot complete produces
no evidence about anything.
"""

from __future__ import annotations

from io import BytesIO
from pathlib import Path

import pytest
import yaml
from openpyxl import load_workbook

from graph.compile import PlanContext, compile_recipe
from graph.recipes import load_recipe
from graph.resolve import ScopeResolver
from harness.procedures import PROCEDURES

ROOT = Path(__file__).resolve().parents[2]
RECIPES = ROOT / "recipes"
RECIPE = "valuation-input-prep"
FOCUS = "valuation/non-esop/non-repeat"
WORKBOOK = ROOT / "templates" / "Chartwell_Input_Help_v1.xlsx"


def recipe(name: str) -> dict:
    return yaml.safe_load((RECIPES / f"{name}.yaml").read_text(encoding="utf-8"))


def skills_of(doc: dict) -> list[str]:
    return [s["skill"] for s in doc["steps"] if s.get("kind") == "skill"]


def plan_for(scope_path: str = FOCUS):
    resolved = ScopeResolver().resolve(scope_path)
    return compile_recipe(
        load_recipe(RECIPE),
        PlanContext(
            engagement_id="E1",
            scope=resolved,
            taxonomy_version="coa-taxonomy@1.0.0",
            precedent_snapshot="snap-1",
        ),
    )


# ============================================================ it can finish


def test_the_recipe_compiles_against_the_focus_scope() -> None:
    plan = plan_for()
    assert [s.id for s in plan.steps] == [
        "intake", "mask", "timing", "tagging", "model", "release"
    ]


def test_every_skill_it_names_has_a_mechanical_half() -> None:
    """The exit criterion, and the reason the other three recipes cannot finish. A skill
    with no procedure raises `KeyError: the skill has no mechanical half yet` the moment a
    run reaches it, and the run stops there with the step named."""
    for skill in skills_of(recipe(RECIPE)):
        assert skill in PROCEDURES, f"{skill} is a skeleton; a run would stop at it"


def test_the_recipe_ends_at_a_release_step_that_depends_on_a_mapping() -> None:
    """`release` builds from the `model` step's output. `full-ddr-intake` has a release
    step and no model step, so even if its skeletons were written it would fail there."""
    doc = recipe(RECIPE)
    steps = {s["id"]: s for s in doc["steps"]}
    # A floor, not an inventory. This asserted `== ["model"]`, and P111 found that the
    # release handler also asks for `timing` - it resolves every figure_ref against a
    # catalogue built from timing's output - so the exact list was a fact about the day
    # this was written. What P101 cared about is that release depends on a mapping at all,
    # which is why `full-ddr-intake` cannot finish; B-066 again.
    assert "model" in steps["release"]["needs"]
    assert steps["model"]["skill"] == "model-load"


@pytest.mark.parametrize("name", ["full-ddr-intake", "monthly-refresh", "mature-company-valuation"])
def test_the_recipes_that_came_before_still_cannot_finish(name: str) -> None:
    """Asserted rather than assumed, because P101's whole premise is that they cannot, and
    a phase built on a premise nobody checked is how this build has gone wrong before.

    They are not deleted: they are the shapes the firm will want once the skills behind
    them exist, and `mature-company-valuation` is the one the compiler's parallel-branch
    handling is exercised by.
    """
    missing = [s for s in skills_of(recipe(name)) if s not in PROCEDURES]
    assert missing, f"{name} now has every procedure - P101's premise has expired"
    assert "qc-tieout" in missing


# ============================================================ what it writes


def test_the_focus_scope_resolves_to_the_input_help_workbook() -> None:
    assert ScopeResolver().resolve(FOCUS).template.file == "Chartwell_Input_Help_v1.xlsx"


def test_a_repeat_engagement_resolves_to_it_too() -> None:
    """The binding is at the segment layer, so both dispositions get it. A repeat client
    still fills an input-help workbook - what changes for a repeat is where the *structure*
    comes from, which is B-048's other half and not this phase."""
    assert (
        ScopeResolver().resolve("valuation/non-esop/repeat").template.file
        == "Chartwell_Input_Help_v1.xlsx"
    )


def test_binding_the_focus_scope_did_not_move_anything_else() -> None:
    """A binding at one layer changes every scope under it and nothing above. ESOP is
    parked rather than gone, and a phase aimed at non-ESOP that quietly re-pointed ESOP
    would be the kind of change nobody notices until an ESOP run writes the wrong file."""
    assert (
        ScopeResolver().resolve("valuation/esop/repeat").template.file
        == "Chartwell_ESOP_Model_v9.xlsx"
    )
    assert (
        ScopeResolver().resolve("consulting/tbd/repeat").template.file
        == "Chartwell_Master_Model_v14.xlsx"
    ), "consulting still falls through to the firm default"


# ============================================================ it is not a model


def test_the_workbook_says_on_its_face_that_it_is_not_the_model() -> None:
    """The first sheet is what opens. A file shaped like the firm's model that only admits
    otherwise on the second tab is one filename away from being used as one."""
    book = load_workbook(WORKBOOK)
    try:
        assert book.sheetnames[0] == "Not the model"
        said = str(book["Not the model"]["A1"].value)
        assert "not the Chartwell model" in said
    finally:
        book.close()


def test_the_workbook_computes_nothing() -> None:
    """The difference between the thing the analyst runs and the thing this build fills in
    front of him. The model templates beside it carry Historicals and DCF sheets of
    formulas over their named ranges; a workbook that computed anything would be the thing
    version 1 is defined against."""
    book = load_workbook(WORKBOOK)
    try:
        formulas = [
            c.value
            for sheet in book.worksheets
            for row in sheet.iter_rows()
            for c in row
            if isinstance(c.value, str) and c.value.startswith("=")
        ]
        assert formulas == []
    finally:
        book.close()


def test_the_model_templates_do_compute__which_is_how_they_differ() -> None:
    """The other side of the same claim, so that "no formulas" is a property of this
    workbook rather than of the synthesizer having stopped writing them."""
    book = load_workbook(ROOT / "templates" / "Chartwell_Master_Model_v14.xlsx")
    try:
        assert any(
            isinstance(c.value, str) and c.value.startswith("=")
            for sheet in book.worksheets
            for row in sheet.iter_rows()
            for c in row
        )
    finally:
        book.close()


def test_every_declared_input_has_a_named_range_to_be_written_into() -> None:
    """Bind by named range, never by cell address - the blueprint's rule, and the reason a
    row inserted into the Input tab does not silently move every figure by one."""
    schema = yaml.safe_load(
        (ROOT / "templates" / "inputs" / "Chartwell_Input_Help_v1@1.0.yaml").read_text(
            encoding="utf-8"
        )
    )
    book = load_workbook(WORKBOOK)
    try:
        names = set(book.defined_names)
        for entry in [*schema["required"], *schema.get("optional", [])]:
            assert entry["name"] in names, entry["name"]
        for stamp in schema["stamps"]:
            assert stamp in names, stamp
    finally:
        book.close()


def test_it_does_not_declare_the_inputs_version_one_cannot_produce() -> None:
    """Terminal WACC, target net working capital and the selected multiple belong to
    `cost-of-capital`, `dcf-build` and `gpc-scrub` - all three outside version 1 and all
    three skeletons. Declaring them would put three inputs on every run's unmapped list
    that nothing was ever going to fill."""
    schema = yaml.safe_load(
        (ROOT / "templates" / "inputs" / "Chartwell_Input_Help_v1@1.0.yaml").read_text(
            encoding="utf-8"
        )
    )
    declared = {e["name"] for e in [*schema["required"], *schema.get("optional", [])]}
    for model_only in ("WACC_Terminal", "NWC_Target", "EV_EBITDA_Multiple"):
        assert model_only not in declared


def test_the_synthesizer_rebuilds_it_exactly() -> None:
    """Deterministic: the same schema produces the same workbook, so the file on disk is
    the schema rather than something somebody edited by hand once."""
    import sys

    sys.path.insert(0, str(ROOT / "tools"))
    from synthesize_templates import build_input_help

    schema = yaml.safe_load(
        (ROOT / "templates" / "inputs" / "Chartwell_Input_Help_v1@1.0.yaml").read_text(
            encoding="utf-8"
        )
    )
    rebuilt = load_workbook(BytesIO(build_input_help(schema)))
    on_disk = load_workbook(WORKBOOK)
    try:
        assert rebuilt.sheetnames == on_disk.sheetnames
        assert set(rebuilt.defined_names) == set(on_disk.defined_names)
    finally:
        rebuilt.close()
        on_disk.close()


# ============================================================ what it asks for


def test_it_requires_only_what_a_valuation_can_begin_on() -> None:
    """The three earlier recipes each demand a trial balance. The first real client folder
    had none, and a required type nobody sends is a run that never releases."""
    assert recipe(RECIPE)["intake"]["required_types"] == ["financial_statements"]


def test_the_request_list_is_not_declared_as_a_second_run_input() -> None:
    """Written the other way round first, and the plan assertions refused the run:

        PROCESS.PLAN_REFUSED at intake - intake needs 'request_schedule' and nothing has
        produced it (have: engagement_profile, source_documents)

    The run boundary supplies two names. A third has no producer, so a recipe declaring one
    cannot execute - which is a second, independent reason `full-ddr-intake` has never run,
    on top of its two skeleton skills.

    The DD list is not lost by this. It arrives the way every other document does: in the
    folder, classified as `request_schedule` by the reader, which is also the only way it
    arrives in real life - nobody hands a run a request list separately from the box of
    documents it came in.
    """
    doc = recipe(RECIPE)
    assert {i["name"] for i in doc["inputs"]} == {"source_documents"}
    assert "request_schedule" not in doc["intake"]["required_types"]


def test_the_earlier_recipe_declares_an_input_nothing_produces() -> None:
    """The defect above, asserted where it actually lives, so that "full-ddr-intake cannot
    run" is a fact about the file rather than a claim in a docstring."""
    declared = {i["name"] for i in recipe("full-ddr-intake")["inputs"]}
    assert "request_schedule" in declared
    assert "engagement_profile" not in declared


def test_nothing_outside_version_one_crept_into_it() -> None:
    """EBITDA normalization is a judgement about add-backs; the other three are the model's
    own work. CLAUDE.md puts all of them outside version 1, and a recipe is the easiest
    place for one to arrive unnoticed."""
    named = set(skills_of(recipe(RECIPE)))
    for outside in ("ebitda-normalization", "qc-tieout", "gpc-scrub", "cost-of-capital",
                    "dcf-build"):
        assert outside not in named
