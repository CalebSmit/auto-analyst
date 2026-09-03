"""Fixtures for P17: files built in the test, a run to read them into, and the store.

The PDF builder writes a minimal, valid PDF by hand so a test can put text on a page at
any font size - which is how "instructions in 1pt white" is built rather than described.
"""

from __future__ import annotations

import hashlib
import io
from collections.abc import Iterator
from pathlib import Path
from typing import Any

import pytest

from guards.untrusted import UntrustedBoundary
from intake.store import IntakeRow
from ops.events import EventBus
from ops.models import RunState
from ops.run_store import SqlRunStore, open_store
from ops.source_parse import SourceParseStore
from registries.contracts import ContractRegistry

ROOT = Path(__file__).resolve().parents[2]
ENGAGEMENT = "HEND-2026-04"
RUN_ID = "run_p17"


@pytest.fixture(scope="session")
def contracts() -> ContractRegistry:
    return ContractRegistry.load()


@pytest.fixture
def store(tmp_path: Path) -> Iterator[SqlRunStore]:
    opened = open_store(f"sqlite:///{(tmp_path / 'p17.db').as_posix()}")
    yield opened
    opened.close()


@pytest.fixture
def bus(store: SqlRunStore, contracts: ContractRegistry) -> EventBus:
    return EventBus(store, contracts=contracts)


@pytest.fixture
def started(store: SqlRunStore) -> str:
    store.start_run(
        run_id=RUN_ID,
        engagement_id=ENGAGEMENT,
        plan_hash="sha256:" + "7c" * 32,
        versions={
            "model_pin": "claude-sonnet-5 @ effort=high",
            "skills": ["a@1.0.0"],
            "contracts": "1.0.0",
            "constraints": "1.0.0",
            "taxonomy": "coa@1.0.0",
        },
        tieout_tolerance=1.0,
    )
    store.transition(RUN_ID, RunState.RUNNING)
    return RUN_ID


@pytest.fixture
def parse_store(store: SqlRunStore) -> SourceParseStore:
    return SourceParseStore(store)


@pytest.fixture
def boundary(bus: EventBus) -> UntrustedBoundary:
    return UntrustedBoundary(bus)


# ---------------------------------------------------------------- builders


def build_pdf(lines: list[tuple[str, float]], *, empty: bool = False) -> bytes:
    """A one-page PDF. `lines` are (text, font size); a size under 4pt is what a person
    cannot read. `empty=True` makes a page with no text layer - a scan."""

    def esc(s: str) -> str:
        return s.replace("\\", "\\\\").replace("(", "\\(").replace(")", "\\)")

    y = 760
    ops: list[str] = []
    for text, size in lines:
        colour = "1 1 1 rg" if size < 4 else "0 0 0 rg"
        ops.append(f"BT {colour} /F1 {size:g} Tf 72 {y} Td ({esc(text)}) Tj ET")
        y -= 16
    content = "" if empty else "\n".join(ops)
    stream = content.encode("latin-1", errors="replace")
    objects = [
        b"<< /Type /Catalog /Pages 2 0 R >>",
        b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
        b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] /Contents 4 0 R "
        b"/Resources << /Font << /F1 5 0 R >> >> >>",
        b"<< /Length " + str(len(stream)).encode() + b" >>\nstream\n" + stream + b"\nendstream",
        b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
    ]
    out = io.BytesIO()
    out.write(b"%PDF-1.4\n")
    offsets = []
    for n, body in enumerate(objects, start=1):
        offsets.append(out.tell())
        out.write(f"{n} 0 obj\n".encode() + body + b"\nendobj\n")
    xref = out.tell()
    out.write(f"xref\n0 {len(objects) + 1}\n".encode())
    out.write(b"0000000000 65535 f \n")
    for off in offsets:
        out.write(f"{off:010d} 00000 n \n".encode())
    out.write(
        f"trailer\n<< /Size {len(objects) + 1} /Root 1 0 R >>\nstartxref\n{xref}\n%%EOF\n".encode()
    )
    return out.getvalue()


def build_xlsx(
    sheets: dict[str, list[list[Any]]],
    *,
    hidden: tuple[str, ...] = (),
    chart: str | None = None,
) -> bytes:
    from openpyxl import Workbook

    book = Workbook()
    book.remove(book.active)
    for name, rows in sheets.items():
        ws = book.create_sheet(name)
        for row in rows:
            ws.append(row)
        if name in hidden:
            ws.sheet_state = "hidden"
    if chart:
        # A real chart sheet, with a chart on it: openpyxl cannot read back an empty one.
        from openpyxl.chart import BarChart, Reference

        first = book.worksheets[0]
        bar = BarChart()
        bar.add_data(Reference(first, min_col=2, min_row=1, max_row=min(first.max_row, 5)))
        book.create_chartsheet(chart).add_chart(bar)
    out = io.BytesIO()
    book.save(out)
    return out.getvalue()


def build_docx(
    paragraphs: list[str],
    tables: list[list[list[str]]] = (),  # type: ignore[assignment]
    *,
    hidden: list[str] = (),  # type: ignore[assignment]
) -> bytes:
    """A Word file. `hidden` paragraphs are set in 1pt - text a person cannot read."""
    import docx
    from docx.shared import Pt

    document = docx.Document()
    for text in paragraphs:
        document.add_paragraph(text)
    for text in hidden:
        run = document.add_paragraph().add_run(text)
        run.font.size = Pt(1)
    for rows in tables:
        table = document.add_table(rows=len(rows), cols=len(rows[0]))
        for r, cells in enumerate(rows):
            for c, text in enumerate(cells):
                table.cell(r, c).text = text
    out = io.BytesIO()
    document.save(out)
    return out.getvalue()


TB_ROWS: list[list[Any]] = [
    ["Account", "FY2025", "FY2024"],
    ["Revenue", 4_812_300, 4_401_900],
    ["Cost of sales", 2_905_100, 2_688_400],
    ["Officer salary", 240_000, 235_000],
    ["Officer bonus", 60_000, 50_000],
    ["Auto allowance", 18_400, 18_400],
    ["Health premium", 12_600, 12_100],
    ["Total officer comp", 331_000, 315_500],
    ["Rent", 96_000, 96_000],
]


def intake_row(filename: str, data: bytes, kind: str) -> IntakeRow:
    return IntakeRow(
        engagement_id=ENGAGEMENT,
        source_id="sha256:" + hashlib.sha256(data).hexdigest(),
        filename=filename,
        path=f"Clients/Henderson Manufacturing/2026 ESOP/{filename}",
        watch_id="w-val",
        document_type=None,
        kind=kind,
        mime="application/octet-stream",
        bytes=len(data),
        source_version=1,
        received_at="2026-09-03T04:00:00.000000Z",
        state="pending",
    )
