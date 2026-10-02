"""P129: reading the firm's model, and drafting the binding it needs.

The one thing standing between this build and Chartwell's own model is a file nobody has
handed over and four hundred lines of YAML nobody has written. `BUILD/OPEN.md` has carried
it since P24; `Chartwell_Input_Help_v1@1.0.yaml` says in its own comment that when the real
template arrives *"this schema is re-derived from its Input tab"*. This is the tool that
re-derives it, and what the tests here hold it to is the difference between a useful draft
and a dangerous one.

**The draft has to be loadable once a person has read it.** The strongest test in this
file: generate a draft, strip the markers the way a reviewer would, and hand it to
`TemplateInputs.load`. A draft in nearly the right shape is worse than no draft, because
the person finding out is the one who thought the work was done.

**And not before.** The markers are not decoration. A schema that loaded straight out of
this tool would bind the firm's model by inference from the spelling of its named ranges,
and the first symptom would be a figure written into the wrong row of a workbook that
opens perfectly.

**It reads what is there, and says what it could not follow.** A name pointing at a
missing sheet, at two destinations, at a whole column - each is reported, never dropped.
A dropped name is an input nobody knows exists.
"""

from __future__ import annotations

import json
from io import BytesIO
from pathlib import Path

import pytest
import yaml
from openpyxl import Workbook, load_workbook
from openpyxl.workbook.defined_name import DefinedName

from deliver.registry import TemplateInputs, version_key
from tools.read_template import (
    NEEDS_A_PERSON,
    as_yaml,
    draft,
    inventory,
    main,
    names_in,
)

ROOT = Path(__file__).resolve().parents[2]


def a_model_template() -> bytes:
    """A workbook shaped like a firm's model: an Input tab of empty named cells, a
    calculation sheet of formulas, and the awkward cases a real file has in it."""
    wb = Workbook()
    inputs = wb.active
    assert inputs is not None
    inputs.title = "Input"
    rows = [
        ("Revenue_FY1", "B2"),
        ("Revenue_FY2", "B3"),
        ("Officer_Comp_FY1", "B4"),
        ("NWC_Target", "B5"),
    ]
    for name, ref in rows:
        wb.defined_names.add(DefinedName(name, attr_text=f"Input!${ref[0]}${ref[1:]}"))

    # A cell somebody has already typed into: not an input, because a template is empty
    # by definition and this one is not.
    inputs["B6"] = 0.21
    wb.defined_names.add(DefinedName("Tax_Rate", attr_text="Input!$B$6"))

    calc = wb.create_sheet("DCF")
    calc["B2"] = "=Input!B2-Input!B4"
    wb.defined_names.add(DefinedName("EBITDA_FY1", attr_text="DCF!$B$2"))

    # A name over a block, which is real and is not one cell.
    wb.defined_names.add(DefinedName("Historicals", attr_text="Input!$B$2:$D$4"))
    # A name local to a sheet, which is how a model's Input tab usually defines them.
    inputs.defined_names.add(DefinedName("Wages_FY1", attr_text="Input!$B$7"))

    buffer = BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


# -------------------------------------------------------------- reading what is there


def test_every_defined_name_is_resolved_to_a_sheet_a_cell_and_what_it_holds() -> None:
    names, skipped = names_in(a_model_template())
    by_name = {n.name: n for n in names}
    assert set(by_name) >= {
        "Revenue_FY1",
        "NWC_Target",
        "Tax_Rate",
        "EBITDA_FY1",
        "Historicals",
        "Wages_FY1",
    }
    assert by_name["Revenue_FY1"].sheet == "Input"
    assert by_name["Revenue_FY1"].ref == "B2"
    assert by_name["Revenue_FY1"].holds == "empty"
    assert by_name["Tax_Rate"].holds == "number"
    assert by_name["EBITDA_FY1"].holds == "formula"
    assert by_name["Historicals"].holds == "range" and by_name["Historicals"].cells == 9
    assert skipped == ()


def test_a_name_local_to_a_sheet_is_found() -> None:
    """A model's Input tab usually defines its names locally, and a reader that asked only
    the workbook would report a file with four hundred names as having none."""
    names, _ = names_in(a_model_template())
    assert "Wages_FY1" in {n.name for n in names}


def test_only_an_empty_single_cell_reads_as_an_input() -> None:
    """Empty is the signal, and it is a good one: a template is empty by definition. A
    cell holding a formula is the model's own work, and writing into it would replace the
    model rather than fill it in."""
    names, _ = names_in(a_model_template())
    inputs = {n.name for n in names if n.looks_like_an_input}
    assert inputs == {"Revenue_FY1", "Revenue_FY2", "Officer_Comp_FY1", "NWC_Target", "Wages_FY1"}
    assert "EBITDA_FY1" not in inputs, "a formula cell is the model's work"
    assert "Tax_Rate" not in inputs, "a cell with a number in it is already filled in"
    assert "Historicals" not in inputs, "a block is not one destination"


def test_the_stamps_this_build_writes_itself_are_never_inputs() -> None:
    wb = load_workbook(BytesIO(a_model_template()))
    wb.defined_names.add(DefinedName("Template_Version", attr_text="Input!$A$1"))
    wb.defined_names.add(DefinedName("Run_Id", attr_text="Input!$A$2"))
    buffer = BytesIO()
    wb.save(buffer)
    names, _ = names_in(buffer.getvalue())
    reserved = {n.name: n for n in names if n.name in ("Template_Version", "Run_Id")}
    assert len(reserved) == 2
    assert not any(n.looks_like_an_input for n in reserved.values()), (
        "a mapping may never name a stamp; proposing one as an input would invite it to"
    )


def test_a_sheet_filter_narrows_it_and_says_so_when_it_matches_nothing() -> None:
    names, skipped = names_in(a_model_template(), sheet="Input")
    assert {n.sheet for n in names} == {"Input"}
    assert "EBITDA_FY1" not in {n.name for n in names}

    _, missed = names_in(a_model_template(), sheet="Nowhere")
    assert any("no defined name points at sheet 'Nowhere'" in note for note in missed)
    assert any("DCF" in note for note in missed), "and it names the sheets there are"


def test_a_name_this_cannot_follow_is_reported_rather_than_dropped() -> None:
    """A dropped name is an input nobody knows exists."""
    wb = load_workbook(BytesIO(a_model_template()))
    wb.defined_names.add(DefinedName("Whole_Column", attr_text="Input!$B:$B"))
    wb.defined_names.add(DefinedName("Gone", attr_text="Deleted!$B$2"))
    buffer = BytesIO()
    wb.save(buffer)

    names, skipped = names_in(buffer.getvalue())
    assert "Whole_Column" not in {n.name for n in names}
    assert "Gone" not in {n.name for n in names}
    assert any("Whole_Column" in note for note in skipped)
    assert any("Gone" in note and "has not got" in note for note in skipped)


def test_the_inventory_is_data_something_else_can_read() -> None:
    """The roll-forward's Excel half needs this as much as the schema does, and CLAUDE.md
    says it is the thing that is not knowable from a photograph."""
    found = inventory(a_model_template(), sheet="Input")
    assert found["sheets"] == ["Input"]
    assert found["inputs"] == 5
    rows = {r["name"]: r for r in found["names"]}
    assert rows["Revenue_FY1"]["input"] is True
    assert rows["Tax_Rate"]["input"] is False
    assert json.dumps(found), "and it serialises"


# ------------------------------------------------------- the draft, and its markers


def test_every_proposed_line_is_marked() -> None:
    drafted = draft(
        a_model_template(), file="Model.xlsx", template_version="1.0", sheet="Input"
    )
    assert drafted.proposed
    for one in drafted.inputs:
        assert one["line"].startswith(f"{NEEDS_A_PERSON}: "), one
    assert len(drafted.proposed) == len(drafted.inputs)


def test_a_name_with_no_period_in_it_says_so_rather_than_being_given_one() -> None:
    drafted = draft(
        a_model_template(), file="Model.xlsx", template_version="1.0", sheet="Input"
    )
    by_name = {i["name"]: i for i in drafted.inputs}
    assert by_name["Revenue_FY1"]["period"] == "FY1"
    assert by_name["Revenue_FY1"]["line"] == f"{NEEDS_A_PERSON}: Revenue"
    assert by_name["Officer_Comp_FY1"]["line"] == f"{NEEDS_A_PERSON}: Officer Comp"
    assert NEEDS_A_PERSON in by_name["NWC_Target"]["period"]
    assert "no period in the name" in by_name["NWC_Target"]["period"]


def test_the_draft_will_not_load_as_it_stands() -> None:
    """The markers are not decoration. A schema that loaded straight out of this tool
    would bind the firm's model by inference from the spelling of its named ranges."""
    drafted = draft(
        a_model_template(), file="Model.xlsx", template_version="1.0", sheet="Input"
    )
    raw = yaml.safe_load(as_yaml(drafted))
    assert all(e["line"].startswith(NEEDS_A_PERSON) for e in raw["required"])
    lines = {e["line"] for e in raw["required"]}
    assert len(lines) > 1, "and the markers are not all one string somebody could sed away"


def test_a_draft_a_person_has_read_loads_into_the_registry(tmp_path: Path) -> None:
    """The strongest test here. Generate the draft, strip the markers the way a reviewer
    would, and hand it to the registry. A draft in nearly the right shape is worse than
    no draft, because the person who finds out is the one who thought it was done."""
    drafted = draft(
        a_model_template(), file="Model.xlsx", template_version="1.0", sheet="Input"
    )
    raw = yaml.safe_load(as_yaml(drafted))
    for entry in raw["required"]:
        entry["line"] = entry["line"].removeprefix(f"{NEEDS_A_PERSON}: ")
        entry["period"] = (
            "FY1" if NEEDS_A_PERSON in entry["period"] else entry["period"]
        )

    key = version_key("Model.xlsx", "1.0")
    (tmp_path / f"{key}.yaml").write_text(yaml.safe_dump(raw), encoding="utf-8")
    schema = TemplateInputs.load(key, folder=tmp_path)

    assert schema.key == key
    assert schema.file == "Model.xlsx"
    assert set(schema.required) == set(drafted.proposed)
    assert "Template_Version" in schema.stamps and "Run_Id" in schema.stamps
    assert schema.scale == 1


def test_the_scale_is_not_guessed_at(tmp_path: Path) -> None:
    """The firm's convention is that statements go into the model in thousands. A workbook
    states a layout and never its own units, so defaulting to the convention would put the
    one assumption nobody can see underneath every figure in the file."""
    plain = draft(a_model_template(), file="M.xlsx", template_version="1.0", sheet="Input")
    assert plain.schema["scale"] == 1
    said = draft(
        a_model_template(), file="M.xlsx", template_version="1.0", sheet="Input", scale=1000
    )
    assert said.schema["scale"] == 1000
    assert "1 here because a workbook states a layout" in as_yaml(plain)


def test_the_header_names_the_file_the_schema_has_to_be_called() -> None:
    """The filename is the key a run pins. A draft saved under the wrong name is a
    template the registry reports as unregistered, which reads as the system being broken
    rather than as a file in the wrong place."""
    drafted = draft(
        a_model_template(), file="Model Template - MASTER 2026.xlsx",
        template_version="2026.1", sheet="Input",
    )
    assert "templates/inputs/Model Template - MASTER 2026@2026.1.yaml" in as_yaml(drafted)


# ------------------------------------------------------------------- the command line


def test_the_command_line_reads_a_workbook_and_writes_a_draft(
    tmp_path: Path, capsys: pytest.CaptureFixture[str]
) -> None:
    book = tmp_path / "Model.xlsx"
    book.write_bytes(a_model_template())
    out = tmp_path / "draft.yaml"
    assert main([str(book), "--sheet", "Input", "--version", "1.0", "--out", str(out)]) == 0

    said = capsys.readouterr().out
    assert "look like inputs" in said
    assert NEEDS_A_PERSON in said
    assert "templates/inputs/Model@1.0.yaml" in said
    body = yaml.safe_load(out.read_text(encoding="utf-8"))
    assert body["template"] == "Model.xlsx"
    assert body["required"]


def test_it_will_not_write_over_a_draft_somebody_has_been_through(
    tmp_path: Path, capsys: pytest.CaptureFixture[str]
) -> None:
    book = tmp_path / "Model.xlsx"
    book.write_bytes(a_model_template())
    out = tmp_path / "draft.yaml"
    out.write_text("# reviewed by a person\n", encoding="utf-8")

    assert main([str(book), "--out", str(out)]) == 2
    assert "will not write over it" in capsys.readouterr().err
    assert out.read_text(encoding="utf-8") == "# reviewed by a person\n"


def test_the_inventory_is_available_on_its_own(
    tmp_path: Path, capsys: pytest.CaptureFixture[str]
) -> None:
    book = tmp_path / "Model.xlsx"
    book.write_bytes(a_model_template())
    assert main([str(book), "--inventory"]) == 0
    body = json.loads(capsys.readouterr().out)
    assert body["names"] and "Input" in body["sheets"]


def test_a_workbook_that_does_not_open_is_refused_rather_than_read_as_empty(
    tmp_path: Path, capsys: pytest.CaptureFixture[str]
) -> None:
    """An empty answer from a file that could not be opened reads as a model with no
    named ranges, which is the one wrong conclusion available here."""
    broken = tmp_path / "Model.xlsx"
    broken.write_bytes(b"not a workbook")
    assert main([str(broken)]) == 2
    assert "could not be read" in capsys.readouterr().err
    assert main([str(tmp_path / "nothing.xlsx")]) == 2


def test_a_workbook_with_no_named_ranges_is_a_non_zero_exit(
    tmp_path: Path, capsys: pytest.CaptureFixture[str]
) -> None:
    """Nothing to draft is a finding about the file: this build binds by named range and
    never by cell address, so a model with no names is a model it cannot write into."""
    wb = Workbook()
    book = tmp_path / "Nameless.xlsx"
    wb.save(book)
    assert main([str(book)]) == 1


def test_the_real_master_template_drafts_into_something_that_loads(tmp_path: Path) -> None:
    """Against the workbook this repository actually has, which is the synthetic stand-in
    for Chartwell's own. The point is not the stand-in: it is that the whole path - read
    the file, draft, review, load - works on a workbook nobody wrote this tool against.
    """
    book = ROOT / "templates" / "Chartwell_Master_Model_v14.xlsx"
    if not book.is_file():  # pragma: no cover - the session fixture builds it
        pytest.skip("the synthetic templates have not been built in this session")

    drafted = draft(
        book.read_bytes(), file=book.name, template_version="14.2", sheet="Inputs"
    )
    assert len(drafted.proposed) >= 40, "the master template declares its inputs"

    raw = yaml.safe_load(as_yaml(drafted))
    for entry in raw["required"]:
        entry["line"] = entry["line"].removeprefix(f"{NEEDS_A_PERSON}: ")
        if NEEDS_A_PERSON in entry["period"]:
            entry["period"] = "FY1"
    key = version_key(book.name, "14.2")
    (tmp_path / f"{key}.yaml").write_text(yaml.safe_dump(raw), encoding="utf-8")

    schema = TemplateInputs.load(key, folder=tmp_path)
    assert len(schema.required) == len(drafted.proposed)
    # And every name it proposed is a name the workbook really has, which is the property
    # that makes the draft worth reviewing rather than worth rewriting.
    real = set(load_workbook(BytesIO(book.read_bytes())).defined_names)
    for ws in load_workbook(BytesIO(book.read_bytes())).worksheets:
        real |= set(ws.defined_names)
    assert set(schema.required) <= real
