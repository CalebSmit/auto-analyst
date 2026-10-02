"""Workbooks to test the formula audit against, built rather than committed.

The same reasoning as `tests/docs.py`: the firm's real model is Chartwell's and a client's
workbook could not live here in any case - `fixtures/` is masked, synthetic or public,
never a document somebody sent. So the workbooks these tests run against are generated,
and they are generated to contain the mistakes that actually happen in a model:

* a **total whose range stops one row short** of the block above it, which is the defect
  the analyst named and the reason this module exists;
* a **total whose range swallows another total**, so a block is counted twice;
* a **typed number in a row of formulas**, which is the plug that stops moving when the
  model does;
* a **formula that breaks its row's pattern**, which is the copy-paste that landed one
  row off;
* a **cached total that does not equal the cells it adds**, which is the file somebody
  edited and did not recalculate - and which needs cached values, so it needs
  `with_cached_values` below;
* a **clean block**, because a check that cannot be silent is not a check. Most of any
  real workbook is correct and every correct cell is a chance for a false positive.

Anything these do not contain is something the tests do not prove.

**Why `with_cached_values` exists.** openpyxl writes a formula cell as `<f>..</f><v></v>`:
the formula, and an empty slot where Excel would have put the last computed value. So a
workbook openpyxl wrote has no cached values at all, `data_only=True` reads `None`
everywhere, and the footing check has nothing to compare - correctly, which is why the
audit reports it as skipped rather than passed. Testing the check therefore means putting
values in that slot, which is a small edit to the sheet XML inside the zip. Deliberately
not done with a library: the point is to produce the one thing openpyxl cannot, and a
helper that went through openpyxl again would produce the same empty slots.
"""

from __future__ import annotations

import re
import zipfile
from collections.abc import Mapping
from io import BytesIO

from openpyxl import Workbook

__all__ = ["a_model", "a_tidy_model", "with_cached_values"]


def _bytes(wb: Workbook) -> bytes:
    buffer = BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


def a_tidy_model() -> bytes:
    """A small workbook with nothing wrong with it.

    Three year columns, a gross profit line, an expense block and a total that covers it,
    an operating income line - all consistent, all footing. The audit must be silent on
    this one, and the reason it is here is that silence is the half of a check that is
    easy to lose.
    """
    wb = Workbook()
    ws = wb.active
    assert ws is not None
    ws.title = "IS"
    for col, label in (("B", "FY1"), ("C", "FY2"), ("D", "FY3")):
        ws[f"{col}1"] = label

    ws["A2"], ws["B2"], ws["C2"], ws["D2"] = "Revenue", 1000, 1100, 1210
    ws["A3"], ws["B3"], ws["C3"], ws["D3"] = "Cost of sales", 400, 430, 470
    ws["A4"] = "Gross profit"
    for col in "BCD":
        ws[f"{col}4"] = f"={col}2-{col}3"

    ws["A6"], ws["B6"], ws["C6"], ws["D6"] = "Salaries", 100, 110, 120
    ws["A7"], ws["B7"], ws["C7"], ws["D7"] = "Rent", 50, 52, 54
    ws["A8"], ws["B8"], ws["C8"], ws["D8"] = "Other operating", 25, 26, 27
    ws["A9"] = "Total operating expense"
    for col in "BCD":
        ws[f"{col}9"] = f"=SUM({col}6:{col}8)"

    ws["A10"] = "Operating income"
    for col in "BCD":
        ws[f"{col}10"] = f"={col}4-{col}9"
    return _bytes(wb)


def a_model() -> bytes:
    """One workbook holding one of each defect, each in its own column or row.

    Laid out so that a finding names a cell and the cell names the defect:

    ======  ===========================================================
    `B9`    a total whose range stops a row short of the block
    `D9`    a total whose range reaches up across the blank row into the
            gross profit line, and starts one row below a number
    `D10`   a typed number in a row that otherwise calculates
    `C15`   a formula one row off the pattern its neighbours share
    `B21`   a total whose range includes the subtotal inside it
    ======  ===========================================================

    The three totals on row 9 also disagree with each other - three different ranges
    across three year columns - which is `run_has_no_agreement`.

    `C9` is the cell to give a wrong saved value to with `with_cached_values`: its range
    is right, so a footing failure there can only be the file being stale, which is the
    thing that check is for.
    """
    wb = Workbook()
    ws = wb.active
    assert ws is not None
    ws.title = "IS"
    for col, label in (("B", "FY1"), ("C", "FY2"), ("D", "FY3")):
        ws[f"{col}1"] = label

    ws["A2"], ws["B2"], ws["C2"], ws["D2"] = "Revenue", 1000, 1100, 1210
    ws["A3"], ws["B3"], ws["C3"], ws["D3"] = "Cost of sales", 400, 430, 470
    ws["A4"] = "Gross profit"
    for col in "BCD":
        ws[f"{col}4"] = f"={col}2-{col}3"

    ws["A6"], ws["B6"], ws["C6"], ws["D6"] = "Salaries", 100, 110, 120
    ws["A7"], ws["B7"], ws["C7"], ws["D7"] = "Rent", 50, 52, 54
    ws["A8"], ws["B8"], ws["C8"], ws["D8"] = "Other operating", 25, 26, 27
    ws["A9"] = "Total operating expense"
    ws["B9"] = "=SUM(B6:B7)"  # short by a row: B8 sits between the range and the total
    ws["C9"] = "=SUM(C6:C8)"
    ws["D9"] = "=SUM(D4:D8)"  # reaches across the blank row into the gross profit line

    ws["A10"] = "Operating income"
    ws["B10"] = "=B4-B9"
    ws["C10"] = "=C4-C9"
    ws["D10"] = 999  # the plug

    # A margin block, four columns wide, where one cell divides by the wrong row.
    ws["A13"], ws["B13"], ws["C13"], ws["D13"], ws["E13"] = "Revenue", 1000, 1100, 1210, 1330
    ws["A14"], ws["B14"], ws["C14"], ws["D14"], ws["E14"] = "Gross profit", 600, 670, 740, 810
    ws["A15"] = "Gross margin"
    ws["B15"] = "=B14/B13"
    ws["C15"] = "=C14/C14"  # one row off: the copy that landed wrong
    ws["D15"] = "=D14/D13"
    ws["E15"] = "=E14/E13"

    # A block with its own subtotal, and a grand total that counts the subtotal as well
    # as the lines beneath it.
    ws["A18"], ws["B18"] = "Current assets - cash", 200
    ws["A19"], ws["B19"] = "Current assets - receivables", 300
    ws["A20"], ws["B20"] = "Total current assets", "=SUM(B18:B19)"
    ws["A21"], ws["B21"] = "Total assets", "=SUM(B18:B20)"
    return _bytes(wb)


def with_cached_values(data: bytes, cached: Mapping[str, Mapping[str, float]]) -> bytes:
    """Return `data` with Excel's saved-value slot filled for the named formula cells.

    `cached` is keyed by sheet name, then by cell. A cell that is not a formula cell, or
    a sheet that is not in the workbook, is a mistake in the test and raises - a helper
    that quietly did nothing would make a test pass by checking nothing, which is this
    repository's most repeated defect in its own suite.
    """
    source = zipfile.ZipFile(BytesIO(data))
    order = {name: i for i, name in enumerate(_sheet_names(source))}
    unknown = set(cached) - set(order)
    if unknown:
        raise AssertionError(f"no such sheet(s) in this workbook: {sorted(unknown)}")

    out = BytesIO()
    with zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as target:
        for item in source.infolist():
            body = source.read(item.filename)
            match = re.fullmatch(r"xl/worksheets/sheet(\d+)\.xml", item.filename)
            if match:
                index = int(match.group(1)) - 1
                sheet = next((n for n, i in order.items() if i == index), "")
                if sheet in cached:
                    body = _fill(body, cached[sheet], sheet)
            target.writestr(item, body)
    return out.getvalue()


def _sheet_names(source: zipfile.ZipFile) -> list[str]:
    """Sheet names in workbook order.

    `name` is not the first attribute on the tag - openpyxl writes an `xmlns:r` before it
    - and a pattern that assumed it was matched nothing and reported every sheet as
    missing. The position of `sheetN.xml` in this list is its index, which holds for a
    workbook openpyxl wrote and is the only kind this helper is given.
    """
    book = source.read("xl/workbook.xml").decode("utf-8")
    return re.findall(r'<sheet\b[^>]*?\sname="([^"]+)"', book)


def _fill(body: bytes, values: Mapping[str, float], sheet: str) -> bytes:
    text = body.decode("utf-8")
    for cell, value in values.items():
        pattern = re.compile(
            rf'(<c r="{re.escape(cell)}"[^>]*>\s*<f>[^<]*</f>\s*)<v>\s*</v>'
        )
        text, count = pattern.subn(rf"\g<1><v>{value}</v>", text)
        if count != 1:
            raise AssertionError(
                f"{sheet}!{cell} is not a formula cell with an empty value slot "
                f"(matched {count} times). Either the cell holds no formula or "
                f"openpyxl has changed how it writes one."
            )
    return text.encode("utf-8")

