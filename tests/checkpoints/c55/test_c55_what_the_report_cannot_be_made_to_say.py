"""C55 - What the report cannot be made to say.

P91 and P92 put the report on the dashboard: find a named place, move a body of Word
content into it, and do that from a button, one appendix at a time. Each phase passed its
own tests. What a checkpoint asks is whether the two halves still agree about the things
that span them.

Three of those, and the third is the one that matters most:

1. **Binding is by name, everywhere.** Not in `anchors.py` and by index in `setup.py`.
2. **An update is one appendix.** The registry says which; the update honours it; the page
   shows exactly those.
3. **Nothing can write the certification.** Not because a guard refuses, but because no
   path exists - and that is a property of the registry, the update and the page together,
   which is why it belongs here rather than in either phase.
"""

from __future__ import annotations

import ast
from pathlib import Path

import pytest
from docx import Document

from registries.errors import RegistryError
from report import APPENDICES, preview, set_up, update, what_would_change
from report.anchors import find_anchor
from report.appendices import updatable
from tests.docs import a_report, appendix_source, text_of

ROOT = Path(__file__).resolve().parents[3]


@pytest.fixture
def desk(tmp_path: Path) -> dict[str, Path]:
    last_year = tmp_path / "Valuation 2025.docx"
    a_report().save(str(last_year))
    source = tmp_path / "appendix-data.docx"
    appendix_source().save(str(source))
    return {
        "last_year": last_year,
        "this_year": set_up(last_year, tmp_path / "Valuation 2026.docx"),
        "source": source,
    }


# ============================================================ by name, everywhere


def test_no_function_in_the_package_takes_a_block_index() -> None:
    """`B44` is not a contract, and neither is "the third paragraph". A report is edited in
    prose: somebody adds a sentence and every index below it moves by one, silently."""
    for module in ("anchors.py", "appendices.py", "setup.py", "transplant.py"):
        tree = ast.parse((ROOT / "report" / module).read_text(encoding="utf-8"))
        for node in ast.walk(tree):
            if not isinstance(node, ast.FunctionDef) or node.name.startswith("_"):
                continue
            for argument in node.args.args + node.args.kwonlyargs:
                if argument.arg in ("index", "position", "paragraph_number", "n"):
                    raise AssertionError(
                        f"{module}:{node.lineno} {node.name} takes {argument.arg!r}"
                    )


def test_every_appendix_in_the_registry_binds_to_a_name() -> None:
    for entry in APPENDICES:
        assert entry.anchor
        assert not entry.anchor.isdigit(), entry.id


def test_a_report_that_has_been_edited_still_binds_correctly(desk: dict[str, Path]) -> None:
    """The whole argument for binding by name, run end to end: insert a paragraph above
    the appendices - the ordinary act of editing a report - and the update still lands in
    the right place. An index-based binding would be off by one from here on."""
    document = Document(str(desk["this_year"]))
    body = document.element.body
    first = next(b for b in body.iterchildren() if isinstance(b.tag, str))
    # Built in this document rather than borrowed from another: an element belongs to one
    # tree, and lxml refuses to have it in two.
    added = document.add_paragraph("A sentence somebody added to the introduction.")._p
    first.addnext(added)
    document.save(str(desk["this_year"]))

    change = update(desk["this_year"], "economic", desk["source"])
    after = text_of(Document(change.written_to))
    assert "A sentence somebody added to the introduction." in after
    assert "Economic conditions in the period under review." in after
    # The report's own heading stays - it carries the style and the bookmark - and the
    # new content goes under it, before the next appendix.
    assert after.index("Appendix C") < after.index("Economic conditions in the period")
    assert after.index("Economic conditions in the period") < after.index("Appendix E")
    assert "Last year's economic conditions, still in place." not in after


def test_a_missing_anchor_refuses_in_both_halves(tmp_path: Path) -> None:
    """The panel reports it, the update refuses it, and neither guesses."""
    thin = tmp_path / "thin.docx"
    Document().save(str(thin))
    source = tmp_path / "s.docx"
    appendix_source().save(str(source))

    rows = {r["appendix"]: r for r in what_would_change(thin)}
    assert not any(r["found"] for r in rows.values())

    with pytest.raises(RegistryError) as refused:
        update(thin, "economic", source)
    assert refused.value.diagnostic.code == "CONFIG.BINDING_MISSING"


# ============================================================ one appendix


def test_the_page_offers_exactly_what_the_registry_says_can_change(
    desk: dict[str, Path],
) -> None:
    """Two lists of appendices would drift, and the drift looks like a button that does
    nothing."""
    offered = {row["appendix"] for row in what_would_change(desk["this_year"])}
    assert offered == {a.id for a in updatable()}


def test_updating_each_of_the_three_leaves_the_other_two_alone(
    desk: dict[str, Path],
) -> None:
    was_there = {
        "economic": "Last year's economic conditions, still in place.",
        "indices": "Last year's summary of major market indices, still in place.",
        "ma": "Last year's summary of m&a markets, still in place.",
    }
    for which in ("economic", "indices", "ma"):
        change = update(desk["this_year"], which, desk["source"])
        after = text_of(Document(change.written_to))
        assert was_there[which] not in after, which
        for other, sentence in was_there.items():
            if other != which:
                assert sentence in after, (which, other)


def test_an_appendix_can_be_updated_again_next_quarter(desk: dict[str, Path]) -> None:
    """The defect this checkpoint found. The update used to remove the anchor's own block,
    so the first press deleted the heading it had bound to and the appendix became
    unreachable - which a unit test on one press could not see, and which would have
    surfaced at work as "the button worked last time"."""
    at = desk["this_year"]
    for _ in range(3):
        at = Path(update(at, "economic", desk["source"]).written_to)
    after = text_of(Document(at))
    assert "Appendix C - Economic Conditions" in after
    assert after.count("Economic conditions in the period under review.") == 1


def test_a_preview_and_the_update_it_previews_agree(desk: dict[str, Path]) -> None:
    for which in ("economic", "indices", "ma"):
        shown = preview(desk["this_year"], which, desk["source"]).as_dict()
        done = update(desk["this_year"], which, desk["source"]).as_dict()
        for field in ("replaced_blocks", "arriving_blocks", "styles", "left_behind"):
            assert shown[field] == done[field], (which, field)


# ============================================================ the certification


def test_there_is_no_path_from_any_button_to_the_certification(
    desk: dict[str, Path],
) -> None:
    """Held across all three surfaces, because "no path exists" is not a property of any
    one of them."""
    # The registry does not offer it.
    assert "certification" not in {a.id for a in updatable()}
    # The panel does not list it.
    assert "certification" not in {r["appendix"] for r in what_would_change(desk["this_year"])}
    # The update refuses it if somebody constructs the call by hand.
    with pytest.raises(RegistryError):
        update(desk["this_year"], "certification", desk["source"])
    # And the page has no button for it.
    from admin.page import PAGE

    assert "certification" in PAGE, "it is explained"
    assert 'data-appendix="certification"' not in PAGE, "but not offered"


def test_the_certification_is_still_there_after_every_update(desk: dict[str, Path]) -> None:
    """Carried forward in the copy, which is the whole mechanism: it survives because
    nothing touches it, not because something protects it."""
    at = desk["this_year"]
    for which in ("economic", "indices", "ma"):
        at = Path(update(at, which, desk["source"]).written_to)
    after = text_of(Document(at))
    assert "Appendix H - Certification" in after
    assert "Last year's certification, still in place." in after


def test_the_blueprint_s_line_is_kept_by_construction_here() -> None:
    """`Nothing concludes, signs or certifies` - and in this package that is true because
    there is no code, rather than because a check refuses. The registry is where it is
    decided, so that is where the test looks."""
    from report.appendices import appendix

    certification = appendix("certification")
    assert certification.letter == "H"
    assert certification.updated is False
    assert "appraiser signs" in certification.note


# ============================================================ the move, once more


def test_a_document_that_has_been_through_the_whole_flow_is_still_valid(
    desk: dict[str, Path],
) -> None:
    """Three updates in sequence, then open it. Each phase proved one move; this proves
    they compose, which is where a partname collision or a numbering clash would show."""
    import zipfile

    at = desk["this_year"]
    for which in ("economic", "indices", "ma"):
        at = Path(update(at, which, desk["source"]).written_to)

    with zipfile.ZipFile(at) as package:
        names = package.namelist()
    assert len(names) == len(set(names)), "a duplicate zip entry is a corrupt .docx"

    again = Document(str(at))
    R = "{http://schemas.openxmlformats.org/officeDocument/2006/relationships}"
    referenced = {
        str(value)
        for node in again.element.body.iter()
        for key, value in node.attrib.items()
        if key.startswith(R) and str(value).startswith("rId")
    }
    assert referenced <= set(again.part.rels), sorted(referenced - set(again.part.rels))


def test_three_lists_from_three_updates_are_three_lists(desk: dict[str, Path]) -> None:
    """The failure the remapping exists for, at the scale it would really happen: three
    appendices each bringing a numbered list from the same source document. Without the
    remap they would share a numId and Appendix F would start at item 7."""
    from tests.docs import list_ids_in

    at = desk["this_year"]
    for which in ("economic", "indices", "ma"):
        at = Path(update(at, which, desk["source"]).written_to)

    ids = list_ids_in(Document(str(at)))
    arriving = [i for i in ids if ids.count(i) == 3]
    assert len(set(arriving)) == 3, f"three separate lists expected, got {sorted(set(ids))}"


def test_the_report_package_is_reached_and_no_longer_excused() -> None:
    """P91 built it and nothing called it; P92 joined it. The allowlist entries that
    covered the gap are gone, which is what resolved looks like."""
    import sys

    sys.path.insert(0, str(ROOT))
    from tests.p82.test_nothing_a_run_cannot_reach import ALLOWED, unreachable

    assert not [m for m in unreachable() if m.startswith("report")]
    assert not [m for m in ALLOWED if m.startswith("report")]


def test_the_anchor_kinds_the_registry_relies_on_both_work(desk: dict[str, Path]) -> None:
    """Headings today, bookmarks when somebody marks one. Both have to resolve, because a
    firm that bookmarks its appendices and a firm that does not are the same firm on
    different years of the same document."""
    from tests.docs import bookmark

    document = Document(str(desk["this_year"]))
    heading = next(
        p for p in document.paragraphs if p.text == "Appendix F - Summary of M&A Markets"
    )
    bookmark(heading, "Appendix F - Summary of M&A Markets", number=70)
    document.save(str(desk["this_year"]))

    assert find_anchor(Document(str(desk["this_year"])), APPENDICES[2].anchor).kind == "bookmark"
    assert find_anchor(Document(str(desk["this_year"])), APPENDICES[0].anchor).kind == "heading"
    change = update(desk["this_year"], "ma", desk["source"])
    assert "Economic conditions in the period under review." in text_of(
        Document(change.written_to)
    )
