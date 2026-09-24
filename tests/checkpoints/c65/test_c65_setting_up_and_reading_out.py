"""C65 - the two ends of the analyst's own test.

Both phases came from doing it rather than from reading the code.

* **P114** is the folder I filed one level too shallow. The planner skipped every file,
  nothing was swept, nothing was unclaimed, and preflight reported that every swept folder
  mapped to an engagement - true, and useless. A green Setup page over thirteen client
  files and silence.
* **P115** is the column I found by opening the workbook the run finally produced. It had a
  header, forty-four empty cells, and a page in front of it promising that each one said
  which file and which row.

They sit at opposite ends of one run and they are the same failure twice: **the system
knew something and did not say it.** The depth was in the watch; the locator was in the
figure envelope. Neither reached the person.

So the question a checkpoint asks here is whether the two ends agree about one engagement:
that a folder which cannot be swept is refused before anything is filed, and that a folder
which can produces a workbook where every figure says where it came from.
"""

from __future__ import annotations

import hashlib
from pathlib import Path
from typing import Any

import pytest
from openpyxl import load_workbook

from deliver.builder import FROM_THE_PROFILE, NOT_FOUND, ModelBuilder, waiver_for
from deliver.registry import SOURCE_SUFFIX, RegisteredTemplate, TemplateInputs
from intake.config import Watch
from intake.sweep import SweepPlanner, folder_too_shallow
from tools.write_template import read_named

ROOT = Path(__file__).resolve().parents[3]
HELP = ROOT / "templates" / "Chartwell_Input_Help_v1.xlsx"

CLIENT = "Kestrel Precision Works"
ENG = "KSPW-2026-01"
RIGHT = "Kestrel Precision/2026 Valuation"
WRONG = "Kestrel Precision"
DOCS = {"sha256:" + "a" * 64: "Consolidated Income Statements_Redacted.xlsx"}


def watch(depth: int = 2) -> Watch:
    return Watch(
        id="w-clients", kind="pattern", name="Clients", root="Clients", recipe="r", depth=depth
    )


def plan_over(folder: str, *paths: str) -> Any:
    from intake.roster import RosterEntry, RosterState

    entry = RosterEntry(
        engagement_id=ENG,
        client=CLIENT,
        scope_path="valuation/non-esop/non-repeat",
        watch_id="w-clients",
        destination_id="d-out",
        folder=folder,
        state=RosterState.ACTIVE,
    )
    roster = type("R", (), {"entries": lambda self: (entry,), "get": lambda self, i: entry})()
    config = type("C", (), {"watches": lambda self: (watch(),)})()
    files = type(
        "F", (), {"list": lambda self, p: [type("Ref", (), {"path": x})() for x in paths]}
    )()
    return SweepPlanner(roster, config, files).plan()


# ============================================================ the folder end


def test_the_folder_i_typed_wrong_is_refused_before_anything_is_filed() -> None:
    """Provable from two numbers on one form, which is why it belongs at registration
    rather than in a report somebody reads after the silence."""
    assert folder_too_shallow(WRONG, watch()) is not None
    assert folder_too_shallow(RIGHT, watch()) is None


def test_and_the_files_i_dropped_there_are_counted() -> None:
    """The second half. The folder can be right and the files still land above it, and
    that one is a person's judgement rather than a refusal."""
    wrong = plan_over(WRONG, f"Clients/{WRONG}/Consolidated Income Statements.xlsx")
    assert wrong.targets == () and len(wrong.shallow) == 1

    right = plan_over(RIGHT, f"Clients/{RIGHT}/Consolidated Income Statements.xlsx")
    assert right.shallow == () and [t.engagement_id for t in right.targets] == [ENG]


# ============================================================ the workbook end


def input_help() -> RegisteredTemplate:
    schema = TemplateInputs.load("Chartwell_Input_Help_v1@1.0")
    data = HELP.read_bytes()
    return RegisteredTemplate(
        file=schema.file,
        template_version=schema.template_version,
        content_hash=f"sha256:{hashlib.sha256(data).hexdigest()}",
        named_ranges=frozenset(load_workbook(HELP).defined_names),
        inputs=schema,
        path=HELP,
    )


def envelope(value: float, locator: str = "IS!B7") -> dict[str, Any]:
    return {
        "value": value,
        "unit": "currency",
        "currency": "USD",
        "scale": 1000,
        "period": {"type": "FY", "end": "2023-12-31"},
        "source_id": "sha256:" + "a" * 64,
        "locator": locator,
        "source_version": 1,
        "derivation": None,
        "masked": True,
        "confidence": 0.9,
        "produced_by": "timing-normalization@0.1.0",
        "model_pin": "claude-sonnet-5 @ effort=high",
    }


def delivered(waived: tuple[str, ...] = (), *, benchmark: bool = False) -> Any:
    template = input_help()
    bound = [n for n in template.inputs.required if n not in waived]
    figures: dict[str, Any] = {
        f"timing:{n}": envelope(float(i), f"IS!B{i + 3}") for i, n in enumerate(bound, start=1)
    }
    if benchmark and bound:
        figures[f"timing:{bound[-1]}"] = 265.0  # the profile's own number, not a document
    mapping = {
        "template_mapping": {
            "template_version": template.template_version,
            "bindings": [{"named_range": n, "figure_ref": f"timing:{n}"} for n in bound],
            "unmapped_template_inputs": [
                {"named_range": n, "why": "nothing in the folder answered it"} for n in waived
            ],
        }
    }
    return ModelBuilder().build(
        mapping,
        figures,
        template,
        run_id="run_c65",
        engagement_id=ENG,
        step_id="release",
        waived=(waiver_for("release", waived),) if waived else (),
        documents=DOCS,
    )


def test_every_figure_in_the_workbook_says_where_it_came_from() -> None:
    built = delivered()
    schema = input_help().inputs
    back = read_named(built.data, list(schema.sources))
    for name in schema.required:
        assert DOCS["sha256:" + "a" * 64] in str(back[f"{name}{SOURCE_SUFFIX}"])


def test_a_figure_that_came_from_the_profile_says_that_instead() -> None:
    built = delivered(benchmark=True)
    said = read_named(built.data, list(input_help().inputs.sources))
    assert FROM_THE_PROFILE in {str(v) for v in said.values()}


# ============================================================ where the two meet


def test_a_waived_input_says_not_found_and_claims_no_source() -> None:
    """Where P112, P113 and P115 have to agree about one row. The Value cell says the
    figure was not found, the Source cell claims nothing, and the front page explains both
    - three statements that have to be the same statement."""
    name = input_help().inputs.required[0]
    built = delivered(waived=(name,))
    assert built.writes[name] == NOT_FOUND
    assert not read_named(built.data, [f"{name}{SOURCE_SUFFIX}"])[f"{name}{SOURCE_SUFFIX}"]

    sheet = load_workbook(HELP).worksheets[0]
    page = " ".join(str(c.value) for row in sheet.iter_rows() for c in row if c.value)
    assert NOT_FOUND.upper() in page.upper()
    assert "Source column" in page


def test_the_page_promises_only_what_the_workbook_does() -> None:
    """The claim that started P113 and the claim that started P115, both now true of the
    file they are printed in."""
    sheet = load_workbook(HELP).worksheets[0]
    page = " ".join(str(c.value) for row in sheet.iter_rows() for c in row if c.value)
    assert "which file and which row" in page
    built = delivered()
    back = read_named(built.data, list(input_help().inputs.sources))
    said = [str(v) for v in back.values() if v]
    assert said and all("!" in text for text in said), "a file and a row, not just a file"


def test_the_numbers_still_lock(world: Any) -> None:
    """Neither phase may cost the thing the whole delivery rests on."""
    built = delivered()
    assert built.lock.passed
    assert read_named(built.data, list(built.writes)) == built.writes
    del world  # the P21 world is imported so this suite runs under the same fixtures


@pytest.mark.parametrize("folder", [WRONG, "", "a/b/c"])
def test_no_folder_shape_makes_the_planner_guess(folder: str) -> None:
    """The rule both halves share: report, never repair. Guessing that
    `Clients/Kestrel Precision` meant `Clients/Kestrel Precision/2026` would be the system
    inventing an engagement folder, which is the same class of mistake as inventing a
    figure."""
    plan = plan_over(folder, f"Clients/{folder}/IS.xlsx" if folder else "Clients/IS.xlsx")
    assert all(t.engagement_id == ENG for t in plan.targets)
    for target in plan.targets:
        assert target.location.endswith(folder) or folder == ""
