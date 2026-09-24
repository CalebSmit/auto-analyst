"""C61 - what a person can set up without being told how.

P101 and P102 answer one question between them: can the analyst point this at a folder and
press go. Each was checkable alone. What neither could ask is whether the *setup* holds
together - and setting a version-1 run going turned out to need four things in order, three
of which were found by watching a run quarantine rather than by reading anything:

1. a vault key, or nothing masks and the pool never starts;
2. a template pin on the engagement, or `model` is refused by the policy gate;
3. that template registered, because a run does not write into a workbook nobody validated;
4. the binding in the **config document**, which `initialise` seeds from
   `templates/bindings.yaml` - editing the file alone changes nothing already running.

Every one of those is a correct guard. Together they are a sequence nobody could infer, so
what this checkpoint holds is that each refusal *names the thing to fix*: a guard that stops
a run without saying which screen to open is a guard that gets worked around.

And one rule across both phases: **nothing spends money unless a person pressed something.**
P101 added the first step that costs anything; P102 added the only button that does.
"""

from __future__ import annotations

import ast
from pathlib import Path

import pytest
import yaml

from admin.connection import probe_model
from graph.recipes import load_recipe, recipe_names
from graph.resolve import ScopeResolver
from harness.procedures import PROCEDURES

ROOT = Path(__file__).resolve().parents[3]
FOCUS = "valuation/non-esop/non-repeat"


class Secrets:
    def has(self, name: str) -> bool:
        return False


class Stand_in:
    name = "fake"

    def complete(self, request: object) -> object:  # pragma: no cover - never called
        raise AssertionError("the probe called a stand-in adapter")


# ================================================ the setup refuses in words, not codes


@pytest.mark.parametrize(
    ("what", "must_say"),
    [
        ("vault", "AA_VAULT_KEY"),
        ("template", "Template Registry"),
        ("pin", "pinned no version"),
    ],
)
def test_each_refusal_names_the_thing_to_fix(what: str, must_say: str) -> None:
    """The four-step setup is only discoverable because each guard says what is missing.

    Held on the text because that is the whole mechanism: `PRIVACY.VAULT_UNAVAILABLE`
    tells somebody nothing, and "AA_VAULT_KEY is not set - generate one with..." tells them
    everything. A diagnostic that named only its code would leave the sequence to be
    rediscovered by whoever sets this up next.
    """
    # Each is asserted where the words are written rather than where the failure is
    # raised. `registries/errors.py` carries the "what to do" for every code, which is the
    # half a person reads, and the first version of this test looked for it in the module
    # that constructs the diagnostic - where only the code lives.
    sources = {
        "vault": [ROOT / "privacy" / "vault.py"],
        "template": [ROOT / "registries" / "errors.py", ROOT / "admin" / "preflight.py"],
        "pin": [ROOT / "harness" / "policy.py"],
    }
    assert any(
        must_say in path.read_text(encoding="utf-8") for path in sources[what]
    ), f"{must_say!r} is in none of {[p.name for p in sources[what]]}"


def test_initialise_seeds_the_bindings_rather_than_leaving_them_empty() -> None:
    """The fourth step, and the one with no error message at all: a binding added to
    `templates/bindings.yaml` changes nothing about a system that is already initialised.
    `initialise` seeding from the checked-in file is what makes the file the record."""
    source = (ROOT / "admin" / "console.py").read_text(encoding="utf-8")
    body = source[source.index('"model_bindings": []') : source.index("def add_watch")]
    assert "ScopeResolver().bindings" in body
    assert "seed_model_bindings" in body


# ================================================ the path a run takes


def test_the_focus_scope_has_a_recipe_whose_every_skill_is_built() -> None:
    """The two halves of P101's premise, asked together: version 1 has a recipe, and the
    recipe has no step that stops for being unbuilt."""
    assert "valuation-input-prep" in recipe_names()
    recipe = load_recipe("valuation-input-prep")
    named = [s.skill for s in recipe.steps if s.kind == "skill"]
    assert named
    for skill in named:
        assert skill in PROCEDURES, skill


def test_the_focus_scope_writes_something_that_is_not_the_model() -> None:
    template = ScopeResolver().resolve(FOCUS).template
    assert template.file == "Chartwell_Input_Help_v1.xlsx"
    assert "Model" not in template.file.removeprefix("Chartwell_")


def test_the_deliverable_and_the_model_template_are_different_files() -> None:
    """Stated as a property of the tree rather than of one binding, because the failure
    would be somebody registering the input-help schema over the master one and nothing
    noticing until a run wrote into the firm's model."""
    schemas = sorted((ROOT / "templates" / "inputs").glob("*.yaml"))
    by_file = {}
    for path in schemas:
        loaded = yaml.safe_load(path.read_text(encoding="utf-8"))
        by_file.setdefault(loaded["template"], []).append(loaded.get("kind"))
    assert by_file["Chartwell_Input_Help_v1.xlsx"] == ["input_help"]
    assert by_file["Chartwell_Master_Model_v14.xlsx"] == [None]


# ================================================ nothing spends without a person


def test_the_connection_test_calls_nothing_when_it_cannot_help() -> None:
    """The only button in the build that spends money, asked to prove it does not spend it
    pointlessly: a stand-in adapter and a missing key both stop it before the call."""
    report = probe_model(Stand_in(), Secrets())
    assert not report.connected
    assert [s.status for s in report.steps] == ["fail", "not reached", "not reached", "not reached"]


def test_no_run_path_can_reach_the_probe() -> None:
    """A run that could call it would pay for a self-test nobody asked for. Checked on the
    imports of every module a run passes through, rather than on the driver alone."""
    for module in (
        "workers/driver.py",
        "workers/pool.py",
        "graph/execute.py",
        "harness/dispatch.py",
        "intake/__init__.py",
    ):
        tree = ast.parse((ROOT / module).read_text(encoding="utf-8"))
        named: set[str] = set()
        for node in ast.walk(tree):
            if isinstance(node, ast.Import):
                named.update(a.name for a in node.names)
            elif isinstance(node, ast.ImportFrom) and node.module:
                named.add(node.module)
        assert "admin.connection" not in named, module


def test_the_spend_limit_is_still_capped_by_the_firm_ceiling() -> None:
    """The control the analyst asked for by name. P101 made a run able to spend; this is
    the thing that decides how much, and it may be lowered per engagement and never raised
    past the firm's own number."""
    source = (ROOT / "admin" / "console.py").read_text(encoding="utf-8")
    body = source[source.index("def set_budget") : source.index("def set_review_level")]
    assert "if usd > self.firm_ceiling:" in body
    assert "AdminRefused" in body or "raise" in body


# ================================================ what version 1 still will not do


def test_version_one_does_not_conclude_anything() -> None:
    """The blueprint's rule, asked of the phase that first produced a deliverable. The
    workbook carries inputs and no formulas; a file that computed a value would be the one
    thing version 1 is defined against."""
    from openpyxl import load_workbook

    book = load_workbook(ROOT / "templates" / "Chartwell_Input_Help_v1.xlsx")
    try:
        assert not [
            c.value
            for sheet in book.worksheets
            for row in sheet.iter_rows()
            for c in row
            if isinstance(c.value, str) and c.value.startswith("=")
        ]
    finally:
        book.close()


def test_the_recipe_leaves_out_everything_outside_version_one() -> None:
    recipe = load_recipe("valuation-input-prep")
    named = {s.skill for s in recipe.steps if s.kind == "skill"}
    for outside in ("qc-tieout", "gpc-scrub", "cost-of-capital", "dcf-build",
                    "ebitda-normalization"):
        assert outside not in named
