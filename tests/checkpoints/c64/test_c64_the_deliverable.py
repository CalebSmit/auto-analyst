"""C64 - what was left once a run could reach the writer.

* **P111** gave the release step the figures as well as the mapping. It had declared
  `needs: [model]` and asked for `timing` too, and failed five steps into the first real
  folder.
* **P112** built the waiver the refusal had been promising since it was written. Before
  it, a mapping that was not total meant no workbook at all - the person accepted the
  item, the step re-ran, recomputed the same list and asked again.
* **P113** rewrote the page an analyst reads first, because P108 and P112 had between them
  made two of its sentences false.

Each is green alone. The disagreement a checkpoint has to look for is that all three are
about **the same moment**: the one step that turns a run into a file somebody opens. P111
decides what that step is handed, P112 decides whether it is allowed to write at all, and
P113 decides what the file says about itself. A change to any one of them can make another
a lie without failing its own tests - which is exactly how P113 came to exist.

So: does the step that writes get what it needs, may it write a partial workbook when a
person says so, and does the page describe the workbook that comes out?
"""

from __future__ import annotations

import ast
from pathlib import Path
from typing import Any

import pytest
import yaml
from openpyxl import load_workbook

from deliver.builder import NOT_FOUND, waiver_for
from graph.execute import Blocked
from tests.p21.conftest import mapping_for
from tools.synthesize_templates import WHAT_IT_IS
from tools.write_template import read_named

ROOT = Path(__file__).resolve().parents[3]
RECIPE = ROOT / "recipes" / "valuation-input-prep.yaml"
WORKBOOK = ROOT / "templates" / "Chartwell_Input_Help_v1.xlsx"
DRIVER = ROOT / "workers" / "driver.py"



def build(world: Any, *, waived: tuple[str, ...] = (), drop: tuple[str, ...] = ()) -> Any:
    mapping = mapping_for(world)
    if drop:
        mapping = {
            **mapping,
            "template_mapping": {
                **mapping["template_mapping"],
                "bindings": [
                    b
                    for b in mapping["template_mapping"]["bindings"]
                    if b["named_range"] not in drop
                ],
            },
        }
    return world["builder"].build(
        mapping,
        world["inputs"].figures(),
        world["template"],
        run_id="run_c64",
        engagement_id="ENG-C64",
        step_id="release",
        waived=waived,
    )


# ============================================================ handed what it asks for


def test_the_release_step_is_handed_both_halves() -> None:
    """P111. The mapping says where each figure goes; the catalogue built from timing's
    output says what each figure is. A step handed one of them has a workbook it cannot
    fill in."""
    recipe = yaml.safe_load(RECIPE.read_text(encoding="utf-8"))
    release = next(s for s in recipe["steps"] if s["id"] == "release")
    assert set(release["needs"]) >= {"model", "timing"}


def test_the_handler_still_asks_for_exactly_those() -> None:
    """The other side of the same fact. A recipe updated to hand over something the
    handler stopped wanting is a recipe describing last month's code."""
    source = DRIVER.read_text(encoding="utf-8")
    release = next(
        n
        for n in ast.walk(ast.parse(source))
        if isinstance(n, ast.FunctionDef) and n.name == "_release"
    )
    body = ast.get_source_segment(source, release) or ""
    assert '_output(attempt, "model")' in body
    assert "self._figures(attempt, context)" in body


# ============================================================ may it write


def test_a_partial_mapping_is_a_question_before_it_is_a_workbook(world: Any) -> None:
    """P112 did not loosen the refusal. Nothing is written because a mapping was
    incomplete and nobody has looked at it."""
    dropped = (world["template"].inputs.required[0],)
    with pytest.raises(Blocked) as blocked:
        build(world, drop=dropped)
    assert blocked.value.item_id == waiver_for("release", dropped)


def test_and_a_workbook_once_somebody_has(world: Any) -> None:
    """The livelock, gone. This is the difference between an analyst getting a file with
    gaps and getting nothing at all, which on a real folder is most runs."""
    dropped = (world["template"].inputs.required[0],)
    built = build(world, drop=dropped, waived=(waiver_for("release", dropped),))
    assert built.data and built.waived == dropped


def test_the_waiver_is_of_the_question_that_was_asked(world: Any) -> None:
    """Where P111 and P112 could disagree. P111 changed what the step is handed, which
    changes what the mapping can bind, which changes the list of unfilled inputs - and a
    waiver that survived that change would be a person's yes applied to a question they
    never saw."""
    two = tuple(world["template"].inputs.required[:2])
    with pytest.raises(Blocked):
        build(world, drop=two, waived=(waiver_for("release", two[:1]),))


# ============================================================ what the file says


def test_a_waived_input_says_so_in_the_workbook(world: Any) -> None:
    dropped = (world["template"].inputs.required[0],)
    built = build(world, drop=dropped, waived=(waiver_for("release", dropped),))
    assert read_named(built.data, [dropped[0]])[dropped[0]] == NOT_FOUND


def test_and_the_front_page_says_to_expect_it() -> None:
    """P113, tied to P112. The page is the only part of the file an analyst reads rather
    than checks, and it has to name the same word the cell holds - "unavailable" on the
    page and "not found" in the cell is a person looking for something that is not
    there."""
    sheet = load_workbook(WORKBOOK).worksheets[0]
    said = " ".join(str(c.value) for row in sheet.iter_rows() for c in row if c.value)
    assert NOT_FOUND.upper() in said.upper()
    assert NOT_FOUND.upper() in "\n".join(WHAT_IT_IS).upper()


def test_the_page_and_the_writer_cannot_drift(world: Any) -> None:
    """The mechanism, not the coincidence. The page imports the writer's constant, so a
    change to the word changes both."""
    source = (ROOT / "tools" / "synthesize_templates.py").read_text(encoding="utf-8")
    assert "from deliver.builder import NOT_FOUND" in source

    dropped = (world["template"].inputs.required[0],)
    built = build(world, drop=dropped, waived=(waiver_for("release", dropped),))
    in_the_cell = read_named(built.data, [dropped[0]])[dropped[0]]
    assert str(in_the_cell).upper() in "\n".join(WHAT_IT_IS).upper()


def test_the_page_still_refuses_to_be_the_model() -> None:
    """Unchanged through all three, and the reason the sheet is first."""
    sheet = load_workbook(WORKBOOK).worksheets[0]
    said = " ".join(str(c.value) for row in sheet.iter_rows() for c in row if c.value)
    assert "not the Chartwell model" in said
    assert "Nothing here is a valuation" in said


# ============================================================ the record of all three


def test_what_went_out_unfilled_is_named_in_the_delivery(world: Any) -> None:
    """A workbook with gaps is only safe if the gaps are on the record. The run says which
    inputs, the cell says so to whoever opens it, and the page says what the cell means."""
    dropped = (world["template"].inputs.required[0],)
    built = build(world, drop=dropped, waived=(waiver_for("release", dropped),))
    assert built.waived == dropped
    assert '"waived": list(built.waived),' in DRIVER.read_text(encoding="utf-8")


def test_the_lock_still_covers_the_whole_write(world: Any) -> None:
    """Across all three phases. More is written than before - the markers - and the
    evidence that the write was faithful has to cover them too."""
    dropped = (world["template"].inputs.required[0],)
    built = build(world, drop=dropped, waived=(waiver_for("release", dropped),))
    assert built.lock.passed
    for name in built.writes:
        assert name in built.locked_text
