"""C56 - A list nobody wrote, tracked against a folder nobody labelled.

P94 gave a new engagement a request list. P95 reads a folder against it and writes the
tracker. Separately each is checkable; what a checkpoint asks is whether the two halves
still agree once a real sequence runs through them - a client sending nothing, then some
of it, then a corrected version, with a person writing in the workbook in between.

The two claims that span both phases:

1. **A guess stays labelled a guess**, from the YAML through the report to the page. The
   whole list is invented, and a coverage number taken from it is worth checking rather
   than acting on.
2. **A refresh never costs a person work.** If notes, `In Model?` or `Not Applicable`
   could be lost, they would be re-entered once and then never again.
"""

from __future__ import annotations

import ast
from pathlib import Path
from typing import Any

from openpyxl import Workbook, load_workbook

from intake.request_list import Status, for_scope
from intake.tracker import THEIRS, Tracker, look_at, read_tracker, write_tracker
from ports.storage.local_folder import LocalFolderStorage
from tests.p17.conftest import build_pdf

ROOT = Path(__file__).resolve().parents[3]


def tracker_for(folder: Path, held: dict[str, dict[str, str]] | None = None) -> Tracker:
    listing = for_scope("valuation/non-esop")
    assert listing is not None
    return Tracker(listing, look_at(LocalFolderStorage(folder)), held or {})


def rows_of(path: Path) -> dict[str, dict[str, Any]]:
    book = load_workbook(path)
    try:
        sheet = book["Request list"]
        header = [c.value for c in sheet[1]]
        return {
            row[0]: dict(zip(header, row, strict=False))
            for row in sheet.iter_rows(min_row=2, values_only=True)
        }
    finally:
        book.close()


def type_into(path: Path, ref: str, column: str, value: str) -> None:
    book = load_workbook(path)
    try:
        sheet = book["Request list"]
        header = [c.value for c in sheet[1]]
        for row in sheet.iter_rows(min_row=2):
            if row[0].value == ref:
                row[header.index(column)].value = value
        book.save(path)
    finally:
        book.close()


def a_trial_balance(path: Path, cash: int) -> None:
    book = Workbook()
    sheet = book.active
    sheet.title = "TB"
    sheet.append(["Account", "Debit", "Credit"])
    sheet.append(["1000 Cash", cash, 0])
    book.save(path)


# ============================================================ the sequence that really happens


def test_an_engagement_goes_from_nothing_to_tracked_without_anybody_writing_a_list(
    tmp_path: Path,
) -> None:
    """The whole point of the pair, run as a sequence rather than as two unit tests.

    Monday: the folder is empty and the engagement has no request list at all. That is not
    a failure state - it is every new engagement, and it is how the first real test will
    arrive.
    """
    inbox = tmp_path / "inbox"
    inbox.mkdir()
    tracker = tmp_path / "DD request list.xlsx"

    # Monday. Nothing has arrived, and there was never a list.
    write_tracker(tracker_for(inbox), tracker)
    monday = rows_of(tracker)
    assert monday["A.3"]["Status"] == Status.OUTSTANDING.value
    assert len(monday) > 30, "a list came into existence"

    # Tuesday. The client sends three things, named the way clients name them.
    a_trial_balance(inbox / "TB 2025 final.xlsx", 1000)
    (inbox / "FS 2021-2025.pdf").write_bytes(
        build_pdf([("Consolidated Balance Sheet", 12), ("Total assets", 10)])
    )
    (inbox / "holiday photos.zip").write_bytes(b"PK")
    write_tracker(tracker_for(inbox, read_tracker(tracker)), tracker)

    tuesday = rows_of(tracker)
    assert tuesday["A.3"]["Status"] == Status.PROVIDED.value
    assert tuesday["A.1"]["Status"] == Status.PROVIDED.value
    assert tuesday["C.2"]["Status"] == Status.OUTSTANDING.value, "nobody sent an org chart"

    # Wednesday. He works through it and writes in the workbook.
    type_into(tracker, "D.5", "Not Applicable", "Y")
    type_into(tracker, "D.5", "Notes", "no union - confirmed with the controller")
    type_into(tracker, "A.3", "In Model?", "Y")
    type_into(tracker, "A.1", "Notes", "FY21 is unaudited - ask")

    # Thursday. The client re-sends a corrected trial balance and he presses update.
    a_trial_balance(inbox / "TB 2025 final.xlsx", 1100)
    write_tracker(tracker_for(inbox, read_tracker(tracker)), tracker)

    thursday = rows_of(tracker)
    assert thursday["D.5"]["Notes"] == "no union - confirmed with the controller"
    assert thursday["D.5"]["Not Applicable"] == "Y"
    assert thursday["D.5"]["Status"] == Status.NOT_APPLICABLE.value
    assert thursday["A.3"]["In Model?"] == "Y"
    assert thursday["A.1"]["Notes"] == "FY21 is unaudited - ask"
    assert thursday["A.3"]["Status"] == Status.PROVIDED.value


def test_the_file_nobody_asked_for_is_still_being_reported_at_the_end(
    tmp_path: Path,
) -> None:
    """Across refreshes, not just on the first. It is the one thing in the folder nobody
    has an answer for, and it is the sort of thing that quietly stops being mentioned."""
    inbox = tmp_path / "inbox"
    inbox.mkdir()
    (inbox / "holiday photos.zip").write_bytes(b"PK")
    tracker = tmp_path / "t.xlsx"
    for _ in range(3):
        write_tracker(tracker_for(inbox, read_tracker(tracker)), tracker)

    book = load_workbook(tracker)
    try:
        listed = [r[0] for r in book["Not matched"].iter_rows(min_row=2, values_only=True)]
    finally:
        book.close()
    assert listed == ["holiday photos.zip"]


# ============================================================ a guess stays labelled one


def test_the_label_survives_from_the_yaml_to_the_page(tmp_path: Path) -> None:
    """Three places have to agree, and each is a place somebody could read the number and
    stop: the file, the report the route returns, and the words on the screen."""
    document = (
        ROOT / "intake" / "request_lists" / "valuation-non-esop.yaml"
    ).read_text(encoding="utf-8")
    assert "starting list" in document.lower()

    inbox = tmp_path / "inbox"
    inbox.mkdir()
    assert tracker_for(inbox).report()["is_a_starting_point"] is True

    from admin.page import PAGE

    assert "is_a_starting_point" in PAGE
    assert "starting point, not the firm" in PAGE


def test_the_default_list_owes_nothing_to_anybody_the_register_knows() -> None:
    """Held here as well as in P94 because this is the file that ships. `interview.py
    scan` polices what leaves `knowledge/`; nothing polices what arrives in `intake/`."""
    register = ROOT / "knowledge" / "identifiers.txt"
    names = [
        line.strip()
        for line in register.read_text(encoding="utf-8").splitlines()
        if line.strip() and not line.startswith("#")
    ]
    assert names
    text = (ROOT / "intake" / "request_lists" / "valuation-non-esop.yaml").read_text(
        encoding="utf-8"
    ).lower()
    assert not [n for n in names if n.lower() in text]


# ============================================================ what the pair may not do


def test_the_two_halves_use_one_list_and_not_two(tmp_path: Path) -> None:
    """The drift that would look like a button doing nothing: the page showing items the
    update does not know about, or the reverse."""
    inbox = tmp_path / "inbox"
    inbox.mkdir()
    tracker = tmp_path / "t.xlsx"
    write_tracker(tracker_for(inbox), tracker)

    listing = for_scope("valuation/non-esop")
    assert listing is not None
    assert set(rows_of(tracker)) == {str(i.ref) for i in listing.items}


def test_neither_half_reaches_a_model_an_ocr_engine_or_a_run() -> None:
    """A tracker is a reading of a folder, not a run: it takes no engagement lock, passes
    no policy gate and produces nothing anybody signs. It must not start being one by
    accident, and the cheapest way to hold that is on the imports."""
    for module in ("intake/tracker.py", "intake/request_list.py"):
        tree = ast.parse((ROOT / module).read_text(encoding="utf-8"))
        named: set[str] = set()
        for node in ast.walk(tree):
            if isinstance(node, ast.Import):
                named.update(a.name for a in node.names)
            elif isinstance(node, ast.ImportFrom) and node.module:
                named.add(node.module)
        assert named, module
        for forbidden in ("ports.llm", "ports.ocr", "harness", "skills", "graph", "ops.run_store"):
            assert not any(
                n == forbidden or n.startswith(forbidden + ".") for n in named
            ), (module, forbidden)


def test_a_person_s_columns_and_the_tool_s_do_not_overlap() -> None:
    """The whole preservation rule in one line. A column in both sets would be written by
    one and expected by the other, and whichever ran last would win."""
    from intake.tracker import DERIVED

    assert not set(THEIRS) & set(DERIVED)


def test_the_tracker_is_reached_by_something_a_person_can_press() -> None:
    """P68 shipped this machinery unreached and it stayed that way for twenty-six phases.
    This is the join, asserted rather than assumed."""
    import sys

    sys.path.insert(0, str(ROOT))
    from tests.p82.test_nothing_a_run_cannot_reach import ALLOWED, unreachable

    for module in ("intake.tracker", "intake.request_list"):
        assert module not in unreachable(), module
        assert module not in ALLOWED, f"{module} is reached and should not be excused"

    console = (ROOT / "admin" / "console.py").read_text(encoding="utf-8")
    assert "def dd_tracker" in console
