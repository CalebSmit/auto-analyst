"""P137: the three reads, against documents shaped like the work rather than the fixtures.

Wrote a normalisation workbook with the two mistakes the analyst named - *"summing too many
rows or one too few rows"* - a page of appendix commentary written the way a model writes
it, and a three-year table with a margin that collapses and receivables that balloon. Then
ran `aa review` over all three and read what came back.

**The formula audit found exactly what it was asked to**, named the cell, named the rule,
printed the formula, and said which checks it could not run and why. Nothing to do.

The other two said things that were wrong.

**The prose read reported a bullet list as a sentence missing its full stop.** The check
excludes a list, and the rule looked only at the first line of a Markdown block - so the
commonest shape in a report,

    Key considerations include:
    - owner compensation normalised to a market rate
    - rent adjusted to an arm's length rate

is one block that begins with a sentence, and every one of them was a finding. It also said
`ends in 'k'`, which is the last character of the sentence quoted directly above it.

**The reasonableness read asked twelve questions about figures that cannot exist.** Given a
table with signed amounts - a cost of sales of -10,900 rather than 10,900, which is what a
trial balance looks like - it reported a gross margin of 173.6 points and asked *"What
changed in the business - and is it in the forecast?"* A margin above 100 points is not a
business event; it is a statement about the table. Six ratios, two comparisons each, every
one of them a difference between two impossible figures, every one in a confident sentence.

Both are the same failure: a read asserting something it had not established. The fixtures
were in the shape the checks were written against, so neither had been seen.
"""

from __future__ import annotations

import sys
from pathlib import Path
from typing import Any

ROOT = Path(__file__).resolve().parents[2]
if str(ROOT) not in sys.path:
    sys.path.insert(0, str(ROOT))

from registries.taxonomy import load_taxonomy  # noqa: E402
from tools.prose_checks import check as prose_check  # noqa: E402
from tools.prose_checks import paragraphs_of  # noqa: E402
from tools.reasonableness import lines_from, run_questions  # noqa: E402

TAXONOMY = "coa-taxonomy@1.0.0"
YEARS = ("FY2023", "FY2024", "FY2025")
MOVE = {"points": 5.0, "days": 10.0, "x": 0.5}


# --------------------------------------------------------------------------- the prose


def stops(text: str) -> list[str]:
    found = prose_check(paragraphs_of(text)).findings
    return [f.says for f in found if f.check == "paragraph_without_a_stop"]


def test_a_list_introduced_by_a_sentence_is_a_list() -> None:
    """The shape this was found on. One Markdown block, a lead-in and three bullets, and
    the block opens with prose - which is why looking at the first line could not see it."""
    assert not stops(
        "Key considerations include:\n"
        "- owner compensation normalised to a market rate\n"
        "- rent adjusted to an arm's length rate\n"
        "- one-time legal costs added back\n"
    )


def test_a_numbered_list_introduced_by_a_sentence_is_a_list_too() -> None:
    assert not stops("We made three adjustments:\n1. owner pay\n2. rent\n3. legal costs\n")


def test_a_sentence_below_a_list_is_still_checked() -> None:
    """Which is the reason the *last* line decides and not any line. A closing sentence
    under a set of bullets is prose, and it is exactly where a tail gets lost in an edit."""
    assert stops("- owner pay normalised\n- rent adjusted\nThese are the only two we made\n")


def test_a_paragraph_that_lost_its_tail_is_still_found() -> None:
    assert stops("Revenue grew steadily across the three years we examined\n")


def test_an_ordinary_paragraph_and_a_heading_are_left_alone() -> None:
    assert not stops("Revenue grew steadily across the three years we examined.\n")
    assert not stops("# Valuation Summary\n")


def test_the_finding_says_what_is_missing_rather_than_what_is_there() -> None:
    """`ends in 'k'` names the last character of the sentence printed directly above it."""
    (said,) = stops("Revenue grew steadily across the three years we examined\n")
    assert said == "no full stop at the end"
    assert "ends in" not in said


# -------------------------------------------------------------------- the reasonableness


def table(**signs: float) -> list[Any]:
    """A three-year table. `signs` flips a category's sign, so one fixture serves both the
    shape the ratios are defined on and the shape a trial balance arrives in."""
    rows = {
        "revenue": (12000, 13250, 14800),
        "cost_of_sales": (7200, 7900, 10900),
        "salaries": (2100, 2250, 2400),
        "rent": (480, 495, 510),
        "accounts_receivable": (1450, 1600, 4300),
        "inventory": (900, 980, 1050),
        "cash": (300, 420, 180),
        "accounts_payable": (1100, 1180, 1260),
        "accrued_liabilities": (240, 260, 280),
    }
    return lines_from(
        {
            "lines": [
                {
                    "category": name,
                    "label": name.replace("_", " ").capitalize(),
                    "locator": f"IS!{i + 2}",
                    "amounts": {
                        y: v * signs.get(name, 1.0)
                        for y, v in zip(YEARS, values, strict=True)
                    },
                }
                for i, (name, values) in enumerate(rows.items())
            ]
        }
    )


def asked(lines: list[Any]) -> dict[str, dict[str, Any]]:
    found = run_questions(lines, load_taxonomy(TAXONOMY), move=MOVE, years=len(YEARS))
    return {f["check"]: f for f in found}


def test_a_margin_above_a_hundred_points_is_a_finding_about_the_table() -> None:
    """Not a question about the business. Profit cannot exceed revenue, so the arithmetic
    that produced 173.6 was given something other than magnitudes."""
    found = asked(table(cost_of_sales=-1.0, salaries=-1.0, rent=-1.0))
    assert "gross margin" in found, "the impossible value is not reported at all"
    says = found["gross margin"]["evidence"]
    assert "above 100.0" in says
    assert "cannot produce" in says
    assert "What changed in the business" not in says, "it is still being asked as a question"


def test_it_names_the_rows_responsible() -> None:
    """A guess at the cause sends somebody looking. A named row is one they can open.

    The rows and not every period of them: a line that is signed is signed in every year,
    and three periods of the same row is one fact printed three times. One signed table
    trips four of the six ratios, so a reader meets this sentence four times in a row -
    everything the same each time is said in as few words as it can be.
    """
    found = asked(table(cost_of_sales=-1.0, salaries=-1.0, rent=-1.0))
    says = found["gross margin"]["evidence"]
    assert "Cost of sales is negative here" in says
    assert says.count("Cost of sales") == 1, "the same row, once per period it is negative in"


def test_it_agrees_with_the_number_of_rows_it_names() -> None:
    """"Cost of sales, Rent is negative here". The second time this module has joined a
    list with code that did not know how many things were in it."""
    found = asked(table(cost_of_sales=-1.0, salaries=-1.0, rent=-1.0))
    assert "Cost of sales and Rent are negative here" in found["operating margin"]["evidence"]


def test_no_movement_is_asked_about_an_impossible_ratio() -> None:
    """Every one would be a difference between two figures that do not mean anything, and
    each would be stated in the same confident sentence as a real one."""
    found = asked(table(cost_of_sales=-1.0, salaries=-1.0, rent=-1.0))
    moves = [name for name in found if name.startswith("gross margin ")]
    assert not moves, f"it still asks how an impossible ratio moved: {moves}"


def test_a_negative_days_or_current_ratio_is_the_same_finding() -> None:
    """Both are magnitudes over magnitudes, so neither can be below zero."""
    found = asked(table(accounts_payable=-1.0, accrued_liabilities=-1.0))
    says = found["the current ratio"]["evidence"]
    assert "below 0.0" in says and "cannot produce" in says


def test_a_table_in_the_right_shape_is_unchanged() -> None:
    """The whole check has to be invisible on a table that is correct, or it is a tax on
    every good run to catch a bad one."""
    found = asked(table())
    for ratio in ("gross margin", "operating margin", "the current ratio"):
        assert ratio not in found, f"{ratio} is reported on a table with nothing wrong with it"
    # And the real question the fixture was built to raise is still asked.
    assert any(name.startswith("days receivable ") for name in found), (
        "receivables trebling in the last year is no longer asked about"
    )


def test_a_margin_can_still_be_negative() -> None:
    """A business can lose more than it earns, and a reviewer wants to be asked about that
    rather than told it is impossible. Only the upper bound is arithmetic."""
    lines = lines_from(
        {
            "lines": [
                {
                    "category": "revenue",
                    "amounts": dict(zip(YEARS, (1000.0, 900.0, 800.0), strict=True)),
                },
                {
                    "category": "cost_of_sales",
                    "amounts": dict(zip(YEARS, (900.0, 1100.0, 1600.0), strict=True)),
                },
            ]
        }
    )
    found = asked(lines)
    assert "gross margin" not in found, "a loss-making business is being called impossible"


def test_the_refusal_says_what_a_line_looks_like() -> None:
    """`cannot convert dictionary update sequence element #0 to a sequence` sends somebody
    to read a traceback instead of the line they typed. `amounts` as a list rather than a
    mapping is the ordinary mistake."""
    import pytest

    with pytest.raises(ValueError) as refused:
        lines_from({"lines": [{"category": "revenue", "amounts": [1, 2, 3]}]})
    said = str(refused.value)
    assert '"amounts"' in said and "keyed by period" in said
    assert "magnitudes" in said, "the convention that matters is not mentioned"


# ------------------------------------------------------------------------ and the grammar


def test_the_pairing_rule_says_it_in_a_sentence() -> None:
    """"in every period, related-party rent appear together" - the rule's name is a thing,
    so the verb was agreeing with the wrong noun."""
    from registries.taxonomy import load_taxonomy as _load

    lines = lines_from(
        {
            "lines": [
                {"category": "inventory", "amounts": {y: 100.0 for y in YEARS}},
                {"category": "cost_of_sales", "amounts": {y: 500.0 for y in YEARS}},
            ]
        }
    )
    found = {f["check"]: f for f in run_questions(lines, _load(TAXONOMY))}
    said = found["inventory and cost of sales"]["evidence"]
    assert "appear together" not in said, said
    assert "is accounted for on both sides" in said


def test_the_formula_audit_still_finds_both_mistakes_the_analyst_named() -> None:
    """The read that needed nothing done to it, asserted so that it keeps needing nothing.

    A total one row short and a total that reaches up over a subtotal, which is how he
    described it: *"it's summing too many rows or one too few rows"*. The second one is
    written as reaching over `Gross profit` rather than over `Revenue`, because that is the
    version the audit can actually prove - a range that covers a cell which is itself a
    total double-counts, and a range that happens to include one more ordinary row is a
    judgement about what the rows mean.
    """
    from io import BytesIO

    from openpyxl import Workbook

    from tools.formula_audit import audit

    wb = Workbook()
    sheet = wb.active
    assert sheet is not None
    sheet.title = "IS"
    for row, (label, *values) in enumerate(
        [
            ("Revenue", 12000, 13250, 14800),
            ("Cost of sales", 7200, 7900, 8600),
            ("Gross profit", None, None, None),
            ("Salaries", 2100, 2250, 2400),
            ("Rent", 480, 495, 510),
            ("Marketing", 310, 340, 395),
            ("Other opex", 260, 275, 290),
        ],
        start=1,
    ):
        sheet.cell(row=row, column=1, value=label)
        for col, value in enumerate(values, start=2):
            if value is not None:
                sheet.cell(row=row, column=col, value=value)
    for col in ("B", "C", "D"):
        sheet[f"{col}3"] = f"=SUM({col}1:{col}2)"
    sheet["B8"] = "=SUM(B4:B7)"
    sheet["C8"] = "=SUM(C4:C6)"  # one row short
    sheet["D8"] = "=SUM(D3:D7)"  # one row too many, and that row is itself a total
    buffer = BytesIO()
    wb.save(buffer)

    findings = audit(buffer.getvalue()).findings
    wrong = {f"{f.sheet}!{f.cell}": f.check for f in findings if f.severity == "wrong"}
    assert wrong.get("IS!C8") == "total_misses_rows", wrong
    assert wrong.get("IS!D8") == "total_counts_a_total", wrong
