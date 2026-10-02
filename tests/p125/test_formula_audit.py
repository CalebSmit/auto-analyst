"""The formula audit: the mistakes a spreadsheet makes that a spreadsheet cannot see.

Every test here is a defect the analyst named, put into a workbook on purpose and then
looked for - and, as many times again, a correct workbook the audit has to be silent on.
The second half is the half that is easy to lose: a check that fires on everything is
worth nothing, because the finding nobody trusts is the finding nobody reads.

What the tests at the foot hold the module to, beyond finding things:

* every check a finding carries is described in `CHECKS`, so no finding arrives as an
  unexplained id;
* a severity is one of the three the module declares;
* the footing check reports that it did not run rather than passing, when the file has
  no saved values for it to compare - `statement_checks.py` draws the same line, and the
  reason is the same: absence of evidence is the commonest way a checklist lies.
"""

from __future__ import annotations

from tests.workbooks import a_model, a_tidy_model, with_cached_values
from tools.formula_audit import (
    CHECKS,
    audit,
    normalise,
    refs_in,
)

# ------------------------------------------------------------ the defects, found


def test_a_total_that_stops_a_row_short_of_its_block() -> None:
    """`=SUM(B6:B7)` under a block running to row 8. The defect the analyst described:
    one too few rows, a plausible number, no error anywhere in the file."""
    result = audit(a_model())
    short = [f for f in result.findings if f.check == "total_misses_rows"]
    assert [f.cell for f in short] == ["B9"], [str(f) for f in result.findings]
    assert short[0].severity == "wrong"
    assert short[0].evidence["rows_left_out"] == [8]
    assert "row 8" in short[0].says


def test_a_total_whose_range_swallows_another_total() -> None:
    """`=SUM(B18:B20)` where B20 is itself `=SUM(B18:B19)`. The block is counted twice,
    and the grand total is double what it should be while looking entirely ordinary."""
    result = audit(a_model())
    doubled = {f.cell: f for f in result.findings if f.check == "total_counts_a_total"}
    assert set(doubled) == {"B21"}, [str(f) for f in result.findings]
    assert doubled["B21"].evidence["totals_inside"] == [20]
    assert doubled["B21"].severity == "wrong"


def test_a_calculated_line_inside_a_range_is_not_called_a_doubled_total() -> None:
    """`=SUM(D4:D8)` covers D4, which is `=D2-D3` - calculated, but not a total. Summing
    derived cells is ordinary and this check does not claim otherwise; what it reports
    about D9 is the two things it can prove - the range reaches across a blank row, and
    it starts one row below a number."""
    result = audit(a_model())
    about_d9 = {f.check for f in result.findings if f.cell == "D9"}
    assert "total_counts_a_total" not in about_d9
    assert about_d9 == {"total_crosses_a_blank_row", "total_starts_late"}, about_d9


def test_the_three_totals_on_one_row_that_disagree_with_each_other() -> None:
    """Three year columns, three different ranges under the same block. Nothing proves
    which is right, and that nothing agrees is itself the finding."""
    result = audit(a_model())
    disagreeing = [f for f in result.findings if f.check == "run_has_no_agreement"]
    assert [f.cell for f in disagreeing] == ["B9"], [str(f) for f in result.findings]
    assert disagreeing[0].evidence["cells"] == ["B9", "C9", "D9"]


def test_a_typed_number_in_a_row_that_otherwise_calculates() -> None:
    """D10 is 999 where B10 and C10 compute operating income. Two agreeing neighbours
    against one exception is a majority: a three-year block is the layout this matters
    most on, and a quorum of three would have been blind to it."""
    result = audit(a_model())
    plugs = [f for f in result.findings if f.check == "number_where_the_row_calculates"]
    assert "D10" in [f.cell for f in plugs], [str(f) for f in result.findings]
    plug = next(f for f in plugs if f.cell == "D10")
    assert plug.evidence["value"] == 999
    assert plug.evidence["axis"] == "row"


def test_a_formula_one_row_off_the_pattern_its_neighbours_share() -> None:
    """`=C14/C14` among `=B14/B13`, `=D14/D13`, `=E14/E13`. The copy-paste that landed
    wrong: it computes 100% every year and looks like a margin."""
    result = audit(a_model())
    broken = [f for f in result.findings if f.check == "formula_breaks_the_pattern"]
    assert "C15" in [f.cell for f in broken], [str(f) for f in result.findings]
    one = next(f for f in broken if f.cell == "C15")
    assert one.evidence["agree"] == ["B15", "D15", "E15"]
    assert one.evidence["shape_here"] != one.evidence["shape_agreed"]


def test_a_cached_total_that_does_not_equal_what_it_adds() -> None:
    """The file somebody edited and did not recalculate. C9 sums 110 + 52 + 26 = 188;
    the value Excel last saved for it says 999."""
    data = with_cached_values(
        a_model(),
        {"IS": {"C9": 999.0}},
    )
    result = audit(data)
    off = [f for f in result.findings if f.check == "total_does_not_foot"]
    assert [f.cell for f in off] == ["C9"], [str(f) for f in result.findings]
    assert off[0].evidence["stated"] == 999.0
    assert off[0].evidence["recomputed"] == 188.0
    assert off[0].evidence["delta"] == 811.0


def test_a_cached_total_that_does_foot_is_not_reported() -> None:
    """The same cell with the right value. The check has to be able to pass."""
    data = with_cached_values(a_model(), {"IS": {"C9": 188.0}})
    result = audit(data)
    assert [f for f in result.findings if f.check == "total_does_not_foot"] == []


# ------------------------------------------------------------ and silence, where due


def test_a_correct_workbook_produces_nothing() -> None:
    """The half of a check that is easy to lose. Three year columns, a consistent
    expense block, a total that covers it - and no finding at either severity."""
    result = audit(a_tidy_model())
    assert result.findings == (), [str(f) for f in result.findings]
    assert result.clean
    assert result.formulas == 9


def test_a_total_that_correctly_excludes_the_line_above_its_range() -> None:
    """`total_starts_late` fires on a *number* directly above the range, not on a
    calculated line: a range that starts below the subtotal above it is right, and a
    check that could not tell the two apart would fire on every correct model."""
    result = audit(a_tidy_model())
    assert [f for f in result.findings if f.check == "total_starts_late"] == []


def test_arithmetic_around_a_sum_is_not_treated_as_a_footing_failure() -> None:
    """`=SUM(C6:C8)-C7` has the same function set as a bare total and a legitimately
    different answer. Reporting it would be the module being wrong about a right file."""
    from openpyxl import Workbook

    from tests.workbooks import _bytes  # noqa: PLC0415 - built here, for this one case

    wb = Workbook()
    ws = wb.active
    assert ws is not None
    ws["C6"], ws["C7"], ws["C8"] = 110, 52, 26
    ws["C9"] = "=SUM(C6:C8)-C7"
    data = with_cached_values(_bytes(wb), {"Sheet": {"C9": 136.0}})
    result = audit(data)
    assert [f for f in result.findings if f.check == "total_does_not_foot"] == []


# ----------------------------------------------------- what the module says about itself


def test_every_finding_names_a_check_the_module_describes() -> None:
    """A finding that arrives as an id nobody has written a sentence for is a finding a
    reviewer has to guess at."""
    result = audit(a_model(), include_convention=True)
    assert result.findings
    for finding in result.findings:
        assert finding.check in CHECKS, finding.check
        severity, description = CHECKS[finding.check]
        assert finding.severity == severity
        assert description.strip()


def test_the_footing_check_says_it_did_not_run_rather_than_passing() -> None:
    """A workbook openpyxl wrote has no saved values at all, so there is nothing to
    compare a total against. That is reported, not passed over."""
    result = audit(a_tidy_model())
    assert any("total_does_not_foot" in note for note in result.skipped)
    assert any("could not run on 3" in note for note in result.skipped)


def test_a_partly_cached_sheet_says_how_many_totals_it_could_not_foot() -> None:
    """The dangerous middle case. One cell given a saved value and the rest without:
    reporting nothing wrong would be reporting nothing about the two it never looked at,
    which is the failure mode this project calls `could_not_evaluate` elsewhere."""
    result = audit(with_cached_values(a_tidy_model(), {"IS": {"C9": 188.0}}))
    assert [f for f in result.findings if f.check == "total_does_not_foot"] == []
    assert any("ran on 1 total(s) and could not run on 2" in n for n in result.skipped)


def test_asking_for_a_sheet_the_workbook_does_not_have_is_reported() -> None:
    result = audit(a_tidy_model(), sheets=["IS", "Nowhere"])
    assert result.sheets == ("IS",)
    assert any("Nowhere" in note for note in result.skipped)


def test_convention_findings_are_off_unless_asked_for() -> None:
    """A cross-sheet reference by address is correct to report and, on somebody else's
    model, numerous. It is its own severity so it never crowds out the rest."""
    from openpyxl import Workbook

    from tests.workbooks import _bytes  # noqa: PLC0415

    wb = Workbook()
    first = wb.active
    assert first is not None
    first.title = "Inputs"
    first["A1"] = 5
    second = wb.create_sheet("Calc")
    second["A1"] = "=Inputs!A1*2"
    data = _bytes(wb)

    quiet = audit(data)
    assert [f for f in quiet.findings if f.check == "cell_address_across_sheets"] == []

    loud = audit(data, include_convention=True)
    offenders = [f for f in loud.findings if f.check == "cell_address_across_sheets"]
    assert [f.cell for f in offenders] == ["A1"]
    assert offenders[0].severity == "convention"
    assert not quiet.of("convention")
    assert loud.clean  # a convention finding does not make a workbook unclean


def test_a_defined_name_is_not_reported_as_a_cell_address() -> None:
    """`=SUM(Revenue)` is the thing the convention check wants people to do. Reporting
    it would invert the rule."""
    refs = refs_in("=SUM(Revenue)")
    assert len(refs) == 1
    assert refs[0].name == "Revenue"
    assert not refs[0].is_address


# --------------------------------------------------------------- the normalisation


def test_two_cells_doing_the_same_thing_normalise_alike() -> None:
    assert normalise("=C5*D5", row=6, col=3) == normalise("=D5*E5", row=6, col=4)


def test_two_cells_doing_different_things_do_not() -> None:
    assert normalise("=C5*D5", row=6, col=3) != normalise("=D5*E4", row=6, col=4)


def test_an_absolute_reference_keeps_its_address() -> None:
    """Two cells pointing at one fixed cell are doing the same thing, and a relative
    offset would make them look different."""
    assert normalise("=C5*$B$1", row=6, col=3) == normalise("=D5*$B$1", row=6, col=4)
    assert "$B" in normalise("=C5*$B$1", row=6, col=3)


def test_a_whole_column_reference_is_carried_without_bounds() -> None:
    """`=SUM(C:C)` is real and nothing here can say which rows it covers. It is carried
    rather than dropped, so a later check can see it exists."""
    refs = refs_in("=SUM(C:C)")
    assert len(refs) == 1
    assert not refs[0].is_address
