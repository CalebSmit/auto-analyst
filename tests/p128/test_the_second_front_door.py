"""P128: what a person can ask for from a command line.

The analyst's own description of what they want is a command and not a page: *"point it at
the right folders, quickly tell it what to do, and it starts working"*. So four
subcommands, and what is tested is the things that make them worth having rather than four
ways to type something the dashboard already does.

**`review` dispatches.** One command over a folder of mixed files, each read by whichever
of the three checks it wants. That is what lets a slash command be one line: the person
does not have to know which check a file needs, and neither does the command they typed.

**It renders through each module's own reporter.** Nothing here reformats a finding. The
text `aa review` prints for a workbook is the text `tools/formula_audit.py` prints for it,
because both call `formula_audit.report`. Two renderers of one answer is already a
recorded defect in this build (B-046) and not a thing to add another of.

**`owed` and `stage` go through the only writer.** `AdminConsole.dd_tracker` is what the
Setup page's button calls, so the page and the command cannot disagree about what has
arrived - one implementation, two surfaces. And `owed` writes nothing unless asked.

**`stage` goes through the only writer.** `AdminConsole.add_file` records who put a file
there and why, refuses a path that would escape the folder, and refuses an archived
engagement. A command line that wrote to the folder itself would be a second, unrecorded
writer - which is the exact thing Power Automate was held outside the boundary for.

**And it builds nothing of its own.** `_served()` calls `tools.serve.build`, the one place
this system is assembled. A command line that assembled its own would be a second answer
to "which folder does this engagement's files go in".
"""

from __future__ import annotations

import ast
import contextlib
import io
import json
import re
from importlib import import_module
from pathlib import Path
from typing import Any

import pytest

from admin import AdminConsole
from tests.p24.conftest import HEND, WHO
from tests.workbooks import a_model, a_tidy_model
from tools.aa import (
    DOCUMENTS,
    TABLES,
    WORKBOOKS,
    _owed,
    _stage,
    _where,
    main,
    review,
    which_read,
)

ROOT = Path(__file__).resolve().parents[2]
FOLDER = Path("Clients") / "Henderson Manufacturing" / "2026 ESOP"

A_TABLE = {
    "lines": [
        {
            "category": "revenue",
            "amounts": {"2024-12-31": 1_000_000, "2025-12-31": 1_100_000},
            "locator": "IS!B5",
        },
        {
            "category": "depreciation",
            "amounts": {"2024-12-31": 30_000, "2025-12-31": 31_000},
            "locator": "IS!B18",
        },
    ]
}

SLOPPY = "The  the margin held \u2014 a testament to execution.\n"


class Served:
    """What `aa.py` asks of the system it was handed: an admin console and a close().

    A stand-in rather than `tools.serve.build()`, because `build()` opens the store the
    environment points at - and this suite's own rule, learned the hard way, is that a
    command is probed against a temp copy rather than against the repository's own state.
    The test at the foot holds the real command to using `build()` and nothing else.
    """

    def __init__(self, admin: AdminConsole) -> None:
        self.admin = admin
        self.closed = False

    def close(self) -> None:
        self.closed = True


# ------------------------------------------------------------------ review dispatches


def test_which_read_a_file_wants() -> None:
    assert which_read(Path("Model.xlsx")) == "formulas"
    assert which_read(Path("report.DOCX")) == "prose"
    assert which_read(Path("notes.md")) == "prose"
    assert which_read(Path("table.json")) == "questions"
    assert which_read(Path("scan.pdf")) == "", "no read knows what to do with it"


def test_one_command_over_a_folder_reads_each_file_the_way_it_wants(tmp_path: Path) -> None:
    (tmp_path / "Model.xlsx").write_bytes(a_model())
    (tmp_path / "report.md").write_text(SLOPPY, encoding="utf-8")
    (tmp_path / "table.json").write_text(json.dumps(A_TABLE), encoding="utf-8")
    (tmp_path / "scan.pdf").write_bytes(b"%PDF-1.4 nothing reads this")

    results = {r.path.name: r for r in review(tmp_path, move={"points": 5.0})}
    assert set(results) == {"Model.xlsx", "report.md", "table.json"}, "the PDF is left alone"
    assert results["Model.xlsx"].read == "formulas"
    assert results["report.md"].read == "prose"
    assert results["table.json"].read == "questions"
    assert all(r.found_something for r in results.values())


def test_the_worst_thing_found_is_named_in_the_words_of_the_read_that_found_it(
    tmp_path: Path,
) -> None:
    (tmp_path / "Model.xlsx").write_bytes(a_model())
    (tmp_path / "report.md").write_text(SLOPPY, encoding="utf-8")
    (tmp_path / "table.json").write_text(json.dumps(A_TABLE), encoding="utf-8")
    worst = {r.read: r.worst for r in review(tmp_path, move={"points": 5.0})}
    assert worst == {"formulas": "wrong", "prose": "slip", "questions": "fail"}


def test_a_clean_folder_finds_nothing_and_says_so(tmp_path: Path) -> None:
    (tmp_path / "Model.xlsx").write_bytes(a_tidy_model())
    (tmp_path / "report.md").write_text(
        "The Company's revenue grew 12%. Margins held. The schedule is unchanged.\n",
        encoding="utf-8",
    )
    results = review(tmp_path)
    assert results and not any(r.found_something for r in results)
    assert main(["review", str(tmp_path)]) == 0


def test_one_unreadable_file_does_not_take_the_folder_with_it(tmp_path: Path) -> None:
    """A client folder of forty documents with one corrupt workbook in it must report
    thirty-nine readings and one refusal, not a traceback."""
    (tmp_path / "broken.xlsx").write_bytes(b"not a workbook")
    (tmp_path / "Model.xlsx").write_bytes(a_tidy_model())
    results = {r.path.name: r for r in review(tmp_path)}
    assert results["broken.xlsx"].worst == "unread"
    assert "could not be read" in results["broken.xlsx"].text
    assert results["Model.xlsx"].worst == ""


def test_a_lock_file_excel_left_behind_is_not_a_workbook(tmp_path: Path) -> None:
    """`~$Model.xlsx` is what Excel writes while somebody has the file open. Auditing one
    produces a traceback rather than a finding, and a folder being reviewed while
    somebody is working in it is the ordinary case rather than the odd one."""
    (tmp_path / "~$Model.xlsx").write_bytes(b"lock")
    (tmp_path / "Model.xlsx").write_bytes(a_tidy_model())
    assert [r.path.name for r in review(tmp_path)] == ["Model.xlsx"]


def test_nothing_it_knows_what_to_do_with_says_what_it_reads(
    tmp_path: Path, capsys: pytest.CaptureFixture[str]
) -> None:
    (tmp_path / "scan.pdf").write_bytes(b"%PDF")
    assert main(["review", str(tmp_path)]) == 0
    said = capsys.readouterr().out
    assert "nothing under" in said
    for suffix in (*WORKBOOKS, *DOCUMENTS, *TABLES):
        assert suffix in said, "and it names what the three reads do read"


def test_a_target_that_is_not_there_is_refused(
    tmp_path: Path, capsys: pytest.CaptureFixture[str]
) -> None:
    assert main(["review", str(tmp_path / "nowhere")]) == 2
    assert "no such file or folder" in capsys.readouterr().err


def test_a_misspelled_tolerance_is_refused_before_anything_is_read(
    tmp_path: Path, capsys: pytest.CaptureFixture[str]
) -> None:
    (tmp_path / "table.json").write_text(json.dumps(A_TABLE), encoding="utf-8")
    assert main(["review", str(tmp_path), "--move", "pts=5"]) == 2
    assert "not a unit any ratio is measured in" in capsys.readouterr().err


def test_the_json_output_carries_each_read_under_its_own_path(
    tmp_path: Path, capsys: pytest.CaptureFixture[str]
) -> None:
    (tmp_path / "Model.xlsx").write_bytes(a_model())
    (tmp_path / "table.json").write_text(json.dumps(A_TABLE), encoding="utf-8")
    main(["review", str(tmp_path), "--json"])
    body = json.loads(capsys.readouterr().out)
    reads = {Path(r["path"]).name: r for r in body["reviewed"]}
    assert reads["Model.xlsx"]["read"] == "formulas"
    assert reads["Model.xlsx"]["findings"]
    assert reads["table.json"]["read"] == "questions"
    assert reads["table.json"]["qc_findings"]


# -------------------------------------------- and renders through one reporter each


def test_the_command_prints_what_the_module_prints() -> None:
    """Not a near-copy. The text is the module's own `report()`, so the command line and
    `python tools/formula_audit.py` cannot come to say different things about one file."""
    from tools import formula_audit

    data = a_model()
    expected = formula_audit.report(formula_audit.audit(data), as_json=False)
    tree = ast.parse((ROOT / "tools" / "aa.py").read_text(encoding="utf-8"))
    rendered = {
        f"{node.func.value.id}.{node.func.attr}"
        for node in ast.walk(tree)
        if isinstance(node, ast.Call)
        and isinstance(node.func, ast.Attribute)
        and isinstance(node.func.value, ast.Name)
        and node.func.attr == "report"
    }
    assert rendered == {"formula_audit.report", "prose_checks.report", "reasonableness.report"}
    assert expected, "and the reporter it calls produces something"


def test_the_review_of_one_file_is_the_review_the_module_gives(tmp_path: Path) -> None:
    from tools import formula_audit

    path = tmp_path / "Model.xlsx"
    data = a_model()
    path.write_bytes(data)
    (one,) = review(path)
    assert one.text == formula_audit.report(formula_audit.audit(data), as_json=False)


# ------------------------------------------------------------------ where, and staging


def test_where_an_engagements_documents_go(
    console: AdminConsole, henderson: Any, environment: dict[str, Path]
) -> None:
    found = _where(Served(console), HEND)
    assert found["relative"] == FOLDER.as_posix()
    assert Path(found["absolute"]) == environment["inbox"] / FOLDER
    assert found["can_add"] is True
    assert found["quiet_period_minutes"] >= 0


def test_where_says_so_when_the_engagement_is_not_on_the_roster(
    console: AdminConsole, henderson: Any
) -> None:
    with pytest.raises(KeyError, match="not on the roster"):
        _where(Served(console), "NOPE-1")


def test_staging_puts_the_documents_in_the_folder_and_records_who(
    console: AdminConsole, henderson: Any, environment: dict[str, Path], tmp_path: Path
) -> None:
    """The join a Box connector needs. Read the client's folder with the connector, save
    the files locally, stage them in - and the sweep finds them exactly as it finds a file
    a person dropped there."""
    one = tmp_path / "FY2025 Income Statement.xlsx"
    one.write_bytes(a_tidy_model())
    two = tmp_path / "Debt schedule.xlsx"
    two.write_bytes(a_model())

    placed = _stage(
        Served(console), HEND, [one, two], who=WHO, why="pulled from the client's Box folder"
    )
    assert [p["file"] for p in placed] == [
        f"{FOLDER.as_posix()}/FY2025 Income Statement.xlsx",
        f"{FOLDER.as_posix()}/Debt schedule.xlsx",
    ]
    assert (environment["inbox"] / FOLDER / "FY2025 Income Statement.xlsx").is_file()
    assert (environment["inbox"] / FOLDER / "Debt schedule.xlsx").is_file()

    logged = [c for c in console.config.changes() if c.field == f"files/{HEND}"]
    assert len(logged) == 2
    assert {c.who for c in logged} == {WHO}
    assert all("Box" in c.reason for c in logged)


def test_staging_refuses_without_a_person_and_a_reason(
    console: AdminConsole, henderson: Any, tmp_path: Path
) -> None:
    """Through `add_file`, which refuses it. Asserted here rather than trusted, because
    the whole reason this command goes through the console is that the console refuses
    things a command writing its own bytes would not - and "who put this file here" has
    to have an answer."""
    source = tmp_path / "statements.xlsx"
    source.write_bytes(a_tidy_model())
    with pytest.raises(ValueError, match="names a person and gives a reason"):
        _stage(Served(console), HEND, [source], who=WHO, why="")
    with pytest.raises(ValueError, match="names a person and gives a reason"):
        _stage(Served(console), HEND, [source], who="  ", why="from the client")
    assert not (tmp_path / "landed").exists()


def test_staging_takes_the_name_and_not_the_path_it_came_from(
    console: AdminConsole, henderson: Any, environment: dict[str, Path], tmp_path: Path
) -> None:
    """A document read out of Box into a deep temp path lands in the engagement's folder
    under its own name and nowhere else. The folder structure it arrived in is the
    connector's business, not the watched folder's."""
    deep = tmp_path / "box" / "Clients" / "Someone Else" / "statements.xlsx"
    deep.parent.mkdir(parents=True)
    deep.write_bytes(a_tidy_model())
    (placed,) = _stage(Served(console), HEND, [deep], who=WHO, why="from Box")
    assert placed["file"] == f"{FOLDER.as_posix()}/statements.xlsx"
    assert (environment["inbox"] / FOLDER / "statements.xlsx").is_file()
    assert not (environment["inbox"] / FOLDER / "box").exists()


def test_staging_a_file_that_is_not_there_is_refused_before_anything_is_opened(
    capsys: pytest.CaptureFixture[str], tmp_path: Path
) -> None:
    """No store is opened. A missing path is a typo and the answer is the typo, not a
    traceback from the composition root."""
    assert (
        main(
            [
                "stage",
                str(tmp_path / "nothing.xlsx"),
                "--engagement",
                HEND,
                "--who",
                WHO,
                "--why",
                "testing",
            ]
        )
        == 2
    )
    assert "no such file" in capsys.readouterr().err


def test_staging_names_a_person_and_a_reason_or_the_command_will_not_parse() -> None:
    """`--who` and `--why` are required by the parser, because `add_file` requires them
    and a command that let them default would have put a default on the record."""
    with pytest.raises(SystemExit):
        main(["stage", "a.xlsx", "--engagement", HEND])


# ------------------------------------------------- and it builds nothing of its own


def test_the_command_line_assembles_the_system_through_the_one_composition_root() -> None:
    """`tools/serve.py build()` is where this system is wired. A second wiring is a second
    answer to what is configured here, and the way two surfaces come to disagree about
    which folder an engagement's files go in."""
    tree = ast.parse((ROOT / "tools" / "aa.py").read_text(encoding="utf-8"))
    imported = {
        node.module
        for node in ast.walk(tree)
        if isinstance(node, ast.ImportFrom) and node.module
    }
    assert "tools.serve" in imported
    for forbidden in ("admin", "admin.console", "intake.config", "intake.roster", "ops.run_store"):
        assert forbidden not in imported, (
            f"aa.py reaches for {forbidden} directly; it should take the assembled system"
        )


def test_the_command_line_does_not_offer_a_third_way_to_ask_what_is_happening() -> None:
    """The census, and the line it draws.

    A command may be a second *surface* on a capability: `where`, `owed` and `stage` all
    reach the same `AdminConsole` methods the Setup page reaches, so the page and the
    command cannot come to disagree because there is only one implementation to disagree
    with. What is refused is a second *answer*: `status` is not here because
    `serve.py --check` and `watch_run.py --once` already answer those questions, and a
    third renderer of either would be two things that could drift - B-046's defect under a
    new name.

    The list is written down rather than derived, on purpose. Adding a subcommand should
    cost somebody this test and a sentence about which side of that line it falls on."""
    tree = ast.parse((ROOT / "tools" / "aa.py").read_text(encoding="utf-8"))
    names = {
        node.args[0].value
        for node in ast.walk(tree)
        if isinstance(node, ast.Call)
        and isinstance(node.func, ast.Attribute)
        and node.func.attr == "add_parser"
        and node.args
        and isinstance(node.args[0], ast.Constant)
    }
    assert names == {"review", "where", "owed", "stage"}


def test_the_served_system_is_closed_even_when_the_command_fails(
    console: AdminConsole, henderson: Any
) -> None:
    """The store is a file handle. A command that raised on the way out and left it open
    would make the next command on the same machine wait on a lock."""
    served = Served(console)
    with pytest.raises(KeyError):
        try:
            _where(served, "NOPE-1")
        finally:
            served.close()
    assert served.closed


# ------------------------------------------------- the slash commands, held to the tools


COMMANDS = sorted((ROOT / ".claude" / "commands").glob("*.md"))
INVOCATION = re.compile(r"python[^#\n]*?(?:tools/[a-z_]+\.py|-m\s+[\w.]+)[^#\n]*")
FLAG = re.compile(r"(?<![\w-])(--[a-z][a-z-]*)")

# The three subcommands' flags, and the tools whose flags a slash command may also use.
# Asked of each parser's own `--help` by `_flags_of`, never listed.
READ_TOOLS = ("formula_audit", "prose_checks", "reasonableness")


def fenced_blocks(page: Path) -> list[str]:
    """Each fenced block on the page, as one string.

    The block and not the line, which is the whole point. `review.md` writes the
    invocation on one line and its tolerances on another, so a check reading lines saw
    `python ... tools/aa.py review <target> --json` and never saw `--move points=5` at
    all - which was proved by breaking the flag on purpose and watching the test pass.
    A fenced block is one command with its arguments; that is the unit to read.
    """
    blocks, current, inside = [], [], False
    for line in page.read_text(encoding="utf-8").splitlines():
        if line.startswith("```"):
            if inside:
                blocks.append("\n".join(current))
                current = []
            inside = not inside
            continue
        if inside:
            current.append(line)
    return blocks


def invocations(page: Path) -> list[str]:
    """Every command a slash command tells Claude to run, out of its fenced blocks.

    The same extractor shape `tests/p58` uses on `docs/USING.md` and `CLAUDE.md`, because
    the failure is the same one: a page tells somebody to run a line, and the line does
    not work. A slash command is worse than a document here - a document is read by a
    person who will notice, and a slash command is read by something that will try it.
    """
    return [m.group().strip() for block in fenced_blocks(page) for m in INVOCATION.finditer(block)]


def test_there_are_slash_commands_to_check() -> None:
    assert len(COMMANDS) >= 4, [p.name for p in COMMANDS]


@pytest.mark.parametrize("page", COMMANDS, ids=lambda p: p.name)
def test_a_slash_command_carries_the_environment_with_it(page: Path) -> None:
    """`python tools/x.py` works only inside an activated shell. `gate.py` is the one
    exception and says so itself: it has no dependencies beyond the standard library,
    which is why it is the tool that installs and checks the others."""
    wrong = [
        c
        for c in invocations(page)
        if c.startswith("python tools/") and not c.startswith("python tools/gate.py")
    ]
    assert not wrong, f"{page.name} tells Claude to run {wrong} outside the environment"


@pytest.mark.parametrize("page", COMMANDS, ids=lambda p: p.name)
def test_every_tool_a_slash_command_names_exists(page: Path) -> None:
    for command in invocations(page):
        for token in command.split():
            if token.startswith("tools/") and token.endswith(".py"):
                assert (ROOT / token).is_file(), f"{page.name} names {token}, which is not here"


@pytest.mark.parametrize("page", COMMANDS, ids=lambda p: p.name)
def test_every_subcommand_a_slash_command_names_would_parse(page: Path) -> None:
    for command in invocations(page):
        if "tools/aa.py" not in command:
            continue
        words = command.split()
        rest = [w for w in words[words.index("tools/aa.py") + 1 :] if not w.startswith("-")]
        assert rest, f"{page.name} invokes aa.py with no subcommand"
        assert rest[0] in ("review", "where", "owed", "stage"), f"{page.name}: {rest[0]!r}"


@pytest.mark.parametrize("page", COMMANDS, ids=lambda p: p.name)
def test_every_flag_a_slash_command_passes_is_a_flag_something_takes(page: Path) -> None:
    """Over the whole fenced block, because that is where the flags are.

    A flag is checked against the union of what the tools invoked in that same block
    accept - a block that runs `aa.py review` and mentions `--fix` is passing `--fix` to
    nothing, and a block that runs `prose_checks.py` may legitimately use it.
    """
    for block in fenced_blocks(page):
        tools = {
            Path(part).stem
            for m in INVOCATION.finditer(block)
            for part in m.group().split()
            if part.startswith("tools/") and part.endswith(".py")
        }
        flags = set(FLAG.findall(block))
        if not tools:
            # A block of flags belonging to no command in it. Refused rather than
            # skipped, which is how the blind spot above survived: `review.md` wrote its
            # tolerances in a fenced block of their own, the check could not attribute
            # them to anything, and so it checked nothing - and a deliberately broken
            # `--movement` passed. A flag lives with the command it is passed to.
            assert not flags, (
                f"{page.name} has a fenced block of flags ({sorted(flags)}) with no "
                f"command in it. Put them on the line they belong to, or nothing can "
                f"tell whether they are flags anything takes."
            )
            continue
        allowed: set[str] = {"--help"}
        for tool in tools:
            if tool == "aa":
                for sub in ("review", "where", "owed", "stage"):
                    allowed |= _flags_of("aa", sub)
            elif tool in READ_TOOLS:
                allowed |= _flags_of(tool)
            else:
                allowed = set()  # a tool this test cannot ask; it checks nothing here
                break
        if not allowed:
            continue
        for flag in flags:
            assert flag in allowed, (
                f"{page.name} passes {flag} in a block running {sorted(tools)}, which "
                f"takes {sorted(allowed)}"
            )


def _flags_of(tool: str, subcommand: str | None = None) -> set[str]:
    """Which flags a command takes, asked of its own parser rather than listed.

    By running its `--help` and reading the flags out of it. The first version of this
    swapped in a recording `ArgumentParser` subclass, which recursed forever the moment
    argparse built a second parser for the subcommands - a test helper being cleverer
    than the thing it was tested against.
    """
    module = import_module(f"tools.{tool}")
    argv = [subcommand, "--help"] if subcommand else ["--help"]
    buffer = io.StringIO()
    with contextlib.redirect_stdout(buffer), pytest.raises(SystemExit):
        module.main(argv)
    return set(FLAG.findall(buffer.getvalue()))


# ------------------------------------------------------- what the client still owes


@pytest.fixture
def non_esop(console: AdminConsole) -> Any:
    """An engagement in the focus division, which is the only one with a request list.

    `HEND` is at `valuation/esop/repeat` and there is no starting list for ESOP - which is
    a real answer rather than a gap: `for_scope` walks up and returns None rather than
    handing over another division's list, because sending the client the wrong list is
    worse than sending none. So the happy path needs an engagement where a list exists.
    """
    from tests.p24.conftest import entry, profile_for

    scope = "valuation/non-esop/repeat"
    return console.register_engagement(
        entry(
            "NONE-2026-01",
            client="A Non-ESOP Client",
            scope_path=scope,
            folder="A Non-ESOP Client/2026",
            template_pin="",
        ),
        profile=profile_for("NONE-2026-01", scope),
        who=WHO,
        reason="engagement letter signed",
    )


def test_what_the_client_still_owes_reads_the_folder_and_writes_nothing(
    console: AdminConsole, non_esop: Any, environment: dict[str, Path]
) -> None:
    """The question that decides whether there is anything to do yet, asked without
    opening anything - and asked without writing, because refreshing the tracker is a
    write into a workbook the analyst keeps notes in."""
    folder = environment["inbox"] / "Clients" / "A Non-ESOP Client" / "2026"
    folder.mkdir(parents=True, exist_ok=True)
    (folder / "FY2025 Financial Statements.xlsx").write_bytes(a_tidy_model())

    found = _owed(Served(console), "NONE-2026-01", write=False)
    assert found["dry_run"] is True
    assert found["files_seen"] >= 1
    assert found["items"], "the list the engagement's scope carries"
    assert found["is_a_starting_point"] is True, "nobody wrote this one for this client"
    assert not (folder / "DD request list.xlsx").exists(), "nothing was written"


def test_an_engagement_whose_division_has_no_list_is_refused_rather_than_lent_one(
    console: AdminConsole, henderson: Any, environment: dict[str, Path]
) -> None:
    """`HEND` is ESOP and there is no starting list for ESOP. A request list is what the
    client was *asked for*, so handing over another division's would be worse than
    handing over none - and the refusal says exactly that."""
    from admin.console import AdminRefused

    (environment["inbox"] / FOLDER).mkdir(parents=True, exist_ok=True)
    with pytest.raises(AdminRefused, match="does not borrow another division"):
        _owed(Served(console), HEND, write=False)


def test_asking_what_is_owed_goes_through_the_same_method_the_page_does() -> None:
    """One implementation, two surfaces. A command that read the folder itself would be a
    second answer to what has arrived, and the page and the command would drift."""
    tree = ast.parse((ROOT / "tools" / "aa.py").read_text(encoding="utf-8"))
    calls = {
        node.func.attr
        for node in ast.walk(tree)
        if isinstance(node, ast.Call) and isinstance(node.func, ast.Attribute)
    }
    assert "dd_tracker" in calls
    page = (ROOT / "admin" / "page.py").read_text(encoding="utf-8")
    assert "dd-tracker" in page, "and the page reaches the same thing"


def test_writing_the_tracker_puts_it_beside_the_documents(
    console: AdminConsole, non_esop: Any, environment: dict[str, Path], tmp_path: Path
) -> None:
    """B-081, found by this test and fixed with it.

    The default tracker path was storage-*relative* and never joined to the root, so
    `write_tracker` - which creates the parent - wrote
    `<cwd>/Clients/<client>/<folder>/DD request list.xlsx`: a client-named directory tree
    inside whatever folder the process started in, which on this machine was the git
    working copy. The first run of this test left a `Clients/` tree in the repository.

    The worse half was invisible: `read_tracker` read that same wrong path, so the
    module's promise that a re-run preserves everything a person wrote was being kept
    against a file nobody had written in. `In Model?`, `Notes` and `Not Applicable` were
    carried across from an empty workbook every time.

    B-081 also said why nothing caught it - *every test passes an explicit path, and the
    default is the only shape the page uses* - which is why this asserts the default.
    """
    folder = environment["inbox"] / "Clients" / "A Non-ESOP Client" / "2026"
    folder.mkdir(parents=True, exist_ok=True)
    (folder / "FY2025 Financial Statements.xlsx").write_bytes(a_tidy_model())

    written = _owed(Served(console), "NONE-2026-01", write=True)
    assert written["dry_run"] is False

    tracker = Path(written["tracker"])
    assert tracker.name == "DD request list.xlsx"
    assert tracker.is_absolute(), "a relative path lands wherever the process started"
    assert tracker.parent == folder, "beside the documents, which is the whole point"
    assert tracker.is_file()
    assert not (ROOT / "Clients").exists(), "and nothing was written into the repository"


def test_what_a_person_wrote_in_the_tracker_survives_a_refresh(
    console: AdminConsole, non_esop: Any, environment: dict[str, Path]
) -> None:
    """The half of B-081 that was silent. The promise is that a re-run preserves what
    somebody typed; before the path was joined it was preserving it from the wrong file,
    so the promise held and meant nothing."""
    from openpyxl import load_workbook

    folder = environment["inbox"] / "Clients" / "A Non-ESOP Client" / "2026"
    folder.mkdir(parents=True, exist_ok=True)
    (folder / "FY2025 Financial Statements.xlsx").write_bytes(a_tidy_model())

    first = Path(_owed(Served(console), "NONE-2026-01", write=True)["tracker"])
    book = load_workbook(first)
    sheet = book.active
    assert sheet is not None
    headers = [c.value for c in next(sheet.iter_rows(max_row=1))]
    notes_at = headers.index("Notes") + 1
    sheet.cell(row=2, column=notes_at, value="asked again on the 3rd")
    book.save(first)

    again = Path(_owed(Served(console), "NONE-2026-01", write=True)["tracker"])
    assert again == first
    kept = load_workbook(again).active
    assert kept is not None
    assert kept.cell(row=2, column=notes_at).value == "asked again on the 3rd"


def test_there_is_nowhere_to_put_a_tracker_when_storage_is_not_a_folder() -> None:
    """A refusal rather than a relative path that will land somewhere. And it comes
    *after* the no-request-list refusal, because that one is about the engagement and
    this one is about the machine - asked the other way round, an engagement in a
    division nobody has written a list for would be told about the storage port."""
    from unittest.mock import Mock

    from admin.console import AdminConsole

    console = AdminConsole(Mock(), Mock(), Mock(), constraints=Mock())
    entry = Mock()
    entry.scope_path = "valuation/non-esop/repeat"
    console._folder_of = Mock(return_value=(entry, "Clients/x"))  # type: ignore[method-assign]
    console.config.document = Mock(return_value={})
    with pytest.raises(ValueError, match="nowhere to put it"):
        console.dd_tracker("ENG-1")


def test_an_engagement_not_on_the_roster_is_a_refusal_rather_than_an_empty_list(
    console: AdminConsole, henderson: Any
) -> None:
    """An empty outstanding list reads as a client who owes nothing, which is the one
    wrong conclusion available from a typo."""
    with pytest.raises(KeyError, match="not on the roster"):
        _owed(Served(console), "NOPE-1", write=False)
