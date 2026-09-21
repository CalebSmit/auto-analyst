"""Word documents to test against, built rather than committed.

The real report has never been seen by anything in this repo, and a client's report could
not live here in any case - `fixtures/` is masked, synthetic or public, never a document
somebody sent. So the documents these tests run against are generated, and they are
generated to contain the things that actually break a copy:

* a **style the destination does not define**, because a paragraph carries a style id and
  a missing definition renders as Normal;
* a **numbered list**, because numbering lives in a different part of the file and a
  naive copy either loses the numbers or joins a list already in the destination;
* an **image**, because it is a separate relationship and an unremapped id points at
  whatever the destination's rId happened to be - usually the letterhead;
* a **bookmark**, because that is what binding by name is supposed to use.

Anything these do not contain is something the tests do not prove.
"""

from __future__ import annotations

import struct
import zlib
from io import BytesIO
from typing import Any

from docx import Document
from docx.oxml.ns import qn

__all__ = [
    "a_report",
    "appendix_source",
    "bookmark",
    "list_ids_in",
    "one_pixel_png",
    "text_of",
]


def one_pixel_png() -> bytes:
    """A valid PNG, built rather than committed so no binary fixture is needed."""

    def chunk(kind: bytes, body: bytes) -> bytes:
        return (
            struct.pack(">I", len(body))
            + kind
            + body
            + struct.pack(">I", zlib.crc32(kind + body) & 0xFFFFFFFF)
        )

    header = struct.pack(">IIBBBBB", 1, 1, 8, 2, 0, 0, 0)
    pixels = zlib.compress(b"\x00\xff\x00\x00")
    return (
        b"\x89PNG\r\n\x1a\n"
        + chunk(b"IHDR", header)
        + chunk(b"IDAT", pixels)
        + chunk(b"IEND", b"")
    )


def bookmark(paragraph: Any, name: str, number: int = 1) -> None:
    """Put a Word bookmark on a paragraph.

    Raw XML because python-docx has no API for it, and bookmarks are the anchor kind that
    matters most - they are what somebody marks deliberately.
    """
    start = paragraph._p.makeelement(qn("w:bookmarkStart"), {})
    start.set(qn("w:id"), str(number))
    start.set(qn("w:name"), name)
    end = paragraph._p.makeelement(qn("w:bookmarkEnd"), {})
    end.set(qn("w:id"), str(number))
    paragraph._p.insert(0, start)
    paragraph._p.append(end)


def a_report(*, bookmarks: bool = False, listed: bool = True) -> Any:
    """A report shaped like one being set up: headings, an appendix section each.

    With `bookmarks`, each appendix heading also carries a bookmark of the same name, so a
    test can prove a bookmark beats a heading rather than assuming it.

    With `listed` - the default - the report has a numbered list of its own, built the
    same way the source's is. That matters more than it looks: both documents then claim
    the same `numId`, which is the collision the remapping exists for. Without it the two
    ids happened not to clash and a test would pass while proving nothing.
    """
    document = Document()
    document.add_heading("Valuation of a Private Company", 0)
    document.add_paragraph("The body of the report, which nothing here touches.")

    if listed:
        mine = _its_own_list(document)
        for point in ("A finding of the report.", "Another finding."):
            _number_it(document.add_paragraph(point, style="List Number"), mine)

    for letter, title in (
        ("C", "Economic Conditions"),
        ("E", "Summary of Major Market Indices"),
        ("F", "Summary of M&A Markets"),
        ("H", "Certification"),
    ):
        name = f"Appendix {letter} - {title}"
        heading = document.add_heading(name, 1)
        if bookmarks:
            bookmark(heading, name, number=ord(letter))
        document.add_paragraph(f"Last year's {title.lower()}, still in place.")
        document.add_heading("A sub-heading inside the appendix", 2)
        document.add_paragraph("Which belongs to the appendix above it, not to the next.")
    return document


def appendix_source(*, listed: bool = True, image: bool = True, style: bool = True) -> Any:
    """The document an appendix's content arrives in: a .docx he points a button at."""
    document = Document()
    document.add_paragraph("Economic conditions in the period under review.")

    if style:
        made = document.styles.add_style("MarketCommentary", 1)
        made.font.italic = True
        document.add_paragraph("A commentary paragraph.", style="MarketCommentary")

    if listed:
        num_id = _its_own_list(document)
        for point in ("Real GDP growth slowed.", "Rates held.", "Spreads widened."):
            paragraph = document.add_paragraph(point, style="List Number")
            _number_it(paragraph, num_id)

    if image:
        document.add_picture(BytesIO(one_pixel_png()))

    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "Index"
    table.cell(0, 1).text = "Change"
    table.cell(1, 0).text = "A broad market index"
    table.cell(1, 1).text = "3.1%"
    return document


def _its_own_list(document: Any) -> str:
    """Give the source document a numbering definition of its own, and return its numId.

    Deliberately not `style="List Number"` alone. A style-only list carries no `numId` on
    the paragraph, so the remapping path - the part that stops Appendix F's first bullet
    becoming item 7 of a list already in the report - would never run and the test would
    prove nothing about it.
    """
    numbering = document.part.numbering_part.element
    abstracts = numbering.findall(qn("w:abstractNum"))
    abstract_id = str(
        max((int(a.get(qn("w:abstractNumId")) or 0) for a in abstracts), default=0) + 1
    )
    abstract = numbering.makeelement(qn("w:abstractNum"), {})
    abstract.set(qn("w:abstractNumId"), abstract_id)
    level = numbering.makeelement(qn("w:lvl"), {})
    level.set(qn("w:ilvl"), "0")
    for tag, value in (("w:start", "1"), ("w:numFmt", "decimal"), ("w:lvlText", "%1.")):
        node = numbering.makeelement(qn(tag), {})
        node.set(qn("w:val"), value)
        level.append(node)
    abstract.append(level)
    first_num = numbering.find(qn("w:num"))
    if first_num is not None:
        first_num.addprevious(abstract)
    else:
        numbering.append(abstract)

    num_id = str(
        max((int(n.get(qn("w:numId")) or 0) for n in numbering.findall(qn("w:num"))), default=0) + 1
    )
    num = numbering.makeelement(qn("w:num"), {})
    num.set(qn("w:numId"), num_id)
    ref = numbering.makeelement(qn("w:abstractNumId"), {})
    ref.set(qn("w:val"), abstract_id)
    num.append(ref)
    numbering.append(num)
    return num_id


def _number_it(paragraph: Any, num_id: str) -> None:
    """Put the numbering reference on the paragraph itself, the way Word does."""
    pPr = paragraph._p.get_or_add_pPr()
    numPr = pPr.makeelement(qn("w:numPr"), {})
    ilvl = pPr.makeelement(qn("w:ilvl"), {})
    ilvl.set(qn("w:val"), "0")
    num = pPr.makeelement(qn("w:numId"), {})
    num.set(qn("w:val"), num_id)
    numPr.append(ilvl)
    numPr.append(num)
    pPr.append(numPr)


def list_ids_in(document: Any) -> list[str]:
    """Every numId a paragraph in the body actually references."""
    return [
        n.get(qn("w:val"))
        for n in document.element.body.iter(qn("w:numId"))
        if n.getparent().tag == qn("w:numPr")
    ]


def text_of(document: Any) -> str:
    """Every bit of text in the body, for asserting what arrived and what did not."""
    W = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"
    return "\n".join(
        "".join(t.text or "" for t in block.iter(f"{W}t"))
        for block in document.element.body.iterchildren()
        if isinstance(block.tag, str)
    )
